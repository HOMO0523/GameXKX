"""Build and verify a timecoded, self-contained tool audio production sample."""
import json
import re
import wave
import subprocess
import hashlib
import shutil
import zipfile
import math
from pathlib import Path

import numpy as np
import imageio_ffmpeg
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
EVIDENCE = ROOT / "Saved/Codex/ToolSfxRequirementSample-20260914"
CLIPS = ["enhance02", "combine02", "reforge_generate01", "reforge_accept01", "reforge_keep01"]
FFMPEG = imageio_ffmpeg.get_ffmpeg_exe()
PACKAGE = ROOT / "Deliverables/音效需求/14_工具操作_v001"
SCENES = [
    ("强化成功", "强化等级 +1 → +2", "强化石 39 → 38", "结果生效时，短促地确认完成。", True),
    ("九件合成", "九格清空 · 新装备入包", "本次材料为普通品质", "整批只响一次，不按九件连响。", True),
    ("洗炼 · 生成候选", "消耗洗炼砂 · 展示新旧对比", "属性还没有替换", "仅保留按钮声，工具完成音不响。", False),
    ("洗炼 · 采用新属性", "选择新属性 · 替换生效", "物理伤害 +7.76%", "选择落定时，复用同一条工具音。", True),
    ("洗炼 · 保留原属性", "再次洗炼后，选择原属性", "物理伤害 +7.76% 保持", "仅保留按钮声，工具完成音不响。", False),
]


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def run(args, **kwargs):
    result = subprocess.run([FFMPEG, "-hide_banner", "-loglevel", "error", "-y", *args], capture_output=True, **kwargs)
    if result.returncode:
        raise RuntimeError(result.stderr.decode("utf-8", "replace")[-3500:])
    return result.stdout


def load(path):
    return json.loads(path.read_text(encoding="utf-8-sig"))


def read_wav(path):
    with wave.open(str(path)) as wav:
        assert wav.getsampwidth() == 2
        channels, rate = wav.getnchannels(), wav.getframerate()
        data = np.frombuffer(wav.readframes(wav.getnframes()), "<i2").reshape(-1, channels).astype(np.float64) / 32768
    return data, rate


def match_wave(recording, reference):
    x = recording.mean(axis=1)[::4]
    y = reference.mean(axis=1)[::4]
    n = 1 << (len(x) + len(y) - 1).bit_length()
    corr = np.fft.irfft(np.fft.rfft(x, n) * np.fft.rfft(y[::-1], n), n)[len(y)-1:len(x)]
    square_sum = np.r_[0.0, np.cumsum(x*x)]
    energy = square_sum[len(y):] - square_sum[:-len(y)]
    denom = np.sqrt(np.maximum(0, energy) * np.sum(y*y))
    scores = np.divide(corr, denom, out=np.zeros_like(corr), where=denom > .0001)
    at = int(scores.argmax())
    return {"correlation": float(scores[at]), "audio_seconds": at / 12000}


def analyze():
    references = {cue: read_wav(ROOT / f"SourceArt/Audio/Essential/Waves/SFX_{cue}_v01.wav")[0] for cue in ("Tool", "Button")}
    result = []
    for name in CLIPS:
        start = load(EVIDENCE / (name + "-audio-start.json"))
        stop = load(EVIDENCE / (name + "-audio-stop.json"))
        capture = load(EVIDENCE / (name + "-capture.json"))
        log = (EVIDENCE / (name + "-ffmpeg.log")).read_text(encoding="utf-8")
        video_epoch = float(re.search(r"start: ([0-9.]+)", log)[1])
        audio, rate = read_wav(EVIDENCE / (name + ".wav"))
        matches = {cue: match_wave(audio, data) for cue, data in references.items()}
        result.append({"name": name, "video_epoch": video_epoch, "audio_epoch": start["wall_before"],
                       "audio_seconds": len(audio) / rate, "wall_seconds": stop["wall_time"] - start["wall_before"],
                       "click_video_seconds": capture["click_before"] - video_epoch,
                       "peak": float(np.abs(audio).max()), "played": stop["audio"]["played"], "matches": matches})
    (EVIDENCE / "audio-analysis.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return result


def visual_events(audio_analysis):
    # These readback areas exclude the clicked button and pointer glow. They
    # detect the resulting value/grid/comparison change in the actual recording.
    rois = [(958, 388, 238, 56), (962, 152, 230, 224), (964, 398, 232, 104),
            (964, 398, 232, 104), (966, 522, 230, 80)]
    events = []
    for index, (record, roi) in enumerate(zip(audio_analysis, rois)):
        x, y, width, height = roi
        raw = subprocess.run([FFMPEG, "-v", "error", "-i", str(EVIDENCE / (record["name"] + "-raw.mp4")),
                              "-vf", f"crop={width}:{height}:{x}:{y}", "-pix_fmt", "gray", "-f", "rawvideo", "-"],
                             check=True, capture_output=True).stdout
        frames = np.frombuffer(raw, np.uint8).reshape(-1, height, width)
        approximate = int(record["click_video_seconds"] * 60)
        baseline = np.median(frames[approximate-55:approximate-20], axis=0)
        difference = np.mean(np.abs(frames.astype(np.float32) - baseline), axis=(1, 2))
        candidates = [f for f in range(approximate-5, min(approximate+80, len(frames)-10))
                      if np.all(difference[f:f+8] > 2.0)]
        assert candidates, (record["name"], float(difference.max()))
        event_frame = candidates[0]
        start_frame = event_frame - 120
        assert start_frame >= 0 and start_frame + 300 <= len(frames), (record["name"], len(frames), event_frame)
        cue = "Tool" if record["played"].get("Tool") else "Button"
        match = record["matches"][cue]
        assert match["correlation"] > .98
        sound_raw_seconds = record["audio_epoch"] - record["video_epoch"] + match["audio_seconds"]
        sound_final_seconds = index * 5 + sound_raw_seconds - start_frame / 60
        events.append(dict(record, index=index, cue=cue, raw_frame_count=len(frames),
                           visual_event_raw_frame=event_frame, trim_start_frame=start_frame,
                           visual_event_final_frame=index*300+120, sound_final_seconds=sound_final_seconds,
                           sound_final_frame=round(sound_final_seconds*60),
                           sound_vs_visual_ms=(sound_raw_seconds-event_frame/60)*1000,
                           roi=list(roi)))
    (EVIDENCE / "timing-analysis.json").write_text(json.dumps(events, ensure_ascii=False, indent=2), encoding="utf-8")
    return events


def sound_events(events):
    tool_reference, _ = read_wav(ROOT / "SourceArt/Audio/Essential/Waves/SFX_Tool_v01.wav")
    button_reference, _ = read_wav(ROOT / "SourceArt/Audio/Essential/Waves/SFX_Button_v01.wav")
    sounds = []
    for record in events:
        clips = [record["cue"]]
        matches = dict(record["matches"])
        if record["cue"] == "Tool":
            data, _ = read_wav(EVIDENCE / (record["name"] + ".wav"))
            y = tool_reference[:, 0]
            at = round(matches["Tool"]["audio_seconds"] * 48000)
            gain = np.dot(data.mean(axis=1)[at:at+len(y)], y) / np.dot(y, y)
            residual = data.copy()
            residual[at:at+len(y)] -= gain * tool_reference
            matches["Button"] = match_wave(residual, button_reference)
            assert matches["Button"]["correlation"] > .90
            clips.append("Button")
        for cue in clips:
            match = matches[cue]
            seconds = record["index"] * 5 + record["audio_epoch"] - record["video_epoch"] + match["audio_seconds"] - record["trim_start_frame"] / 60
            sounds.append({"scene": record["index"] + 1, "cue": cue, "file": f"SFX_{cue}_v01.wav",
                           "seconds": seconds, "frame_contains_start": math.floor(seconds * 60),
                           "correlation": match["correlation"],
                           "method": "overlap residual match" if cue == "Button" and record["cue"] == "Tool" else "waveform match"})
    (EVIDENCE / "sound-events.json").write_text(json.dumps(sounds, ensure_ascii=False, indent=2), encoding="utf-8")
    return sounds


def ass_time(seconds):
    centis = round(seconds * 100)
    return f"{centis//360000}:{centis//6000%60:02}:{centis//100%60:02}.{centis%100:02}"


def subtitles():
    lines = ["[Script Info]", "ScriptType: v4.00+", "PlayResX: 1280", "PlayResY: 720", "WrapStyle: 2", "",
             "[V4+ Styles]", "Format: Name, Fontname, Fontsize, PrimaryColour, SecondaryColour, OutlineColour, BackColour, Bold, Italic, Underline, StrikeOut, ScaleX, ScaleY, Spacing, Angle, BorderStyle, Outline, Shadow, Alignment, MarginL, MarginR, MarginV, Encoding"]
    styles = [("Tag", 18, "&H00BDB1A2", 0), ("Title", 40, "&H00F8F4EC", -1),
              ("Scene", 29, "&H00F8F4EC", -1), ("Body", 21, "&H00D4C9BC", 0),
              ("Rule", 23, "&H0087D7A3", -1), ("MutedRule", 23, "&H009CBFE3", -1),
              ("Small", 16, "&H00ABA092", 0)]
    for name, size, color, bold in styles:
        lines.append(f"Style: {name},Microsoft YaHei,{size},{color},&H00000000,&H00000000,&H00000000,{bold},0,0,0,100,100,0,0,1,0,0,7,0,0,0,1")
    lines += ["", "[Events]", "Format: Layer, Start, End, Style, Name, MarginL, MarginR, MarginV, Effect, Text"]
    def text(start, end, style, x, y, label):
        lines.append(f"Dialogue: 0,{ass_time(start)},{ass_time(end)},{style},,0,0,0,,{{\\pos({x},{y})}}{label}")
    text(0, 25, "Tag", 48, 43, "SFX 14  /  TOOL OPERATIONS")
    text(0, 25, "Title", 48, 85, "工具操作 · 音效样包")
    text(0, 25, "Small", 48, 142, "v001   ·   60 fps   ·   实机画面 / 现用音频")
    text(0, 25, "Small", 520, 144, "背包 / 消耗与产物")
    text(0, 25, "Small", 48, 597, "整包只制作 1 条工具音")
    text(0, 25, "Small", 48, 625, "按钮声是伴随参考，已另附原文件")
    for i, (title, action, value, requirement, tool) in enumerate(SCENES):
        begin, end = i * 5, i * 5 + 5
        text(begin, end, "Tag", 48, 205, f"{i+1:02} / 05")
        text(begin, end, "Scene", 48, 246, title)
        text(begin, end, "Body", 48, 306, action)
        text(begin, end, "Body", 48, 342, value)
        text(begin, end, "Rule" if tool else "MutedRule", 48, 414, "工具完成音 × 1" if tool else "工具完成音 × 0")
        # Manual line breaks keep every requirement inside the left column.
        desc = requirement.replace("，", "，\\N")
        text(begin, end, "Body", 48, 461, desc)
        text(begin, end, "Small", 520, 640, f"画面结果 {begin+2:06.3f} s   /   F{i*300+120:04}")
        text(begin+2, begin+2.8, "Tag", 520, 675, "结果已生效" if tool else "预览 / 保留，无工具完成音")
    (EVIDENCE / "sample.ass").write_text("\n".join(lines), encoding="utf-8-sig")


def make_videos(events):
    PACKAGE.mkdir(parents=True, exist_ok=True)
    rendered = []
    pcm_segments = []
    for record in events:
        name = record["name"]
        start = record["trim_start_frame"] / 60
        audio_start = start + record["video_epoch"] - record["audio_epoch"]
        target = EVIDENCE / (name + "-cut.mp4")
        run(["-i", str(EVIDENCE / (name + "-raw.mp4")),
             "-vf", f"trim=start_frame={record['trim_start_frame']}:end_frame={record['trim_start_frame']+300},setpts=PTS-STARTPTS",
             "-an", "-c:v", "libx264", "-crf", "17", "-preset", "fast", "-threads", "4",
             "-pix_fmt", "yuv420p", "-r", "60", "-fps_mode", "cfr",
             "-t", "5", str(target)])
        rendered.append(target)
        audio, rate = read_wav(EVIDENCE / (name + ".wav"))
        at = round(audio_start * rate)
        assert rate == 48000 and 0 <= at < len(audio)
        part = audio[at:at+240000]
        if len(part) < 240000:
            assert 240000-len(part) < 24000 and float(np.abs(part[-7200:]).max()) < .003
            part = np.pad(part, ((0,240000-len(part)),(0,0)))
        pcm_segments.append(np.rint(part*32768).astype("<i2"))
    # Concatenate PCM before AAC encoding: separately encoded AAC segments can
    # introduce padding at every cut and drift from the video frame timebase.
    with wave.open(str(EVIDENCE / "joined-audio.wav"), "wb") as output:
        output.setparams((2, 2, 48000, 0, "NONE", "not compressed"))
        output.writeframes(np.concatenate(pcm_segments).tobytes())
    concat = EVIDENCE / "concat.txt"
    concat.write_text("\n".join(f"file '{x.name}'" for x in rendered), encoding="utf-8")
    run(["-f", "concat", "-safe", "0", "-i", str(concat), "-c", "copy", str(EVIDENCE / "joined.mp4")])
    subtitles()
    shutil.copy2(Path("C:/Windows/Fonts/consola.ttf"), EVIDENCE / "counter.ttf")
    counter = r"drawtext=fontfile=counter.ttf:fontcolor=0xaab8bf:fontsize=19:x=48:y=673:text='F%{eif\:n\:d\:4}  /  60 fps'"
    run(["-i", "joined.mp4", "-i", "joined-audio.wav", "-map", "0:v:0", "-map", "1:a:0", "-vf", "subtitles=sample.ass," + counter, "-c:v", "libx264", "-preset", "fast", "-crf", "17", "-threads", "4",
         "-r", "60", "-fps_mode", "cfr", "-pix_fmt", "yuv420p", "-c:a", "aac", "-b:a", "192k", "-t", "25", "-movflags", "+faststart",
         str(PACKAGE / "01_演示_现有效果.mp4")], cwd=EVIDENCE)
    run(["-i", str(PACKAGE / "01_演示_现有效果.mp4"), "-map", "0:v:0", "-c:v", "copy", "-an", "-movflags", "+faststart", str(PACKAGE / "02_演示_无声版.mp4")])
    print("Videos rendered", flush=True)


def copy_sources():
    manifest = load(ROOT / "SourceArt/Audio/Essential/manifest.json")
    records = [r for r in manifest["files"] if r["cue"] in ("Tool", "Button")]
    folder = PACKAGE / "参考音频"
    provenance = PACKAGE / "来源与授权"
    folder.mkdir(exist_ok=True)
    provenance.mkdir(exist_ok=True)
    for r in records:
        source = ROOT / r["file"]
        assert sha(source) == r["sha256"]
        target = folder / source.name
        shutil.copy2(source, target)
        pack = provenance / r["pack"]
        pack.mkdir(exist_ok=True)
        original = ROOT / r["source_file"]
        assert sha(original) == r["source_sha256"]
        shutil.copy2(original, pack / original.name)
        for name in ("LICENSE-1.txt", "source-page.html", "source-manifest.json"):
            shutil.copy2(ROOT / "SourceArt/Audio/Essential/Licenses" / r["pack"] / name, pack / name)
        r["project_file"] = r["file"]
        r["file"] = target.relative_to(PACKAGE).as_posix()
        r["project_source_file"] = r["source_file"]
        r["source_file"] = (pack / original.name).relative_to(PACKAGE).as_posix()
        r["package_role"] = "本包制作目标的现用参考" if r["cue"] == "Tool" else "伴随按钮声参考，不在本包新增制作"
    (provenance / "音频清单.json").write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
    return records


def workbook(events, sounds, records):
    wb = Workbook()
    overview = wb.active
    overview.title = "制作需求"
    overview.append(["工具操作音效 · 制作 / 调整样包", "v001"])
    overview.append(["制作数量", "1 条；强化、合成、分解、镶嵌/拆卸、采用洗炼新属性共用。样片展示其中三个典型工具。"])
    overview.append(["期望效果", "短、清楚、有操作完成感；轻量敲击主体，可带很短的灵力尾音。"])
    overview.append(["调整方向", "保留清楚的起音；可弱化采矿/敲石的具象感，使其适合多种工具。此项为建议方向，尚未人工听审定稿。"])
    overview.append(["建议时长", "约 0.25–0.55 秒；一次性播放，不循环。现用参考为 0.520 秒。"])
    overview.append(["触发原则", "以实际操作成功/选择生效为准；一批九件只响一次。生成洗炼候选和保留原属性不触发工具完成音。"])
    overview.append(["与按钮声关系", "按钮声作轻量按下反馈；工具声承担当次操作已完成的反馈，避免两个起音互相遮盖。"])
    overview.append(["建议交付", "1 条 WAV，48 kHz，单声道，24-bit PCM；命名 SFX_Tool_v02.wav。保留尾音自然结束，不带整段前置静音。"])
    overview.append(["参考格式", "现有 WAV 是 48 kHz / 单声道 / 16-bit PCM，由 OGG 源处理；不能当作原生 24-bit 母带。"])
    overview.append(["视频规格", "1280×720，60 fps 固定帧率，25.000 秒，1500 帧；从 F0000 计到 F1499。"])
    overview.append(["时间约定", "秒/帧均相对最终 v001 视频；区间含起点、不含终点。F120 表示 2.000–2.0167 秒。"])
    overview.append(["声画同步", "保留现用实录声轨，用采集时钟对齐；实测起音与画面更新相差约 5–25 ms。制作目标落点以画面结果帧为准，采集精度约 2 帧。"])
    overview.append(["剪接说明", "五段各 5 秒，段落间硬切。最后一段是再次付费生成候选后的另一种选择，省略了第二次生成过程。"])
    overview.append(["状态", "样包待评审；当前参考声音已接入，下一版新音效尚未制作、导入或听审。"])
    for filename in ("01_演示_现有效果.mp4", "02_演示_无声版.mp4"):
        overview.append(["演示视频", filename]); overview.cell(overview.max_row, 2).hyperlink = filename
    timeline = wb.create_sheet("时间轴需求")
    timeline.append(["编号", "场景", "片段秒区间", "片段帧区间", "画面结果秒", "结果帧", "工具音需求", "当前使用音频", "实际起音秒（约）", "起音所在帧", "制作 / 调整要点"])
    for i, scene in enumerate(SCENES):
        these = sorted([s for s in sounds if s["scene"] == i+1], key=lambda s:s["seconds"])
        for j, sound in enumerate(these):
            timeline.append([f"14-{i+1:02}-{j+1}", scene[0], f"[{i*5:.3f}, {(i+1)*5:.3f})", f"[F{i*300:04}, F{(i+1)*300:04})",
                             i*5+2, i*300+120, "播放1次" if scene[4] else "不播放工具音", sound["file"], round(sound["seconds"], 4), sound["frame_contains_start"],
                             scene[3] if sound["cue"] == "Tool" else "伴随按钮反馈；沿用现有按钮声，不另增工具声音类别。"])
    assets = wb.create_sheet("音频清单")
    assets.append(["文件", "用途", "时长秒", "采样率", "声道", "位深", "原文件", "作者", "授权", "来源", "包内位置", "SHA256"])
    for r in records:
        assets.append([r["name"] + ".wav", r["package_role"], r["seconds"], r["sample_rate"], r["channels"], r["pcm_bits"],
                       Path(r["source_file"]).name, r["author"], r["license"], r["source"], r["file"], r["sha256"]])
        assets.cell(assets.max_row, 10).hyperlink = r["source"]
        assets.cell(assets.max_row, 11).hyperlink = r["file"]
    scene_sheet = wb.create_sheet("画面段落")
    scene_sheet.append(["序号", "内容", "起秒", "止秒（不含）", "起帧", "止帧（不含）", "画面结果帧", "应发生的工具反馈"])
    for i, scene in enumerate(SCENES):
        scene_sheet.append([i+1, scene[0], i*5, (i+1)*5, i*300, (i+1)*300, i*300+120, scene[3]])
    widths = {"制作需求": [24, 112], "时间轴需求": [14, 24, 22, 22, 16, 12, 22, 30, 22, 16, 58],
              "音频清单": [30, 46, 13, 13, 10, 10, 28, 15, 15, 48, 45, 69], "画面段落": [10, 30, 14, 18, 14, 18, 18, 62]}
    for sheet in wb:
        sheet.freeze_panes = "B2"; sheet.sheet_view.showGridLines = False
        sheet.auto_filter.ref = sheet.dimensions if sheet.title != "制作需求" else "A1:B1"
        for index, width in enumerate(widths[sheet.title], 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        for row in sheet:
            for cell in row:
                cell.font = Font(name="Microsoft YaHei", size=11, color="26343D")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = Border(bottom=Side(style="hair", color="DDE3E6"))
                if cell.row == 1:
                    cell.fill = PatternFill("solid", fgColor="26343D"); cell.font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
                elif cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F0F5F3")
                if cell.hyperlink:
                    cell.font = Font(name="Microsoft YaHei", size=11, color="207E73", underline="single")
            sheet.row_dimensions[row[0].row].height = 38 if row[0].row == 1 else (64 if sheet.title == "时间轴需求" else 48)
        sheet.page_setup.orientation = "landscape"; sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1; sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = "1:1"
    destination = PACKAGE / "03_时间轴需求.xlsx"
    wb.save(destination)
    check = load_workbook(destination, read_only=False)
    assert check.sheetnames == ["制作需求", "时间轴需求", "音频清单", "画面段落"]
    assert check["时间轴需求"].max_row == 9 and check["音频清单"].max_row == 3
    check.close()


def documents(events, sounds, records):
    text = ["工具操作音效样包 v001", "25.000秒 | 60 fps | 1280×720 | F0000起算", "区间含起点、不含终点。声音落点与画面区间分别列出。", ""]
    for i, scene in enumerate(SCENES):
        text += [f"{i*5:06.3f}–{(i+1)*5:06.3f} 秒 / F{i*300:04}–F{(i+1)*300:04}（末帧不含）  {scene[0]}",
                 f"画面结果：{i*5+2:06.3f} 秒 / F{i*300+120:04}", f"工具音需求：{scene[3]}"]
        for sound in sorted([s for s in sounds if s['scene'] == i+1], key=lambda s:s['seconds']):
            text.append(f"现用音频：{sound['file']} | 实录起音约 {sound['seconds']:.4f} 秒 / 所在帧 F{sound['frame_contains_start']:04}")
        text.append("")
    text += ["同一条 SFX_Tool 在三处复用；洗炼生成候选、保留原属性仅有按钮声。",
             "声音制作方向：短、清楚、有完成感；约0.25–0.55秒，单次，不循环。",
             "弱化采矿敲石的具象联想，可加很短的灵力收尾。此为制作建议，需人工听审。",
             "实录同步约2帧精度；制作目标建议对齐画面结果帧。重剪、变速或换帧率后必须重新标注。"]
    (PACKAGE / "04_时间轴速览.txt").write_text("\n".join(text), encoding="utf-8-sig")
    (PACKAGE / "00_先看这里.md").write_text("""# 工具操作 · 音效制作样包 v001

本包用于交给音效人员制作或调整 **1条共用工具完成音**。现有音频是参考，新版音效尚未制作或接入。按钮声是随片参考，继续单独复用，不增加本包的制作数量。

先看 `01_演示_现有效果.mp4`，再按 `03_时间轴需求.xlsx` 制作。`02_演示_无声版.mp4` 与有声版使用完全相同的视频流，可以直接配音。`04_时间轴速览.txt` 提供与截图示例相近的简版标注。

## 要做的声音

- 短、清楚、有完成感。轻量敲击主体，可以加很短的灵力尾音；建议弱化现用采矿音的具体材质感，让强化、合成、分解、镶嵌/拆卸与采用洗炼新属性能共用。
- 建议0.25–0.55秒，一次性，不循环。整批操作只响一次；起音不要与按钮声争抢。上述为制作方向，需人工听审确认。
- 只在操作成功或采用新属性生效时响。洗炼生成候选、保留原属性、材料不足、只放置物品、切模式均不应触发工具完成音。
- 建议交付 `SFX_Tool_v02.wav`，48kHz、单声道、24-bit PCM。现用参考是16-bit、由OGG源处理，不是24-bit录音母带。

## 视频与时间

25.000秒，1280×720，60fps固定帧率，共1500帧，从F0000到F1499。五段各5秒，段落之间硬切。每段画面结果位于该段第2秒：F0120、F0420、F0720、F1020、F1320。区间含起点、不含终点；帧F表示[F/60,(F+1)/60)秒。

声音使用当前游戏混音器实录，没有用参考WAV重新配音。用视频和音频采集时钟对齐，完整保留静音；视频仅裁取游戏纸面区域并添加说明。需求表同时列出**画面结果落点**与**当前声音实测起音**，后者约有2帧采集精度，请勿把小数位数当作亚帧精度承诺。制作目标建议对齐画面结果帧。

最后一段来自再次生成候选后的“保留原属性”操作，第二次生成过程已剪掉。素材来自临时测试场景；配置物品和切换工具由项目脚本准备，五个录制提交动作均为实际鼠标点击。样片只覆盖典型流程，不是全工具功能验收视频。

## 当前使用文件

|文件|本包角色|来源原文件|
|---|---|---|
|`参考音频/SFX_Tool_v01.wav`|待调整的共用工具音，0.520秒|Kenney Impact Sounds / impactMining_000.ogg|
|`参考音频/SFX_Button_v01.wav`|随片按钮音，约0.0662秒|Kenney Interface Sounds / click_001.ogg|

原始OGG、作者发布页快照、原包CC0授权原文和逐文件来源放在`来源与授权`。来源页面：[Impact Sounds](https://kenney.nl/assets/impact-sounds)、[Interface Sounds](https://kenney.nl/assets/interface-sounds)。这两份音频的授权不覆盖本项目画面、美术或其他内容。

## 验收建议

对照无声版按表配音，再与现有效果版比较操作落点和反馈力度。确认三次工具提交使用同一条声音、两个洗炼非提交场景没有工具完成音；连续操作时不刺耳、不产生拖长尾音。人工听审通过后再替换项目音频。

文件改名或重剪视频后应同步升版并重建时间轴。`05_文件核验.json`记录媒体规格与文件哈希，支持核对交付是否完整。
""", encoding="utf-8")


def verify_and_zip(events, sounds):
    audible = PACKAGE / "01_演示_现有效果.mp4"
    silent = PACKAGE / "02_演示_无声版.mp4"
    # Compare the encoded elementary video stream, not two container hashes.
    streams = []
    for video in (audible, silent):
        data = run(["-i", str(video), "-map", "0:v:0", "-c", "copy", "-bsf:v", "h264_mp4toannexb", "-f", "h264", "-"])
        streams.append(hashlib.sha256(data).hexdigest())
    assert streams[0] == streams[1]
    decoded = run(["-i", str(audible), "-map", "0:v:0", "-f", "framemd5", "-"])
    frame_lines = [line for line in decoded.decode().splitlines() if line and not line.startswith("#")]
    assert len(frame_lines) == 1500, len(frame_lines)
    final_visual_frames = []
    for record in events:
        x, y, width, height = record["roi"]
        raw = run(["-i", str(audible), "-vf", f"crop={width}:{height}:{x}:{y}", "-pix_fmt", "gray", "-f", "rawvideo", "-"])
        frames = np.frombuffer(raw, np.uint8).reshape(-1, height, width)
        expected = record["visual_event_final_frame"]
        baseline = np.median(frames[expected-55:expected-20], axis=0)
        delta = np.mean(np.abs(frames.astype(np.float32)-baseline), axis=(1,2))
        observed = next(f for f in range(expected-10,expected+10) if np.all(delta[f:f+8] > 2))
        assert observed == expected, (record["name"], expected, observed)
        final_visual_frames.append(observed)
    # Decode sound and ensure only the five intended event windows contain audio.
    wave_file = EVIDENCE / "final-audio.wav"
    run(["-i", str(audible), "-vn", "-c:a", "pcm_s16le", str(wave_file)])
    audio, rate = read_wav(wave_file)
    assert abs(len(audio)/rate - 25) < .03
    peak = np.abs(audio).max(axis=1)
    mask = np.zeros(len(peak), dtype=bool)
    for i in range(5):
        mask[int((i*5+1.8)*rate):int((i*5+2.7)*rate)] = True
    assert float(peak[~mask].max()) < .003, float(peak[~mask].max())
    files = [{"file": p.relative_to(PACKAGE).as_posix(), "bytes": p.stat().st_size, "sha256": sha(p)}
             for p in sorted(PACKAGE.rglob("*")) if p.is_file() and p.name != "05_文件核验.json"]
    report = {"package": "14_工具操作_v001", "video": {"width":1280,"height":720,"fps":60,"frame_count":1500,"seconds":25,
               "frame_index_origin":0,"same_encoded_video_stream":True,"h264_sha256":streams[0]},
              "audio": {"type":"UE mixer recording", "sample_rate":rate,"seconds":len(audio)/rate,"peak":float(peak.max()),
                        "outside_event_windows_peak":float(peak[~mask].max()), "source_files":["SFX_Tool_v01.wav","SFX_Button_v01.wav"]},
              "sound_events":sounds, "visual_result_frames":final_visual_frames, "files":files}
    (PACKAGE / "05_文件核验.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    zip_path = PACKAGE.with_suffix(".zip")
    with zipfile.ZipFile(zip_path,"w",zipfile.ZIP_DEFLATED) as archive:
        for path in sorted(PACKAGE.rglob("*")):
            if path.is_file():archive.write(path,path.relative_to(PACKAGE.parent))
    with zipfile.ZipFile(zip_path) as archive:
        assert archive.testzip() is None
        assert len(archive.namelist()) == len(files)+1
    print(json.dumps({"package":str(PACKAGE),"zip":str(zip_path),"files":len(files)+1,"zip_bytes":zip_path.stat().st_size,"frames":1500,"seconds":25},ensure_ascii=False),flush=True)


if __name__ == "__main__":
    events = visual_events(analyze())
    sounds = sound_events(events)
    make_videos(events)
    records = copy_sources()
    workbook(events, sounds, records)
    documents(events, sounds, records)
    verify_and_zip(events, sounds)
