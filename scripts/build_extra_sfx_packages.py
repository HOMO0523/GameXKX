"""Six supplemental briefs with authentic pictures and clearly labelled reference dubbing."""
import argparse,hashlib,json,shutil,urllib.request,zipfile
from pathlib import Path
import numpy as np
from openpyxl import load_workbook
from essential_sfx_requirement_catalog import SPECS,folder
from build_essential_sfx_packages import ff,wav,sheets,ass_time

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'Saved/Codex/ExtraSfxRequirements-20260915'
DEST=ROOT/'Deliverables/音效需求'
LIB=ROOT/'SourceArt/Audio/Essential'
SPECS_EXTRA=[s for s in SPECS if 16<=s['number']<=21]
RECIPES={
 'Buff':('kenney-interface','confirmation_003.ogg',1.0,.48,.46),
 'Reject':('kenney-interface','error_001.ogg',1.4,.24,.36),
 'ChestOpen':('kenney-rpg','doorOpen_1.ogg',2.0,.65,.48),
 'Critical':('kenney-impact','impactPunch_heavy_002.ogg',1.0,.44,.64),
 'PageSwitch':('kenney-rpg','bookFlip2.ogg',2.7,.19,.38),
 'LevelUp':('kenney-jingles','jingles_PIZZI02.ogg',1.0,1.10,.48),
}

def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def write(path,value):path.write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf-8')
def load(path):return json.loads(path.read_text(encoding='utf-8-sig'))

def archive(pack):
 meta=load(LIB/'Licenses'/pack/'source-manifest.json');entry=meta['downloads'][0]
 target=OUT/'downloads'/entry['file'];target.parent.mkdir(parents=True,exist_ok=True)
 if not target.exists():
  cached=ROOT/'Saved/Codex/CardDrawRequirement-20260915'/entry['file']
  if cached.is_file():shutil.copy2(cached,target)
  else:
   with urllib.request.urlopen(entry['url'],timeout=45) as response:target.write_bytes(response.read())
 assert sha(target)==entry['sha256'],str(target)
 return target,meta

def reference(spec):
 cue=spec['cue'];pack,filename,tempo,maximum,peak=RECIPES[cue];source=ROOT/'SourceArt/Audio/References'/cue
 source.mkdir(parents=True,exist_ok=True);package=DEST/folder(spec);(package/'参考音频').mkdir(parents=True,exist_ok=True)
 a,meta=archive(pack)
 with zipfile.ZipFile(a) as z:
  members=[n for n in z.namelist() if Path(n).name==filename];assert len(members)==1,members
  original=source/filename;original.write_bytes(z.read(members[0]))
 raw=ff(['-i',original,'-af',f'atempo={tempo},highpass=f=40,lowpass=f=8000','-ar','48000','-ac','1','-f','f32le','-'])
 samples=np.frombuffer(raw,'<f4').astype(np.float64);assert np.isfinite(samples).all() and np.abs(samples).max()>0
 threshold=max(.002,float(np.abs(samples).max())*.025);active=np.flatnonzero(np.abs(samples)>threshold);begin=max(0,int(active[0])-96)
 end=min(len(samples),int(active[-1])+961,begin+round(maximum*48000));part=samples[begin:end].copy();part-=part.mean()
 part[:96]*=np.linspace(0,1,96);fade=min(960,len(part)//5);part[-fade:]*=np.linspace(1,0,fade)
 gain=peak/float(np.abs(part).max());part*=gain
 name=f'SFX_{cue}_ref_v01';target=source/(name+'.wav');wav(target,part[:,None])
 record={'cue':cue,'variant':1,'name':name,'author':meta['author'],'pack':pack,'source':meta['source'],'license':'CC0-1.0','license_url':meta['license_url'],
  'file':'参考音频/'+target.name,'source_file':'来源与授权/'+pack+'/'+filename,'project_file':target.relative_to(ROOT).as_posix(),
  'project_source_file':original.relative_to(ROOT).as_posix(),'sha256':sha(target),'source_sha256':sha(original),'source_archive_sha256':sha(a),
  'seconds':round(len(part)/48000,6),'sample_rate':48000,'channels':1,'pcm_bits':16,'primary':True,'used_in_video':True,'reference_only':True,
  'role':'新增参考配音，尚未接入游戏','processing':{'tempo':tempo,'excerpt_after_tempo_seconds':begin/48000,'maximum_seconds':maximum,'highpass_hz':40,'lowpass_hz':8000,'fade_in_ms':2,'fade_out_ms':fade/48,'peak':peak,'gain':gain}}
 assert spec['duration'][0]<=record['seconds']<=spec['duration'][1],(cue,record['seconds'])
 write(source/'manifest.json',record);shutil.copy2(target,package/record['file'])
 license_dir=package/'来源与授权'/pack;license_dir.mkdir(parents=True,exist_ok=True);shutil.copy2(original,package/record['source_file'])
 for f in (LIB/'Licenses'/pack).iterdir():
  if f.is_file():shutil.copy2(f,license_dir/f.name)
 write(package/'来源与授权/音频清单.json',[record])
 return record,part

def caption(spec,plan,work):
 header='''[Script Info]
ScriptType: v4.00+
PlayResX: 1280
PlayResY: 720
WrapStyle: 2
[V4+ Styles]
Format: Name, Fontname, Fontsize, PrimaryColour, SecondaryColour, OutlineColour, BackColour, Bold, Italic, Underline, StrikeOut, ScaleX, ScaleY, Spacing, Angle, BorderStyle, Outline, Shadow, Alignment, MarginL, MarginR, MarginV, Encoding
Style: Title,Microsoft YaHei,28,&H00F1F3EC,&H00000000,&H00000000,&H00000000,-1,0,0,0,100,100,0,0,1,0,0,7,0,0,0,1
Style: Body,Microsoft YaHei,18,&H00D0D8D1,&H00000000,&H00000000,&H00000000,0,0,0,0,100,100,0,0,1,0,0,7,0,0,0,1
Style: Cue,Microsoft YaHei,20,&H008FD8B3,&H00000000,&H00000000,&H00000000,0,0,0,0,100,100,0,0,1,0,0,7,0,0,0,1
[Events]
Format: Layer, Start, End, Style, Name, MarginL, MarginR, MarginV, Effect, Text
'''
 lines=[]
 def line(start,end,style,x,y,text):lines.append(f'Dialogue: 0,{ass_time(start)},{ass_time(end)},{style},,0,0,0,,{{\\pos({x},{y})}}{text}')
 line(0,plan['seconds'],'Title',28,19,f"{spec['number']:02} / {spec['title']} · 音效需求")
 line(0,plan['seconds'],'Body',870,26,'参考配音 · 游戏待接入')
 for e in plan['events']:
  section=plan['sections'][e['section']]
  if spec['cue']=='ChestOpen':
   for k,label in enumerate(['开箱成功','宝箱数量 1 → 0','破军头冠进入背包','对齐物品首次揭示']):line(0,plan['seconds'],'Body',60,210+k*38,label)
  else:line(section['offset'],section['offset']+section['duration'],'Body',28,651,e['label'])
  line(e['seconds'],min(plan['seconds'],e['seconds']+.85),'Cue',720,688,f"触发点 {e['seconds']:.3f} s / F{e['target_frame']:04}")
 if spec['cue']=='Reject':
  # A brief annotation marks an otherwise visually silent rejected attempt.
  line(plan['events'][0]['seconds'],plan['events'][0]['seconds']+1.3,'Cue',28,48,'需求标注：此刻尝试出牌，被拒绝')
 (work/'caption.ass').write_text(header+'\n'.join(lines)+'\n',encoding='utf-8-sig')

def build(spec,edit):
 record,part=reference(spec);package=DEST/folder(spec);work=OUT/('build-'+spec['cue']);work.mkdir(exist_ok=True)
 sections=[];events=[];clips=[];offset=0
 for i,segment in enumerate(edit['segments']):
  source=ROOT/segment['source'];start=segment['source_frame']-60;length=segment.get('length',240);assert start>=0
  filters=[f'trim=start_frame={start}:end_frame={start+length}','setpts=PTS-STARTPTS']
  if edit.get('crop'):
   x,y,w,h=edit['crop'];filters += [f'crop={w}:{h}:{x}:{y}','scale=1280:570:force_original_aspect_ratio=decrease','pad=1280:570:(ow-iw)/2:(oh-ih)/2:color=0x11181c','pad=1280:720:0:72:color=0x11181c']
  target=work/f'clip-{i}.mp4';ff(['-i',source,'-vf',','.join(filters),'-an','-c:v','libx264','-crf','18','-preset','fast','-threads','4','-pix_fmt','yuv420p','-r','60','-fps_mode','cfr',target]);clips.append(target)
  sections.append({'name':source.stem,'video':str(source),'source_sha256':sha(source),'offset':offset/60,'duration':length/60,'trim_frame':start,'raw_event_frame':segment['source_frame'],'crop':edit.get('crop'),'capture_api':'Windows.Graphics.Capture'})
  events.append({'cue':spec['cue'],'file':record['name']+'.wav','seconds':(offset+60)/60,'frame_contains_start':offset+60,'target_frame':offset+60,'primary':True,'reference_only':True,'section':i,'label':segment['label'],'timing_basis':edit['timing_basis']})
  offset+=length
 plan={'cue':spec['cue'],'title':spec['title'],'seconds':offset/60,'frame_count':offset,'reference_only':True,'sections':sections,'events':events,'recording_notes':edit['notes']}
 (work/'concat.txt').write_text('\n'.join(f"file '{p.name}'" for p in clips),encoding='utf-8')
 ff(['-f','concat','-safe','0','-i',work/'concat.txt','-c','copy',work/'picture.mp4'])
 mixed=np.zeros((offset*800,2),np.float64)
 for e in events:
  at=round(e['seconds']*48000);mixed[at:at+len(part)]+=part[:,None]*.282
 wav(work/'reference-mix.wav',mixed);caption(spec,plan,work);shutil.copy2('C:/Windows/Fonts/consola.ttf',work/'counter.ttf')
 vf=r"subtitles=caption.ass,drawtext=fontfile=counter.ttf:fontsize=18:fontcolor=0xaebfb5:x=28:y=691:text='F%{eif\:n\:d\:4}  /  60 fps'"
 ff(['-i','picture.mp4','-i','reference-mix.wav','-map','0:v:0','-map','1:a:0','-vf',vf,'-c:v','libx264','-crf','18','-preset','fast','-threads','4','-pix_fmt','yuv420p','-c:a','aac','-b:a','192k','-t',plan['seconds'],'-movflags','+faststart',package/'01_演示_参考配音.mp4'],cwd=work)
 ff(['-i',package/'01_演示_参考配音.mp4','-map','0:v:0','-an','-c:v','copy','-movflags','+faststart',package/'02_演示_无声版.mp4'])
 sheets(spec,plan,[record],package)
 book=load_workbook(package/'03_时间轴需求.xlsx');timeline=book['时间轴需求'];timeline.cell(1,3).value='参考配音落点秒';timeline.cell(1,4).value='成片帧索引'
 for i,e in enumerate(events,2):timeline.cell(i,2).value=e['label'];timeline.cell(i,7).value='新增参考配音，待接入'
 for row in book['制作需求'].iter_rows():
  if row[0].value=='时间含义':row[1].value=edit['timing_basis']+'；60fps成片帧从F0000起，不等同于引擎内部帧。'
  if row[0].value=='状态':row[1].value='新增独立音效需求，参考声音未导入游戏；视频为真实画面加参考配音。'
 book['制作需求'].append(['录制范围',edit['notes']]);book.save(package/'03_时间轴需求.xlsx')
 timeline_text='\n'.join(f"{e['seconds']:.3f} 秒 / F{e['target_frame']:04} — {e['label']} — {e['file']}" for e in events)
 (package/'04_时间轴速览.txt').write_text(f"{spec['title']} v001 · 新增参考配音 / 待接入\n{plan['seconds']}秒 / 60fps / {offset}帧\n\n{timeline_text}\n\n触发：{spec['rule']}\n不触发：{spec['silent']}\n\n取证：{edit['timing_basis']}\n{edit['notes']}\n",encoding='utf-8-sig')
 (package/'00_先看这里.md').write_text(f"# {spec['title']} · 音效制作需求\n\n双击 `00_音效需求.html`，可离线查看视频、按秒/帧定位、试听和下载参考音，也可导入制作好的 WAV 试配并导出单文件评审 HTML。\n\n**新增需求；参考配音，尚未接入游戏。** 视频保留真实游戏画面，只配本类参考声。字幕属于需求标注。\n\n{spec['brief']} 建议 {spec['duration'][0]}–{spec['duration'][1]} 秒。\n\n- 触发：{spec['rule']}\n- 保持安静：{spec['silent']}\n\n{timeline_text}\n\n录制与落点：{edit['timing_basis']}。{edit['notes']}\n\n原始素材 `{filename_for(spec)}` 来自 [{record['author']} 发布页]({record['source']})，CC0-1.0。原文件、授权记录、加工参数和哈希随包提供。参考 WAV 为 48kHz、单声道、16-bit，来自有损 OGG，加工后的声音用于方向和时长参考。正式制作建议交付 48kHz、单声道、24-bit PCM；最终听感待人工确认。音频授权不覆盖游戏画面或随后导入的新成品。\n",encoding='utf-8')
 streams=[ff(['-i',package/n,'-map','0:v:0','-c','copy','-bsf:v','h264_mp4toannexb','-f','h264','-']) for n in ['01_演示_参考配音.mp4','02_演示_无声版.mp4']]
 assert hashlib.sha256(streams[0]).digest()==hashlib.sha256(streams[1]).digest()
 frames=[l for l in ff(['-i',package/'01_演示_参考配音.mp4','-map','0:v:0','-f','framemd5','-']).decode().splitlines() if l and not l.startswith('#')];assert len(frames)==offset
 encoded=np.frombuffer(ff(['-i',package/'01_演示_参考配音.mp4','-vn','-ar','48000','-ac','2','-f','f32le','-']),'<f4').reshape(-1,2)
 n=min(len(encoded),len(mixed));correlation=float(np.corrcoef(encoded[:n].ravel(),mixed[:n].ravel())[0,1]);assert correlation>.99,(spec['cue'],correlation)
 manifest={'package':folder(spec),'primaryCue':spec['cue'],'reference_only':True,'integration_status':'not_integrated','video':{'width':1280,'height':720,'fps':60,'frame_count':offset,'seconds':plan['seconds'],'same_encoded_video_stream':True,'soundtrack':'reference_dub','aac_mix_correlation':correlation},'sound_events':events,'recording_notes':edit['notes'],'files':[{'file':p.relative_to(package).as_posix(),'bytes':p.stat().st_size,'sha256':sha(p)} for p in sorted(package.rglob('*')) if p.is_file() and p.name!='05_文件核验.json']}
 write(package/'05_文件核验.json',manifest);write(work/'delivery-plan.json',plan);write(ROOT/'Saved/Codex/EssentialSfxRequirements-20260914'/(spec['cue']+'-delivery-plan.json'),plan)
 print(json.dumps({'cue':spec['cue'],'seconds':plan['seconds'],'frames':offset,'reference_seconds':record['seconds'],'aac_correlation':correlation},ensure_ascii=False),flush=True)

def filename_for(spec):return RECIPES[spec['cue']][1]

def main():
 p=argparse.ArgumentParser();p.add_argument('--references-only',action='store_true');p.add_argument('--cue',default='');a=p.parse_args()
 edits={} if a.references_only else load(OUT/'edit-list.json')
 for spec in SPECS_EXTRA:
  if a.cue and a.cue!=spec['cue']:continue
  if a.references_only:
   r,_=reference(spec);print(spec['cue'],r['seconds'],r['source_file'],flush=True)
  else:build(spec,edits[spec['cue']])

if __name__=='__main__':main()
