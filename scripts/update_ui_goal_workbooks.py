"""Export the current goal's real project state; never imply pending gameplay is complete."""
import json
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
STATE = ROOT / "docs/production/2026-09-10-ui-guidance-localization-state.json"
OUT = ROOT / "docs/design/2026-09-10-project-progress"
FILE = OUT / "GameXXK_项目进度与验收总表_2026-09-10.xlsx"


def prepare(ws, title, subtitle, headers, widths):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.cell(1, 1, title).font = Font(name="Microsoft YaHei", size=20, bold=True, color="183F3A")
    ws.row_dimensions[1].height = 38
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    ws.cell(2, 1, subtitle).font = Font(name="Microsoft YaHei", size=10, color="62766F")
    ws.row_dimensions[2].height = 30
    for index, (header, width) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(4, index, header)
        cell.font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244C45")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.row_dimensions[4].height = 28
    ws.freeze_panes = "C5"


def finish(ws):
    for row in ws.iter_rows(min_row=5):
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=10, color="243E38")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.fill = PatternFill("solid", fgColor="F3F5EF" if cell.row % 2 else "FFFFFF")
        ws.row_dimensions[row[0].row].height = 60
    ws.auto_filter.ref = f"A4:{get_column_letter(ws.max_column)}{ws.max_row}"
    ws.print_title_rows = "1:4"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def main():
    state = json.loads(STATE.read_text(encoding="utf-8"))
    manifest_path = ROOT / "SourceArt/UI/Relics/gemstyle-20260910/manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    jobs = json.loads((manifest_path.parent / "art-jobs.json").read_text(encoding="utf-8"))["jobs"]
    icons = {entry["slug"]: entry for entry in manifest["icons"]}
    imported = sum(bool(e.get("imported")) for e in icons.values())
    design = json.loads((ROOT / "docs/design/2026-09-10-relic-redesign/relic-design.json").read_text(encoding="utf-8"))
    relics = {row["slug"]: row for row in design["relics"]}
    timestamp = datetime.now(timezone(timedelta(hours=8))).strftime("%Y-%m-%d %H:%M +08:00")
    wb = Workbook()
    ws = wb.active
    ws.title = "01_项目进度"
    prepare(ws, "GameXXK · 界面、本地化与讨伐版本进度", f"{timestamp}  |  {state['overallStatus']}  |  {state['branchDecision']}",
            ["ID", "工作包", "设计状态", "实现状态", "验证状态", "当前进度", "下一步", "证据 / 源文件"],
            [9, 29, 32, 37, 43, 44, 48, 67])
    for item in state["milestones"]:
        progress = f"{len(icons)}/45 已整理透明图；{imported}/45 已导入" if item["id"] == "G5A" else item["progress"]
        ws.append([item["id"], item["name"], item["design"], item["implementation"], item["validation"], progress, item["next"], item["evidence"]])
    finish(ws)

    checks = wb.create_sheet("02_完成门槛")
    prepare(checks, "完成条件 · 勾选必须有对应证据", "来自当前 goal 文档；未勾选即未完成，不能由图片存在或编译成功推断玩法验收通过。",
            ["编号", "状态", "验收条件", "说明"], [10, 16, 135, 47])
    goal = (ROOT / state["goalDocument"]).read_text(encoding="utf-8")
    for index, (mark, description) in enumerate(re.findall(r"^- \[([ xX])\] (.+)$", goal, re.M), 1):
        checks.append([f"A{index:02}", "已完成" if mark.strip() else "未完成", description, "最终通过后更新勾选及证据"])
    finish(checks)

    art = wb.create_sheet("03_遗物素材状态")
    prepare(art, "45件遗物 · 逐件素材状态", "旧30件与新增15件去重；普通/特殊机制保持原决定。首批写实稿已否决，表中只记录宝石同系列画风。",
            ["Slug", "名称", "来源", "机制品质", "素材制作", "视觉复核", "UE导入", "图标源路径", "SHA256"],
            [24, 20, 18, 18, 22, 31, 17, 72, 68])
    for job in jobs:
        icon = icons.get(job["slug"])
        spec = relics.get(job["slug"])
        quality = {"Rare": "稀有", "Epic": "史诗"}.get(spec["quality"], "") if spec else "普通（保留）"
        art.append([job["slug"], job["name"], "旧图重绘" if job["existing"] else "新增物件", quality,
                    "512透明图已整理" if icon else "待完成", icon.get("visualReview", "待复核") if icon else "待复核",
                    "已导入" if icon and icon.get("imported") else "未导入", icon["icon"] if icon else "", icon["sha256"] if icon else ""])
    finish(art)

    deps = wb.create_sheet("04_共享任务交接")
    prepare(deps, "同项目任务 · 基线与交接", "任务标题按原名称记录；构建和编辑器窗口串行交接，用户要求最后统一提交。",
            ["任务名称", "当前状态", "实际范围", "验证情况", "验收依据"], [30, 39, 30, 70, 85])
    for task in state["otherTasks"]:
        deps.append([task["name"], task["state"], task["scope"], task["verification"], task["evidence"]])
    finish(deps)

    OUT.mkdir(parents=True, exist_ok=True)
    temporary = FILE.with_name(FILE.stem + ".pending.xlsx")
    wb.save(temporary)
    reread = load_workbook(temporary, read_only=True, data_only=False)
    actual = {sheet.title: sheet.max_row - 4 for sheet in reread}
    assert actual["01_项目进度"] == len(state["milestones"])
    assert actual["03_遗物素材状态"] == 45
    assert actual["04_共享任务交接"] == len(state["otherTasks"])
    reread.close()
    temporary.replace(FILE)
    report = {"updatedAt": timestamp, "path": str(FILE), "sheets": actual, "relicArt": len(icons), "relicImported": imported,
              "note": "Workbook reload/count validation only; pending gameplay remains pending."}
    (OUT / "workbook-update-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    main()
