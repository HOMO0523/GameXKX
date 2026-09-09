"""Apply verified 1.05 gold/talent calibration without touching other table data."""
import hashlib
import json
import math
import re
import shutil
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
TABLES = ROOT/'docs/design/2026-09-04-project-design-tables'
EVIDENCE = ROOT/'docs/design/2026-09-10-training-economy/runtime-budget.json'
OUT = ROOT/'Saved/Diagnostics/TrainingEconomy105'


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def budget():
    data = json.loads(EVIDENCE.read_text(encoding='utf-8-sig'))
    training = (ROOT/'Source/GameXXK/Private/GameXXKTrainingRules.cpp').read_text(encoding='utf-8')
    talents = (ROOT/'Source/GameXXK/Private/GameXXKTalentRules.cpp').read_text(encoding='utf-8')
    growth = float(re.search(r'PostNormalGoldGrowthPerStage\s*=\s*([\d.]+)', training).group(1))
    base = float(re.search(r'const double Raw\s*=\s*([\d.]+)\s*\*\s*FMath::Pow', talents).group(1))
    assert growth == data['gold_growth'] == 1.05 and base == data['price_base'] == 229
    assert sha(ROOT/'Source/GameXXK/Private/GameXXKTalentCatalog.cpp') == data['catalog_sha256'], 'Talent catalog changed; rerun calibration validation'
    assert data['nodes'] == 453 and data['ranks'] == 2229 and data['total_gold'] == 1880754900
    assert data['waves'] == 388128 and abs(data['days']-45) < .25
    assert len(data['stages']) == 27 and len(data['prices']) == 36
    return data


def stage_rows(data):
    rows = []
    labels = {'Normal':'普通', 'Hard':'困难', 'Hell':'地狱'}
    for entry in data['stages']:
        match = re.fullmatch(r'Training\.(Normal|Hard|Hell)\.(\d+)-(\d+)', entry['id'])
        assert match, entry['id']
        diff, chapter, part = match.groups()
        number = (int(chapter)-1)*3+int(part)
        level = (list(labels).index(diff)*9+number)*5
        gold, xp = int(entry['gold']), int(entry['experience'])
        full = math.floor(gold*4.5+.5)
        rows.append([labels[diff], chapter+'-'+part, level, gold, xp, full,
                     '已校准：普通×10；困难至地狱连续每关×1.05；天赋价格与45天基准见07E/07F',
                     'Source/GameXXK/Private/GameXXKTrainingRules.cpp；docs/design/2026-09-10-training-gold-talent-calibration.md'])
    return rows


def write_sheet(wb, name, headers, rows, widths):
    position = wb.sheetnames.index(name) if name in wb.sheetnames else len(wb.sheetnames)
    if name in wb.sheetnames: wb.remove(wb[name])
    ws = wb.create_sheet(name, position)
    ws.append(headers)
    for row in rows: ws.append(row)
    for row in ws:
        ws.row_dimensions[row[0].row].height = 38 if row[0].row > 1 else 30
        for cell in row:
            cell.font = Font(name='Microsoft YaHei', size=11, color='263F43')
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            cell.fill = PatternFill('solid', fgColor='F1F5F4' if cell.row%2 == 0 else 'FFFFFF')
            if isinstance(cell.value, int): cell.number_format = '#,##0'
    for cell in ws[1]:
        cell.font = Font(name='Microsoft YaHei', size=11, color='FFFFFF', bold=True)
        cell.fill = PatternFill('solid', fgColor='244D57')
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
    ws.freeze_panes='C2'; ws.auto_filter.ref=ws.dimensions
    return ws


def apply_training_economy_update(wb):
    data = budget()
    write_sheet(wb, '07A_挂机金币现状',
                ['难度','关卡','敌人等级','每波基础金币','每波基础经验','满在线金币天赋每波金币','当前状态','源码定位'],
                stage_rows(data), [12,12,15,21,21,31,70,68])
    rows = [
        ['目标周期', '约45天', '24小时在线；稳定刷地狱3-3；平均10秒/波'],
        ['起点和购买顺序', '零金币、零天赋', '先购买可解锁且最便宜的在线金币天赋，再买其他可解锁天赋'],
        ['不计入的因素', '推关/额外收入/其他消费', '不计前期推关时间、宝箱与分解等额外收入，也不计买箱和工具支出'],
        ['当前节点数', int(data['nodes']), '保持原目录、解锁前置与容量门槛'],
        ['升级总次数', int(data['ranks']), '每个节点全部等级'],
        ['旧全树总价', 20532352700, '旧2500价格基数，作为历史对比'],
        ['当前全树总价', int(data['total_gold']), '以实际C++目录和购买接口模拟核对'],
        ['价格公式', '百位四舍五入(229×1.35^档位)', '同节点每级同价；每档价格见07F'],
        ['根/入口每级金币', 200, '0档，根节点与四个分支入口各1级'],
        ['最高档每级金币', 8346700, '第35档'],
        ['金币天赋点满天数', round(data['gold_talent_max_days'],6), '过程中收益由1倍逐渐提高至4.5倍'],
        ['全树点满波数', int(data['waves']), '逐波取整发金币，保留余额，逐点合法购买'],
        ['全树点满天数', round(data['days'],6), '固定对比基准；不是所有玩家从新档起的保证时间'],
        ['地狱末关基础金币/波', 1083, '普通450为起点，连续增长18次'],
        ['末关满金币天赋/波', 4874, '1083×4.5，最后四舍五入'],
        ['末关7波满天赋金币', 34118, '约三局累计可购1个100000金币高级箱'],
        ['主动挑战', '基础金币×2', '不叠在线金币天赋；经验曲线未修改'],
        ['已购买天赋', '等级保留、不自动退款', '本次只改变之后购买的价格'],
        ['验证记录', 'TrainingEconomy105', 'Saved/Diagnostics/TrainingEconomy105；文档见2026-09-10-training-gold-talent-calibration.md'],
    ]
    write_sheet(wb, '07E_挂机与天赋预算', ['项目','金额/数值','口径'], rows, [28,35,100])
    write_sheet(wb, '07F_天赋分档价格', ['档位','每级金币','5级节点合计参考','说明'],
                [[int(r['tier']),int(r['price']),int(r['price'])*5,
                  '根/入口为1级；仓库页节点也是1级，实际按各节点上限计费'] for r in data['prices']], [12,22,25,80])


def media_hashes(path):
    with zipfile.ZipFile(path) as archive:
        return sorted(hashlib.sha256(archive.read(name)).hexdigest() for name in archive.namelist() if name.startswith('xl/media/'))


def update_book(path, affected, apply):
    original = sha(path)
    images = media_hashes(path)
    wb = openpyxl.load_workbook(path)
    unchanged = {ws.title:list(ws.values) for ws in wb if ws.title not in affected}
    apply(wb)
    OUT.mkdir(parents=True, exist_ok=True)
    backup = OUT/(path.stem+'-before-105.xlsx')
    if not backup.exists(): shutil.copy2(path, backup)
    temp = path.with_name(path.stem+'.economy105.tmp.xlsx')
    wb.save(temp); wb.close()
    check = openpyxl.load_workbook(temp)
    assert all(list(check[name].values)==rows for name,rows in unchanged.items()), 'An unrelated sheet changed'
    check.close()
    assert media_hashes(temp)==images, 'Existing embedded art changed'
    assert sha(path)==original, 'Workbook changed concurrently'
    temp.replace(path)
    return {'file':str(path),'updated_sheets':sorted(affected),'unchanged_sheets':len(unchanged),'sha256':sha(path),'images_preserved':True}


def main():
    data = budget()
    equipment = TABLES/'GameXXK_装备设计总表_2026-09-04.xlsx'
    enemy = TABLES/'GameXXK_怪物与阶段数值设计总表_2026-09-04.xlsx'
    reports = [update_book(equipment, {'07A_挂机金币现状','07E_挂机与天赋预算','07F_天赋分档价格'}, apply_training_economy_update)]
    reports.append(update_book(enemy, {'14_挂机金币曲线'}, lambda wb: write_sheet(wb, '14_挂机金币曲线',
        ['难度','关卡','敌人等级','每波基础金币','每波基础经验','满在线金币天赋每波金币','当前状态','源码定位'],
        stage_rows(data), [12,12,15,21,21,31,70,68])))
    (OUT/'workbook-report.json').write_text(json.dumps(reports,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(reports,ensure_ascii=False))


if __name__ == '__main__': main()
