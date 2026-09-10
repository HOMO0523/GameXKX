"""Sync the approved Wind change without regenerating unrelated tuning sheets."""
from pathlib import Path
import copy
import json
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from update_ui_goal_workbooks import prepare, finish

ROOT=Path(__file__).resolve().parents[1]
FILE=ROOT/'docs/design/2026-09-04-project-design-tables/GameXXK_装备设计总表_2026-09-04.xlsx'
DOC='docs/design/2026-09-11-affix-wind-ramp.md'
RETIRED={'Affix.ZhuiFeng.Draw','Affix.ZhuiFeng.SharedEnergy','Affix.ZhuiFeng.ComboCount','Affix.ZhuiFeng.TemporaryCostReduction','Affix.ShanHe.TerrainCostReduction'}
NEW='Affix.ZhuiFeng.LateCardDamage'
COPY=next(e for e in json.loads((ROOT/'Content/Localization/GameXXK/strings.json').read_text(encoding='utf-8'))['entries'] if e.get('namespace')=='EquipmentAffixUI' and e['key']=='WindMomentum')
ZH=COPY['zh-Hans']
EN=COPY['en']

def main():
    wb=load_workbook(FILE)
    allowed={'04_词缀目录','04A_触发与接入审查','04B_品质与数值基线','04D_叠加上限与边际','08_实现差异','04H_追风递增与退休'}
    preserved={(s.title,c.coordinate):c.value for s in wb if s.title not in allowed for row in s for c in row}
    ws=wb['04_词缀目录']
    for row in range(2,ws.max_row+1):
        id=ws.cell(row,9).value
        if id in RETIRED:
            ws.cell(row,4,'已退休；旧战斗快照兼容至结算')
            ws.cell(row,5,'不再生成或洗炼；旧装备载入时转换')
            ws.cell(row,6,'原资源操作不接入；改为未占用数值词缀')
            ws.cell(row,7,'定向迁移；省力优先转乘风；强化/孔/其他词条保留。'+DOC)
    row=next((r for r in range(2,ws.max_row+1) if ws.cell(r,9).value==NEW),ws.max_row+1)
    values=[36,'追风',ZH,'已接伤害路径，冷编译/行为验收中','追风随机池；固定1%','穿戴者当前牌物理/元素直击；全队主动计数；每回合重置','1%=100基点；同名按条数加算，不套用0.5系数或75%递减','BasisPoints',NEW,'Source/GameXXK/Private/GameXXKCardRules.cpp | '+DOC,'Tailwind',EN]
    for col,value in enumerate(values,1):
        c=ws.cell(row,col,value)
        if row>2:c._style=copy.copy(ws.cell(row-1,min(col,10))._style)
        c.alignment=Alignment(wrap_text=True,vertical='top')
    ws.cell(1,11,'英文短名');ws.cell(1,12,'英文效果');ws.column_dimensions['K'].width=18;ws.column_dimensions['L'].width=70
    ws.row_dimensions[row].height=90;ws.auto_filter.ref=f'A1:L{ws.max_row}'
    ws=wb['04A_触发与接入审查']
    for r in range(2,ws.max_row+1):
        if ws.cell(r,9).value in RETIRED:
            ws.cell(r,4,'退休，不再接入旧资源效果');ws.cell(r,7,'旧档转换为未占用数值项，保留品质/数值相对位置');ws.cell(r,8,'用户已确认；转换与运行时测试中')
    r=next((r for r in range(2,ws.max_row+1) if ws.cell(r,9).value==NEW),ws.max_row+1)
    for c,v in enumerate([36,'追风',ZH,'直接伤害路径已接入，测试中','全队第3张起；穿戴者自己的当前牌','不增重放计数，不影响固定伤害/DOT/装备触发','固定1%；两条第3/4张分别+2%/+4%','用户确认；冷编译中',NEW],1):
        ws.cell(r,c,v);ws.cell(r,c).alignment=Alignment(wrap_text=True,vertical='top')
    ws.row_dimensions[r].height=90
    for r in range(2,wb['04B_品质与数值基线'].max_row+1):
        wb['04B_品质与数值基线'].cell(r,11,'产甲已接入递减；追风乘风各品质固定1%并线性叠加（不适用本页区间/折半）；其余分系项待接入，退休5项见04H。')
    ws=wb['04D_叠加上限与边际']
    r=next((r for r in range(2,ws.max_row+1) if ws.cell(r,1).value=='追风乘风例外'),ws.max_row+1)
    for c,v in enumerate(['追风乘风例外','max(0,主动牌序号-2) × 同名条数 × 1%','固定1%直接叠加，不折半、不使用75%递减；详见04H','已接代码；测试中'],1):ws.cell(r,c,v)
    ws=wb['08_实现差异']
    for r in range(2,ws.max_row+1):
        if ws.cell(r,1).value=='分系随机词缀接入':
            ws.cell(r,2,'原30项退休5项、新增乘风1项，当前26项；产甲已通过，乘风已接待测。')
            ws.cell(r,3,'24项仍需触发/运行时接入；不等同固定套装失效。');ws.cell(r,4,'产甲已验证；乘风验证中；24待接入')
    name='04H_追风递增与退休'
    if name in wb:del wb[name]
    ws=wb.create_sheet(name,wb.sheetnames.index('04G_详细属性口径')+1)
    prepare(ws,'追风乘风 · 固定1%递增','用户裁决 2026-09-11；实现/测试进度见项目状态表。',['项目','中文规则','English','状态'],[24,90,100,28])
    rows=[
        ['效果',ZH,EN,'代码已接，待测试'],
        ['1条示例','第1/2张0%，第3/4/5张+1%/+2%/+3%。','Cards 1/2: 0%; cards 3/4/5: +1%/+2%/+3%.','已确认'],
        ['2条示例','第3/4/5张+2%/+4%/+6%。','Cards 3/4/5: +2%/+4%/+6%.','已确认'],
        ['品质','普通至无量，每条固定100基点；品质不增幅。','Every tier grants exactly 1% per step per copy.','已接目录'],
        ['范围','物理/元素/多段/重箭直击；只看出牌者配装。','Own physical, elemental, multi-hit and Heavy Arrow card attacks.','待实际伤害对比'],
        ['排除','固定伤害、自损、持续伤害、反应、装备触发不加。','Excludes fixed damage, self-loss, DOT, reactions and gear procs.','待回归'],
        ['重放','不增加计数，使用当前全队计数；下回合归零。','Replays use the current team count without advancing it; reset each round.','待实际重放对比'],
        ['退休池','追风抽牌/回气/计数/减费；山河地势减费。','Retired random draw, AP, combo-count and cost-discount affixes.','生成已移除'],
        ['旧档','定向转换退休项，优先省力→乘风，保留其余装备字段。','Convert only retired rolls; preserve the rest of each item.','迁移待验收'],
        ['待选洗炼','原/新结果同步转换，不重复扣费；始终可保留或采用。','Migrate both paid choices without spending sand again.','待迁移回归'],
        ['旧战斗','保存的战斗效果快照保留至结算，下场启用新配装。','Saved battles retain their snapshots; new loadouts apply next battle.','兼容规则'],
        ['存档范围','未加持久化字段；沿用原装备结构及存档校验。','No new serialized fields; existing gear schema is retained.','待读回'],
    ]
    for row in rows:ws.append(row)
    finish(ws)
    for (name,cell),value in preserved.items():assert wb[name][cell].value==value
    temp=FILE.with_suffix('.pending.xlsx');wb.save(temp)
    check=load_workbook(temp,read_only=False)
    for (name,cell),value in preserved.items():assert check[name][cell].value==value
    check.close();temp.replace(FILE)
    print(json.dumps({'workbook':str(FILE),'preservedCells':len(preserved),'affixIdentities':36,'activeSetAffixes':26,'retiredResourceAffixes':5},ensure_ascii=True))

if __name__=='__main__':main()
