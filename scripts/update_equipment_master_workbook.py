"""Refresh the equipment master, including the approved 2026-09-09 gem-quality revision."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import shutil
import subprocess
import tempfile
import unicodedata
from datetime import datetime, timedelta, timezone
from pathlib import Path

import openpyxl
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from update_travel_money_master import apply_travel_money_update
from update_training_economy_master import apply_training_economy_update

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "docs/design/2026-09-04-project-design-tables/GameXXK_装备设计总表_2026-09-04.xlsx"
PRIVATE = "Source/GameXXK/Private/"
PUBLIC = "Source/GameXXK/Public/"
FILES = {
    "affix": PRIVATE + "GameXXKAffixCatalog.cpp",
    "tooltip": PRIVATE + "UI/GameXXKEquipmentTooltipPresentation.cpp",
    "equipment": PRIVATE + "GameXXKEquipmentRules.cpp",
    "catalog": PRIVATE + "GameXXKEquipmentCatalog.cpp",
    "types": PRIVATE + "GameXXKEquipmentTypes.cpp",
    "sets": PRIVATE + "GameXXKEquipmentSetCatalog.cpp",
    "battle": PRIVATE + "GameXXKCardRules.cpp",
    "adapter": PRIVATE + "GameXXKCardBattleAdapter.cpp",
    "gems": PRIVATE + "GameXXKGemRules.cpp",
    "gem_limits": PUBLIC + "GameXXKGemRules.h",
    "gem_combat": PRIVATE + "GameXXKCombatGemRules.cpp",
    "resistance": PRIVATE + "GameXXKResistanceRules.cpp",
    "resistance_profiles": PRIVATE + "Data/GameXXKResistanceProfiles.inl",
    "training": PRIVATE + "GameXXKTrainingRules.cpp",
    "chests": PRIVATE + "GameXXKTrainingChestRules.cpp",
    "subsystem": PRIVATE + "MVP/GameXXKMVPSubsystem.cpp",
    "economy": PRIVATE + "GameXXKEquipmentEconomyRules.cpp",
    "tools": PRIVATE + "GameXXKEquipmentToolRules.cpp",
    "shop": PRIVATE + "GameXXKMetaShopRules.cpp",
    "workbench": PRIVATE + "UI/GameXXKDesktopTrainingWorkbenchWidget.cpp",
    "old_shop": PRIVATE + "UI/GameXXKMetaShopWidget.cpp",
    "old_roster": PRIVATE + "UI/GameXXKCompanionRosterWidget.cpp",
    "save": PUBLIC + "MVP/GameXXKSaveMigration.h",
    "limits": PUBLIC + "GameXXKEquipmentRules.h",
    "tool_limits": PUBLIC + "GameXXKEquipmentToolRules.h",
    "shop_limits": PUBLIC + "GameXXKMetaShopRules.h",
    "talents": PRIVATE + "GameXXKTalentRules.cpp",
    "bonus_rules": PRIVATE + "GameXXKEquipmentBonusRules.cpp",
    "bonus_limits": PUBLIC + "GameXXKEquipmentBonusRules.h",
    "detailed_attributes": PRIVATE + "UI/GameXXKCharacterDetailedAttributes.cpp",
    "inventory_ui": PRIVATE + "UI/GameXXKInventoryWindowWidget.cpp",
}
SETS = {"PoJun": "破军", "XuanJia": "玄甲", "QingNang": "青囊", "ZhuiFeng": "追风", "ShiGu": "蚀骨", "ShanHe": "山河", "Starter": "基础"}
QUALITY = ["普通", "稀有", "珍稀", "传奇", "不朽", "至宝", "超凡", "天界", "登神", "宇宙"]
AFFIX_ALPHA = 0.5
AFFIX_BONUS_CAP = 75.0
GEM_QUALITY_BONUSES = [1, 2, 4, 8, 16, 32, 58, 95, 137, 172]
GEM_QUALITY_MULTIPLIERS = [None, 2, 2, 2, 2, 2, 1.8125, 1.625, 1.4375, 1.25]
GEM_HEALTH_MULTIPLIER = 5
QUALITY_TOKENS = ["Common","Rare","Epic","Legendary","Immortal","Treasure","Transcendent","Celestial","Ascendant","Cosmic"]
GEM_TYPES = [
    ("Attack","攻击",1,"平坦值"), ("Defense","防御",1,"平坦值"), ("MaxHealth","生命",5,"平坦值"),
    ("AttackPercent","攻击百分比",0.25,"名义百分比"), ("DefensePercent","防御百分比",0.25,"名义百分比"), ("MaxHealthPercent","生命百分比",0.25,"名义百分比"),
    ("DirectDamage","物理伤害",0.5,"名义百分比"), ("ArmorGain","护甲获得量",0.5,"名义百分比"), ("Healing","治疗效果",0.5,"名义百分比"),
    ("CounterDamage","反击伤害",0.5,"名义百分比"), ("FireDamage","火焰伤害",0.5,"名义百分比"), ("DamageOverTime","持续伤害",0.5,"名义百分比"),
    ("FrostDamage","冰霜伤害",0.5,"名义百分比"), ("LightningDamage","雷击伤害",0.5,"名义百分比"),
    ("FireResistance","火焰抗性",0.5,"名义抗性百分点"), ("FrostResistance","冰霜抗性",0.5,"名义抗性百分点"), ("LightningResistance","雷击抗性",0.5,"名义抗性百分点"),
]


def current_gem_value(src, token, rank):
    registered = set(re.findall(r'case EGameXXKGemType::(\w+):\s*return TEXT',src["gems"]))
    if token not in registered:
        return None
    bonuses, health = current_gem_values(src)
    if token in {"Attack","Defense","MaxHealth"}:
        return bonuses[rank-1] * (health if token == "MaxHealth" else 1)
    constant = "StatPercentBasisPointsPerStep" if token.endswith("Percent") else "MechanicBasisPointsPerStep"
    found = re.search(constant+r"\s*=\s*(\d+)", src["gem_limits"])
    return bonuses[rank-1] * int(found.group(1))/100 if found else None


def current_gem_values(src):
    values = re.search(r"QualityBonuses\[\]\s*=\s*\{([^}]+)\}", src["gems"])
    if values:
        bonuses = [int(v.strip()) for v in values.group(1).split(",") if v.strip()]
        health = int(re.search(r"MaxHealth \? Bonus \* (\d+)", src["gems"]).group(1))
        assert len(bonuses) == 10
        return bonuses, health
    assert "1 << (Rank - 1)" in src["gems"] and "Multiplier * 10" in src["gems"]
    return [1 << n for n in range(10)], 10


def diminishing_bonus(original_percent, cap=AFFIX_BONUS_CAP):
    """Design-model extra bonus in percentage points; not a runtime gameplay implementation."""
    if original_percent < 0 or cap <= 0:
        raise ValueError("The original bonus must be nonnegative and the cap positive.")
    nominal = AFFIX_ALPHA * original_percent
    return cap * nominal / (cap + nominal)

# These are review findings and proposed boundaries, not claims of implemented gameplay.
# Key = the stable equipment modifier kind; all 30 set-affix kinds must be covered.
AUDIT = {
    "DirectDamage": ("补接结算", "穿戴者造成物理伤害时，在统一物理伤害结算中计算一次。", "不将法术、持续伤害或毒爆计入物理；反击是否纳入需按物理收益池明确，随机词缀仍待实现。", "先乘50%名义系数，再按同次伤害收益池的75%上限递减。"),
    "MultiHitDamage": ("先定义多段，再接入", "区分同一张牌对同一目标的多次命中与单次群攻。", "需要明确整张多段牌都加成，还是仅后续命中；不能按伤害包数量把群攻误算为多段。", "触发定义确定后再试算50%，不先改通用区间。"),
    "ArmorBreakStacks": ("重设效果含义", "项目当前使用破绽状态；现有状态枚举没有独立的破甲层数。", "需明确改为额外施加破绽，还是直接削减护甲；不得把两种效果混用。", "整数效果独立定量；不直接把1～10层通用区间除以2。"),
    "VulnerableTargetDamage": ("统一术语并补接", "按该次命中结算前目标是否具有破绽判断。", "总表统一使用当前提示中的破绽，不再混写易伤；明确穿戴者与受益伤害类型。", "百分比幅度可按原值50%试算。"),
    "FirstAttackDamage": ("明确首次计数", "按穿戴者在我方回合的真实主动攻击记录。", "需明确首张攻击牌或第一次命中；取消、无效目标、反击和自动重放不应错误消耗次数。", "每回合一次来源计数；百分比幅度待对照测试。"),
    "ArmorGain": ("已接入75%边际模型", "现行ResolveGeneratedArmorAmount按护甲产生者汇总加成，调用方决定是否允许放大。", "随机词缀汇总后按75%模型递减，固定玄甲2件10%独立加入；衍生护甲保留原来的不重复放大规则。", "六件至宝极品原始合计210%；名义105%，有效43.75%；叠加玄甲2件后53.75%。"),
    "ArmorRetention": ("补接并定义合并规则", "对应我方回合开始时穿戴者的护甲保留。", "现有玄甲4件保留50%通过套装ID结算，并未读取本词缀；需明确与全保留卡牌状态的优先级。", "词缀有效额外收益采用75%模型；与其他来源合并后的保留比例另设合法性规则，不能超过100%。"),
    "CounterDamage": ("补接反应伤害入口", "限定穿戴者实际造成的反击或格挡追加伤害。", "当前玄甲4件按专用套装入口追加伤害；词缀不能自行生成反击，也不能递归触发反应链。", "百分比幅度可按原值50%试算。"),
    "GuardReduction": ("明确承伤对象", "只在穿戴者实际为其他友方承担援护伤害时评估。", "需要在现有援护转移链中确定减伤阶段；同一伤害包不能在原目标与援护者处各减一次。", "词缀有效收益采用75%模型；最终减伤与其他来源如何合并仍需定义。"),
    "LowHealthProtection": ("重设低血保护定义", "需要明确低气血阈值、判定时点和收益是减伤还是护甲。", "当前只有低气血防护标签，没有实际触发；不把模糊保护直接实现为自动回气或无限护盾。", "阈值与效果待定，不伪造已批准数值。"),
    "Healing": ("补接治疗计算", "按穿戴者产生的治疗请求，在一次治疗结算中计算。", "区分请求治疗、实际恢复和溢出；与药方、套装固定回血的关系需明确，避免递归。", "百分比幅度可按原值50%试算。"),
    "Cleanse": ("优先重设计", "现有Cleanse/CleanseFriendlyDamageOverTime通常直接清除三种DoT。", "追加净化数量在现有全清语义下可能没有收益；建议改为净化成功后的明确附加效果，再定数值。", "不能照搬1～10净化数量，也不能用机械减半掩盖无效触发。"),
    "OverhealConversion": ("明确转换目标", "必须以实际溢出治疗为基数，并明确转换为哪种资源、给谁。", "旧青囊套的溢疗转甲方向已被替换；若作为新词缀恢复，需要独立定义，转换结果不再自触发。", "转换比例与每卡/每目标次数待定。"),
    "ManaRecovery": ("补接回复，维持固定上限", "明确是穿戴者产生的内力回复，还是穿戴者收到的回复，两者不能混算。", "装备不增加有效内力上限；不能放大固定套装返还或溢出转甲后又重复计算。", "回复幅度及适用来源待定；不改最大内力冻结。"),
    "EmergencyHealing": ("明确低血治疗条件", "按治疗前受益者的气血比例判定，明确阈值。", "与一般治疗加成在同一结算中合并；无有效治疗不能错误触发后续效果。", "阈值待定；百分比幅度可按原值50%试算。"),
    "Draw": ("重设计为有界触发", "绑定真实主动出牌事件或明确回合事件。", "不能每次抽牌又触发额外抽牌；共享手牌收益需要同名去重和每回合上限。", "优先小整数收益与限次，不能把六件合计36张再简单减半。"),
    "LowCostBonus": ("重新确认玩法方向", "现行追风以全队真实主动出牌数循环为核心，旧低费连打方向已替换。", "若保留低费条件，必须明确支付时费用还是原始费用，以及究竟增幅伤害、护甲或其他收益。", "先确定受益效果，暂不批准笼统低费牌收益倍率。"),
    "SharedEnergy": ("重设计为有界触发", "绑定明确的真实主动牌计数或回合条件。", "同名共享气力收益应设全队限次；自动重放、协战与触发效果不伪造主动出牌。", "优先每次1气与严格限次；不直接读取1～10气的通用整数区间。"),
    "ComboCount": ("替换虚拟计数收益", "现有套装、药方和任务使用真实主动出牌/记录规则。", "不能凭词缀增加已经打出的牌数；建议改为达到真实连打条件后的独立收益。", "新收益待定；旧连击计数不直接接入全队计数器。"),
    "TemporaryCostReduction": ("明确支付前减费", "在主动牌支付费用前应用，限定持有人、次数与有效期。", "需与卡牌减费、山河4件等确定顺序；不得修改已支付费用或把免费重放当一次减费触发。", "优先每次少耗1气与限次，最低费用0。"),
    "Poison": ("补接并限制重复施加", "只增强穿戴者原本就会主动施加中毒的效果。", "建议按每卡每目标去重；自然跳伤、毒爆和套装衍生效果不再次触发追加施加。", "额外点数及每回合上限独立定量。"),
    "Bleed": ("补接并限制重复施加", "只增强穿戴者原本就会主动施加流血的效果。", "建议按每卡每目标去重；多段命中不能无条件重复追加，毒爆与反击不递归。", "额外点数及每回合上限独立定量。"),
    "Burn": ("补接并限制重复施加", "只增强穿戴者原本就会主动施加灼烧的效果。", "建议按每卡每目标去重；区别主动牌、地势与套装来源，防止叠层循环。", "额外点数及每回合上限独立定量。"),
    "DamageOverTime": ("明确伤害范围及归属", "明确自然DoT跳伤、主动毒爆和自动毒爆是否受益。", "需要可靠来源归属；三种DoT本体与蚀伤额外固定伤害不能重复放大。", "统一结算入口后试算百分比减半。"),
    "StatusRetention": ("明确保留时机", "明确保留发生在毒爆消耗、自然消耗还是回合切换。", "现行蚀骨6件首次毒爆不消耗DoT；额外保留不能凭空增加状态，也不能与全保留重复返还。", "保留量不超过实际消耗量，限次与状态范围待定。"),
    "TerrainPower": ("拆清地势收益类别", "区分地势造成的伤害、护甲、治疗与抽牌/气力等资源。", "主动地势牌、回合开始自动地势与山河6件必须区分来源；不能让所有整数资源乘同一百分比。", "按收益类别独立定量，不直接全地势统一乘倍率。"),
    "TerrainCostReduction": ("明确与套装减费关系", "在地势主动牌支付前判定；地势牌指改变或触发地势的主动牌。", "现行山河4件已对首张地势牌少耗1气；需避免同条件重复减费或高阶词缀使所有地势牌免费。", "优先每次少耗1气与限次；是否另选触发张数待定。"),
    "AdjacentAllyPower": ("优先重设效果", "当前标签是相邻队友增益，现行山河套装已转向明确地势事件。", "相邻按阵位、队伍数组还是界面位置均未定义；应改为明确受益者与具体收益，不能依赖UI排序。", "新受益对象和效果待定。"),
    "FormationPower": ("优先重设效果", "阵型收益不是一个已定义的通用数值结算入口。", "需要落到具体伤害、护甲、治疗或真实地势触发；避免与地势效果/全队地势增益多次相乘。", "新效果待定，不能把模糊标签当已实现属性。"),
    "TeamTerrainPower": ("修正团队作用域设计", "需明确全队唯一来源与受益的地势效果。", "当前分系词缀统一按Owner/Passive描述符生成，此标签的全队含义未落实；不能按多个穿戴者连乘。", "团队去重、限次与效果类别先确定，再定幅度。"),
}


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def rows_of(ws):
    return [list(row) for row in ws.iter_rows(values_only=True)]


def visual_width(value) -> int:
    return sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in str(value or ""))


def sheet(wb, name, headers, rows, widths=None, freeze="C2"):
    if name in wb.sheetnames:
        index = wb.sheetnames.index(name)
        wb.remove(wb[name])
        ws = wb.create_sheet(name, index)
    else:
        ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append(list(row))
    widths = widths or [18] * len(headers)
    for column, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(column)].width = width
    thin = Side(style="hair", color="D9E1E4")
    for row in ws:
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=10, color="22343A")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
            if cell.row == 1:
                cell.fill = PatternFill("solid", fgColor="244D57")
                cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
            elif cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F1F6F6")
            if cell.row > 1 and isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0" if isinstance(cell.value, int) else "#,##0.00"
            elif cell.row > 1 and isinstance(cell.value, str) and cell.value.startswith(("未找到", "待实现", "待确认", "未接入")):
                cell.fill = PatternFill("solid", fgColor="FFF1D6")
            elif cell.row > 1 and isinstance(cell.value, str) and cell.value.startswith(("已接入", "一致（")):
                cell.fill = PatternFill("solid", fgColor="E1F0E7")
        lines = max(1, *(sum(max(1, math.ceil(visual_width(part) / max(8, widths[c.column - 1] - 2))) for part in str(c.value or "").split("\n")) for c in row))
        ws.row_dimensions[row[0].row].height = min(300, max(32 if row[0].row == 1 else 29, lines * 15 + 9))
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.print_title_rows = "1:1"
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = ws.dimensions
    return ws


def sources():
    return {key: (ROOT / path).read_text(encoding="utf-8-sig") for key, path in FILES.items()}


def ref(src, key, needle):
    offset = src[key].index(needle)
    return f"{FILES[key]}:{src[key][:offset].count(chr(10)) + 1}"


def parse_affixes(src):
    pattern = r'MakeAffix\(TEXT\("([^"]+)"\),\s*TEXT\("([^"]+)"\),\s*(?:Set|EGameXXKEquipmentSet::\w+),\s*K::(\w+),\s*U::(\w+)\)'
    records = []
    labels = dict(re.findall(r'LABEL\(\s*(\w+)\s*,\s*"([^"]+)"\s*\)', src["tooltip"]))
    for m in re.finditer(pattern, src["affix"]):
        ident, internal_name, kind, unit = m.groups()
        family = ident.split(".")[1]
        group = "历史兼容" if kind in {"MaxMana", "Speed"} and family == "Universal" else ("通用" if family == "Universal" else SETS[family])
        display = "内力上限加成（已停用，当前提示不显示）" if kind == "MaxMana" else labels[kind] + (" +{数值}%" if unit == "BasisPoints" else " +{数值}")
        records.append({"id": ident, "internal_name": internal_name, "kind": kind, "unit": unit, "group": group, "display": display})
    assert len(records) == 35 and len({r["id"] for r in records}) == 35
    assert {r["kind"] for r in records if r["group"] in SETS.values()} == set(AUDIT)
    comparisons = set(re.findall(r'Effect\.ModifierKind\s*==\s*EGameXXKEquipmentModifierKind::(\w+)', src["battle"]))
    assert comparisons == {"ArmorGain", "Invalid"}, "Battle consumers changed; re-audit before writing status claims."
    return records


def validate(wb, src):
    assert "AffixCapBasisPoints = 7500" in src["bonus_limits"]
    assert "AffixNominalDivisor = 2" in src["bonus_limits"]
    assert "Bonus.ApplyTo(BaseArmor)" in src["battle"]
    assert "InventoryDetailedAttributesButton" in src["inventory_ui"]
    affixes = parse_affixes(src)
    actual = rows_of(wb["04_词缀目录"])[1:]
    assert len(actual) == 35
    assert {r[8] for r in actual} == {r["id"] for r in affixes}
    expected = {r["id"]: r["display"] for r in affixes}
    assert all(r[2] == expected[r[8]] for r in actual), "A workbook label does not match the current tooltip source."
    assert not any(r[2] in {a["internal_name"] for a in affixes} for r in actual)
    assert len(rows_of(wb["04A_触发与接入审查"])) == 31
    assert sum(r[3].startswith("已接入产甲") for r in actual) == 1
    assert sum(r[3].startswith("未找到") for r in actual) == 29
    assert wb["05_套装效果"].max_row == 19
    assert wb["03_装备模板"].max_row == 50
    assert wb["02_宝石总表"].max_row == len(GEM_TYPES)*10+1
    gem_differences = 0
    for row in rows_of(wb["02_宝石总表"])[1:]:
        rank, token = row[3], row[1].split(".")[2]
        definition = next(d for d in GEM_TYPES if d[0] == token)
        assert row[6] == GEM_QUALITY_BONUSES[rank - 1] * definition[2]
        assert row[9] == current_gem_value(src,token,rank)
        assert row[10] == (row[9]-row[6] if row[9] is not None else None)
        gem_differences += row[9] is None or row[10] != 0
    assert wb["02A_宝石品质边际"].max_row == 11
    assert len(wb["02A_宝石品质边际"]._charts) == 1
    assert all(r[2:5] == [r[9]+240,r[10]+48,r[11]+48] for r in rows_of(wb["01_百级最终属性"])[1:])
    assert [GEM_QUALITY_BONUSES[i] - GEM_QUALITY_BONUSES[i-1] for i in range(5, 10)] == [16, 26, 37, 42, 35]
    assert all(GEM_QUALITY_BONUSES[i] == (GEM_QUALITY_BONUSES[i-1] * round(GEM_QUALITY_MULTIPLIERS[i]*10000) + 9999)//10000 for i in range(1,10))
    assert all(GEM_QUALITY_BONUSES[i]/GEM_QUALITY_BONUSES[i-1] > GEM_QUALITY_BONUSES[i+1]/GEM_QUALITY_BONUSES[i] for i in range(5,9))
    assert wb["07_掉落等级带"].max_row == 28
    assert wb["07A_挂机金币现状"].cell(28, 4).value == 24980
    assert wb["07B_分解与工具经验"].cell(11, 3).value == 20000000
    assert wb["04D_叠加上限与边际"].cell(2, 2).value == AFFIX_ALPHA
    assert wb["04D_叠加上限与边际"].cell(3, 2).value == AFFIX_BONUS_CAP
    marginal_rows = rows_of(wb["04E_一至六件边际"])[1:]
    assert len(marginal_rows) == 60
    for rank in range(1, 11):
        subset = [r for r in marginal_rows if r[0] == rank]
        assert len(subset) == 6
        assert all(0 < r[7] < AFFIX_BONUS_CAP for r in subset)
        assert all(subset[i][7] < subset[i + 1][7] for i in range(5))
        assert all(subset[i][8] > subset[i + 1][8] > 0 for i in range(5))
        assert math.isclose(sum(r[8] for r in subset), subset[-1][7], abs_tol=1e-9)
    treasure_six = next(r for r in marginal_rows if r[0] == 6 and r[2] == 6)
    assert math.isclose(treasure_six[7], 43.75, abs_tol=1e-9)
    cosmic_six = next(r for r in marginal_rows if r[0] == 10 and r[2] == 6)
    assert math.isclose(cosmic_six[7], 56.61764705882353, abs_tol=1e-9)
    assert len(wb["04F_递减曲线"]._charts) == 1
    recorded_sources = {r[0]: r[2] for r in rows_of(wb["09_来源校验"])[1:]}
    for path in FILES.values():
        assert recorded_sources.get(path) == sha(ROOT / path), f"Source snapshot changed: {path}"
    for ws in wb:
        assert ws.freeze_panes and ws.auto_filter.ref
        for row in ws:
            for cell in row:
                assert cell.data_type != "e", f"Excel error at {ws.title}!{cell.coordinate}"
                if isinstance(cell.value, str):
                    assert "\ufffd" not in cell.value, f"Encoding damage at {ws.title}!{cell.coordinate}"
    return {"affixes": 35, "class_affixes": 30, "class_affixes_with_consumer": 1, "class_affixes_without_consumer": 29, "templates": 49, "set_bonuses": 18, "gems": len(GEM_TYPES)*10, "gem_types":len(GEM_TYPES), "gem_design_runtime_differences": gem_differences, "training_stages": 27, "affix_bonus_cap_percent": AFFIX_BONUS_CAP, "marginal_examples": 60, "sheets": len(wb.sheetnames)}


def update(check_only=False):
    src = sources()
    source_hashes = {FILES[key]: sha(ROOT / FILES[key]) for key in FILES}
    original_hash = sha(BOOK)
    wb = openpyxl.load_workbook(BOOK)
    if check_only:
        print(json.dumps({"ok": True, "file": str(BOOK), **validate(wb, src)}, ensure_ascii=False))
        wb.close()
        return
    stamp = datetime.now(timezone(timedelta(hours=8)))
    date = stamp.strftime("%Y-%m-%d")
    snapshot = {name: rows_of(wb[name]) for name in wb.sheetnames}
    old_treasure = {r[5]: r[6] for r in snapshot["02_宝石总表"][1:] if r[3] == 6}
    # Rebase only gem contributions. Keep the originally frozen numbers for comparison.
    target_rows = snapshot["01_百级最终属性"]
    if len(target_rows[0]) == 9:
        target_rows[0] += ["2026-09-04原设计生命", "2026-09-04原设计攻击", "2026-09-04原设计防御"]
        for row in target_rows[1:]:
            row += row[2:5]
    for row in target_rows[1:]:
        row[2] += 4 * (GEM_QUALITY_BONUSES[5] * GEM_HEALTH_MULTIPLIER - old_treasure["生命"])
        row[3] += 4 * (GEM_QUALITY_BONUSES[5] - old_treasure["攻击"])
        row[4] += 4 * (GEM_QUALITY_BONUSES[5] - old_treasure["防御"])
        row[8] = "2026-09-09宝石品质修订后的设计值；非实测"
    for row in snapshot["06_至宝配装基准"][1:]:
        for kind, value in [("攻击", GEM_QUALITY_BONUSES[5]), ("防御", GEM_QUALITY_BONUSES[5]), ("生命", GEM_QUALITY_BONUSES[5]*GEM_HEALTH_MULTIPLIER)]:
            if row[0] == kind + "至宝宝石":
                row[2], row[3] = f"+{value}{kind}", f"+{4 * value}{kind}"
    original_drop_bands = [r[:6] for r in snapshot["07_掉落等级带"]]
    commit = subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip()
    branch = subprocess.check_output(["git", "branch", "--show-current"], cwd=ROOT, text=True).strip()
    version = int(re.search(r"CurrentSaveVersion\s*=\s*(\d+)", src["save"]).group(1))
    affixes = parse_affixes(src)
    current_affix_rows, audit_rows = [], []
    for index, a in enumerate(affixes, 1):
        is_class = a["group"] in SETS.values()
        if a["kind"] == "MaxMana":
            status, trigger, action = "已停用：有效内力贡献为0", "属性投影忽略；统一装备提示不显示", "保留旧档解析，不进入新随机池"
        elif a["kind"] == "Speed":
            status, trigger, action = "历史兼容属性投影", "仅旧实例兼容；不进入新随机池", "保留旧档解析和迁出路径"
        elif not is_class:
            status, trigger, action = "已接入装备属性投影", "按穿戴者汇总到装备属性预算", "本轮分系审查不自动改动通用属性"
        elif a["kind"] == "ArmorGain":
            status, trigger, action = "已接入产甲增幅（源码核对）", "按护甲产生者匹配；调用方允许放大时生效", AUDIT[a["kind"]][0]
        else:
            status, trigger, action = "未找到战斗数值消费者（源码核对）", "当前仅Owner/Passive描述符；实际触发尚未接入", AUDIT[a["kind"]][0]
        pool = "旧档兼容，不进入新随机池" if a["group"] == "历史兼容" else ("通用随机池" if a["group"] == "通用" else "对应套装随机池")
        evidence = ref(src, "affix", f'TEXT("{a["id"]}")') + "\n" + ref(src, "tooltip", "FString GameXXKEquipmentTooltipPresentation::AffixLine")
        current_affix_rows.append([index, a["group"], a["display"], status, pool, trigger, action, a["unit"], a["id"], evidence])
        if is_class:
            _, boundary, exclusions, budget = AUDIT[a["kind"]]
            if a["unit"] == "BasisPoints":
                budget += "\n叠加方案：α=0.5、有效额外收益上限75%，E=75S/(75+S)，按本次同一收益池汇总后折算；详见04D。"
            audit_rows.append([index, a["group"], a["display"], status, boundary, exclusions, budget,
                "产甲已接入75%模型与最终属性显示" if a["kind"] == "ArmorGain" else "叠加模型已记录；本词缀仍待触发定义与运行时接入", a["id"]])
    sheet(wb, "04_词缀目录", ["序号", "归属", "当前装备提示的效果文本", "当前运行时状态", "生成状态", "现行触发/作用域", "本轮审查处理", "当前数据单位", "稳定词缀ID", "源码定位"], current_affix_rows, [8,12,39,35,25,43,27,18,47,65], "D2")
    sheet(wb, "04A_触发与接入审查", ["目录序号", "归属", "效果描述", "当前接入状态", "需要明确或保留的触发边界", "排除/叠加/现行机制约束", "数值处理建议", "建议状态", "稳定词缀ID"], audit_rows, [10,12,37,31,51,61,52,32,47], "D2")
    range_rows = []
    for q, name in enumerate(QUALITY, 1):
        low = (q + 1) * (q + 2) / 2
        high = low + q + 1
        range_rows.append([q,name,min(q,5),1+max(0,q-5),low,high,low*AFFIX_ALPHA,high*AFFIX_ALPHA,(q+1)//2,q,"分系百分比先乘0.5，再按04D做75%上限的边际递减；本列是名义值。产甲已接入，其余29条仍待接入；整数类继续独立重设计。"])
    ws = sheet(wb, "04B_品质与数值基线", ["阶", "品质", "词缀条数", "宝石孔数", "当前百分比下限", "当前百分比上限", "方案单条名义下限", "方案单条名义上限", "当前整数下限", "当前整数上限", "状态与适用范围"], range_rows, [7,12,12,12,18,18,21,21,16,16,82])
    for row in ws.iter_rows(min_row=2,min_col=5,max_col=8):
        for cell in row: cell.number_format = '0.00"%"'
    sheet(wb, "04C_名称显示入口", ["入口", "代码当前行为", "名称状态", "核对范围", "源码定位"], [
        ["统一装备提示", "使用ModifierLabel + 数值/百分比；最大内力词缀直接隐藏。", "已直接描述效果", "这是总表词缀名称的取值来源。", ref(src,"tooltip","FString GameXXKEquipmentTooltipPresentation::AffixLine")],
        ["桌面背包装备格", "调用统一装备提示Build/Bind。", "已直接描述效果", "当前默认桌面工作台入口。", ref(src,"workbench","SlotButton->SetToolTipText(FText::GetEmpty());GameXXKEquipmentTooltipPresentation::Bind")],
        ["桌面商店结果卡", "调用统一装备提示Build/Bind。", "已直接描述效果", "当前桌面商店购买结果入口。", ref(src,"workbench","GameXXKEquipmentTooltipPresentation::Bind(ResultCard")],
        ["旧商店控件", "仍有读取Affix->DisplayName的格式化代码。", "仍保留旧短名路径", "源码事实；本轮未验证旧控件在当前流程的可达性。", ref(src,"old_shop","*Affix->DisplayName.ToString()")],
        ["旧伙伴名册控件", "仍有读取Affix->DisplayName的格式化代码。", "仍保留旧短名路径", "源码事实；本轮未验证旧控件在当前流程的可达性。", ref(src,"old_roster","*Affix->DisplayName.ToString()")],
        ["词缀目录内部字段", "DisplayName仍保留旧别名；稳定ID用于实例和存档解析。", "内部字段未改名", "总表主名称不再采用旧别名；不把显示更新误写成C++字段已修改。", FILES["affix"]],
    ], [24,64,28,67,70], "B2")

    model_rows = [
        ["基础系数 α", AFFIX_ALPHA, "沿用50%方案；原始词缀百分比先乘此系数。", "产甲运行时已接入；其余29条仍待接入"],
        ["有效额外加成上限 C", AFFIX_BONUS_CAP, "单位为百分比数值，例如75代表+75%；上限是渐近上限，不代表六件必然达到。", "用户选定75%上限"],
        ["原始输入 P", "P=Σpᵢ", "只合计本次实际满足条件、属于同一收益池的原始词缀百分比。", "先确认触发，再参与计算"],
        ["名义合计 S", "S=αP=0.5×Σpᵢ", "先汇总后折算，装备顺序和批次分组不能改变结果。", "不逐件分别递减后再相加"],
        ["有效额外收益 E", "E=C×S/(C+S)=75S/(75+S)", "S=0时E=0；S增加时E持续增加，但始终低于75%。", "α和C只作用于分系随机百分比词缀"],
        ["新增一条的边际收益", "ΔE=C²×ΔS/[(C+S)(C+S+ΔS)]", "已有S越高，同样一条词缀带来的增量越小；ΔS=0.5×新增原始百分比。", "边际的单位为百分点"],
        ["计算半饱和点", "S=C=75 → E=37.5%", "名义合计达到上限数值时，实际收益恰好达到上限的一半。", "用于核对递减曲线"],
        ["计算顺序", "有效触发筛选 → 按收益池相加 → α=0.5 → 递减 → 原有效果结算", "先在高精度中计算，最后按原伤害/治疗/产甲入口的规则取整；不逐件提前取整。", "产甲当前为最终向上取整"],
        ["基点实现等价式", "有效基点=7500B/(15000+B)", "B为合计原始基点（100基点=1%）；避免先除2丢失半个基点，最后才取整。", "共享EquipmentBonusRules已接入产甲与详细属性"],
        ["同一收益池", "同一次结算、同一个收益类别共同使用75%上限", "例如本次适用的直接/多段/破绽/首次攻击增伤进入同一个伤害词缀池，不能各吃一次75%后连乘。", "未定义的地势/阵型收益先确定输出类型"],
        ["个人与团队来源", "个人按穿戴者；团队按唯一来源规则", "同一作用实例只进入一次；多名穿戴者的团队同名收益不能各自递减后再相加绕过上限。", "来源与触发去重仍须先补齐"],
        ["既有套装与天赋", "独立来源按各自现行规则计算", "75%是本次分系词缀有效额外收益的上限；不代表所有套装、天赋和基础值合计只能提高75%。", "不把18条套装机制误当随机词缀"],
        ["减伤/保留/转化等比例", "仍需明确各自的最终比例结算规则", "这类效果不能盲套基础值×(1+E)；禁止出现超过100%的减伤、护甲保留或重复转化。与已有比例的合并方式仍需逐条定。", "75%模型不代替效果定义"],
        ["整数收益", "抽牌、气力、减费、状态层数不套本百分比公式", "继续采用独立的次数与小整数预算，不把0.5张牌或0.5气力当作已完成设计。", "触发重设计仍待完成"],
        ["存档与成色", "保留原始词缀阶数、Magnitude和随机成色", "α与递减在有效计算层应用；原始词缀品质区间、分解成色评价不因战斗折算而被改写。", "旧档迁移与显示实现时需核对"],
        ["显示要求", "详细属性展示当前角色的有效合计", "原始单件数值与全身实际加成区分；未接入词缀显示零贡献与未生效标记。", "详细属性按钮已接入；现有单件提示仍显示原始值；百分比换装边际提示待补"],
        ["数值例：六件至宝极品", "6×35%×0.5=105%；E=43.75%", "每件都带同一收益且本次全部满足触发条件的理论例。", "不是任意六件至宝的固定收益"],
        ["数值例：六件宇宙极品", "6×77%×0.5=231%；E≈56.62%", "继续堆高品质有收益，但不会达到或超过75%。", "60组配装实例见04E"],
        ["产甲例", "基础100＋玄甲2件10%＋上述词缀43.75% → 向上取整为154护甲", "仍依赖原产甲入口允许装备放大；不把衍生护甲反复放大。", "已写入产甲结算与对应回归；实测证据见本轮验收"],
    ]
    sheet(wb, "04D_叠加上限与边际", ["项目", "参数/公式", "解释与边界", "实现状态"], model_rows, [29,75,105,49], "B2")
    marginal_rows = []
    for q,name in enumerate(QUALITY,1):
        original_max=(q+1)*(q+4)/2
        previous=0.0
        for pieces in range(1,7):
            nominal=pieces*original_max*AFFIX_ALPHA
            effective=diminishing_bonus(pieces*original_max)
            marginal_rows.append([q,name,pieces,original_max,original_max*AFFIX_ALPHA,nominal,AFFIX_BONUS_CAP,effective,effective-previous,"同一收益池、每件各一条极品且本次都生效的公式示例；产甲已接入，其他词缀需先落实触发"])
            previous=effective
    ws=sheet(wb,"04E_一至六件边际",["阶","词缀品质","件数","原始单条上限","α后单条名义值","名义合计S","上限C","实际额外收益E","新增这一件的边际","条件与状态"],marginal_rows,[7,14,9,20,23,20,15,24,27,82],"D2")
    for row in ws.iter_rows(min_row=2,min_col=4,max_col=8):
        for cell in row: cell.number_format='0.00"%"'
    for row in ws.iter_rows(min_row=2,min_col=9,max_col=9):
        row[0].number_format='0.00" 个百分点"'
    curve_rows=[[nominal,diminishing_bonus(nominal/AFFIX_ALPHA),AFFIX_BONUS_CAP] for nominal in range(0,601,5)]
    curve_ws=sheet(wb,"04F_递减曲线",["名义合计S（已乘0.5）","有效额外收益E","75%上限"],curve_rows,[27,24,19],"A2")
    for row in curve_ws.iter_rows(min_row=2):
        for cell in row: cell.number_format='0.00"%"'
    chart=LineChart()
    chart.title="75%上限下的边际递减（设计方案）"
    chart.y_axis.title="有效额外收益（%）"
    chart.x_axis.title="名义合计S（%，已乘0.5）"
    chart.y_axis.scaling.min=0
    chart.y_axis.scaling.max=80
    chart.x_axis.tickLblSkip=10
    chart.y_axis.majorUnit=10
    chart.height=13
    chart.width=25
    chart.legend.position="b"
    chart.add_data(Reference(curve_ws,min_col=2,max_col=3,min_row=1,max_row=curve_ws.max_row),titles_from_data=True)
    chart.set_categories(Reference(curve_ws,min_col=1,min_row=2,max_row=curve_ws.max_row))
    chart.series[0].graphicalProperties.line.solidFill="244D57"
    chart.series[0].graphicalProperties.line.width=26000
    chart.series[1].graphicalProperties.line.solidFill="D99A33"
    chart.series[1].graphicalProperties.line.prstDash="dash"
    curve_ws.add_chart(chart,"E2")
    sheet(wb,"04G_详细属性口径",["项目","当前显示值","来源与边界"],[
        ["入口","背包 → 属性 → 详细属性；提供返回基础属性按钮与滚动列表。","跟随当前角色刷新；保存背包会话内的打开状态和滚动位置。"],
        ["暴击率","永久天赋最终概率，0～20%。","直接攻击使用的概率；不加未接入的词缀。"],
        ["暴击伤害倍率","150%基础＋0～50%永久天赋；最终150～200%。","显示最终倍率，不把天赋额外50%错显示成完整倍率。"],
        ["物理伤害加成","物理宝石合并名义值，经75%边际后再与永久天赋相乘，法术另算。","主动直伤、反击、各元素直伤和持续伤害分别显示；随机直伤词缀仍无消费者。"],
        ["护甲获得量加成","随机产甲词缀×0.5与产甲宝石合池边际，再加固定套装。","与战斗共用EquipmentBonusRules；无宝石时6条至宝极品＋玄甲2件仍为+53.75%。"],
        ["回合开始护甲保留","玄甲4件有效50%；无该套装为0%。","与尚未接入的随机护甲保留词缀分开；不计本局临时全保留状态。"],
        ["宝石贡献","固定值相加；属性百分比显示有效百分比与实际属性增量。","百分比作用于裸身＋装备＋固定宝石，再计算天赋。至宝4攻/4防/4生命仍为128/128/640。"],
        ["条件套装","采用当前完整描述，标条件生效；团队效果按现有唯一来源规则。","未出战角色的团队来源不计入当前队伍。"],
        ["其他随机词缀","无有效消费者的加成显示0；装备实际持有时标未生效。","不能把原始值或理论递减值显示成已生效；不擅自补造触发。"],
        ["临时状态","不包括本局状态、牌面倍率、敌方防御及临时地势。","本页是永久配装与天赋视图，不是对任意目标的即时伤害预测。"],
    ],[26,102,115],"B2")

    # Document the chosen quality curve before changing runtime; keep a source comparison.
    current_bonuses, current_health = current_gem_values(src)
    gem_rows = []
    for rank in range(1,11):
        for token,label,scale,unit in GEM_TYPES:
            value = GEM_QUALITY_BONUSES[rank-1]*scale
            row = [len(gem_rows)+1,f"Item.Gem.{token}.{QUALITY_TOKENS[rank-1]}",QUALITY[rank-1]+label+"宝石",rank,QUALITY[rank-1],label,value,unit,
                   "固定值正常相加" if unit=="平坦值" else "R0+(75-R0)S/(75-R0+S)，最终抗性-25～75%" if unit=="名义抗性百分点" else "名义值先汇总，再75S/(75+S)；物理/法术分开"]
            current = current_gem_value(src,token,rank)
            gem_rows.append(row+[current,current-value if current is not None else None,
                "一致（仅目录源码数值；战斗与显示按专项验收）" if current==value else "待实现：新类型或数值尚未同步代码"])
    headers = snapshot["02_宝石总表"][0][:9]
    headers[6], headers[8] = "批准设计值", "设计口径"
    sheet(wb,"02_宝石总表",headers+["当前代码值","代码－设计","当前状态"],gem_rows,[8,44,24,7,12,10,16,12,44,16,16,40],"D2")
    gem_curve_rows = []
    for i, value in enumerate(GEM_QUALITY_BONUSES):
        delta = value - GEM_QUALITY_BONUSES[i-1] if i else None
        gem_curve_rows.append([i+1, QUALITY[i], value, value*GEM_HEALTH_MULTIPLIER, delta,
            delta/GEM_QUALITY_BONUSES[i-1]*100 if i else None,
            "前五档保留" if i < 5 else "升级倍率从×2均匀降到×1.25；总值持续增长",
            "数量不递减；同属性宝石正常相加", GEM_QUALITY_MULTIPLIERS[i]])
    gem_ws = sheet(wb,"02A_宝石品质边际",["阶","品质","单颗攻防","单颗生命","新增攻防","实际相对上一档提升%","设计范围","多颗叠加","设计升级倍率（逐阶向上取整）"],gem_curve_rows,[8,12,16,16,17,27,49,51,38],"C2")
    for row in range(2, 12):
        gem_ws.cell(row, 9).number_format = '"×"0.00##'
    percentage_rows=[]
    for rank,step in enumerate(GEM_QUALITY_BONUSES,1):
        stat=step*.25;mechanic=step*.5
        percentage_rows.append([rank,QUALITY[rank-1],stat,75*stat/(75+stat),mechanic,75*mechanic/(75+mechanic),
            "本表后两列仅为伤害/机制增益；抗性宝石另按佩戴者基础抗性折算"])
    percent_ws=sheet(wb,"02B_百分比宝石数值",["阶","品质","属性宝石名义%","属性单颗有效%","机制宝石名义%","机制单颗有效%","叠加规则"],percentage_rows,[8,12,22,22,22,22,74],"C2")
    for row in percent_ws.iter_rows(min_row=2,min_col=3,max_col=6):
        for cell in row:cell.number_format='0.00"%"'
    resistance_rows=[]
    for rank,step in enumerate(GEM_QUALITY_BONUSES,1):
        nominal=step*.5
        resistance_rows.append([rank,QUALITY[rank-1],nominal,*[base+(75-base)*nominal/(75-base+nominal) for base in [0,10,20,40]],
            "每种抗性宝石分别只提高对应元素抗性；此处列的是最终抗性"])
    rw=sheet(wb,"02D_抗性宝石折算",["阶","品质","名义抗性百分点","基础0%时最终%","基础10%时最终%","基础20%时最终%","基础40%时最终%","范围"],resistance_rows,[8,13,24,24,24,24,24,78],"C2")
    for row in rw.iter_rows(min_row=2,min_col=3,max_col=7):
        for cell in row:cell.number_format='0.00'
    sheet(wb,"02E_物理法术与抗性",["项目","结算口径","当前边界"],[
        ["物理伤害","原始物理伤害逐包减防御，再由护甲吸收，余量扣生命。","直接伤害宝石对外改称物理伤害，内部DirectDamage ID保留。"],
        ["法术伤害","火/冰/雷法术不扣防御，只按对应抗性减免，再由护甲吸收。","与物理宝石分开；不改变原倍率和触发次数。"],
        ["抗性边际","R0+(75-R0)S/(75-R0+S)，减抗后最终夹到-25～75%。","R0为基础抗性；当前正向S来自抗性宝石，未额外配置抗性天赋或随机词缀。"],
        ["灼烧","自然跳伤与主动引燃吃火抗，来源宝石先合池。","继续按持续伤害绕护甲；血/毒/蚀伤保持既有通道。"],
        ["单位基准","13角色模板、21怪物分别配置三抗。","见怪物总表与docs/design/2026-09-09-resistance-discussion；不随等级或难度普涨。"],
        ["内力吸取","金钱鼠8/8/8、巨蟾6/6/6；扣到0为止。","资源效果，不吃伤害增幅或抗性；巨蟾命中后一次且跟随援护目标。"],
    ],[25,100,108],"B2")
    sheet(wb,"02C_宝石收益与表现",["类别","生效范围","边界与归属"],[
        ["固定攻防生命","固定值直接相加。","10档共用每类型一张图标。"],
        ["属性百分比三类","裸身＋装备＋固定宝石，再计算永久天赋。","攻击/防御/生命各自独立75%边际，不共享一个总属性池。"],
        ["物理伤害","本人物理攻击及物理反击；法术另算。","真正适用的物理/反击加成合池；固定独立伤害与持续包分别计算。"],
        ["护甲获得量","本人生成的护甲。","宝石名义值与随机产甲词缀×0.5合池；固定套装独立，复制甲不重复放大。"],
        ["治疗效果","本人主动治疗及其重放。","不增加吸血返还、复活、路线恢复与固定套装治疗。"],
        ["反击伤害","已有反击、玄甲格挡追击。","不增加触发次数；与适用元素先合池。"],
        ["火焰伤害","本人火系直伤、灼烧自然伤害和主动引燃。","灼烧与持续宝石合池；不增加层数。"],
        ["持续伤害","流血、中毒、灼烧及其主动引爆。","自然按各施加者贡献；主动引爆按触发者和原始层数；不放大蚀伤。"],
        ["冰霜伤害","冰系直伤、冰爆和冰系阵赏。","耗甲、返甲、回内不变。"],
        ["雷击伤害","雷系攻击与实际落雷。","标记、雷击和出牌次数不变。"],
        ["品质显示","名称、道具底、tooltip复用装备品质Shader。","镶嵌后的每颗文本按自己品质显示；三种属性百分比独有同尺寸浅色浮雕%。"],
        ["验证状态","已完成14张纹理导入；接入代码与专项测试推进中。","以docs/production/2026-09-09-gem-attributes-progress.md的最终报告为准。"],
    ],[22,79,111],"B2")
    gem_chart = LineChart()
    gem_chart.title = "宝石品质：倍率从×2降到×1.25，总值持续增长"
    gem_chart.y_axis.title, gem_chart.x_axis.title = "攻防数值", "品质"
    for col in (3, 5):
        gem_chart.add_data(Reference(gem_ws,min_col=col,max_col=col,min_row=1,max_row=11),titles_from_data=True)
    gem_chart.set_categories(Reference(gem_ws,min_col=2,min_row=2,max_row=11))
    gem_chart.height, gem_chart.width = 13, 25
    gem_ws.add_chart(gem_chart,"K2")

    slot_names = dict(re.findall(r'\{EGameXXKEquipmentSlot::(\w+), TEXT\("\w+"\), NSLOCTEXT\("GameXXKEquipment", "[^"]+", "([^"]+)"\)\}',src["catalog"]))
    curves = {}
    for slot, body in re.findall(r'case EGameXXKEquipmentSlot::(\w+):(.*?)(?=break;)',src["catalog"],re.S):
        fields = {}
        for stat, args in re.findall(r'Coefficients\.(\w+) = Curve\(([^)]+)\)',body):
            values = [int(x.strip()) for x in args.split(",")]
            first, num, den = (values+[0,0])[:3]
            fields[stat] = first + (19*num//den if num and den else 0)
        if fields: curves[slot] = fields
    legacy = {m[0]: (m[1],m[2],list(map(int,m[3:]))) for m in re.findall(r'MakeLegacyDefinition\(TEXT\("([^"]+)"\), NSLOCTEXT\("GameXXKEquipment", "[^"]+", "([^"]+)"\), EGameXXKEquipmentSlot::(\w+), LegacyStats\((\d+), (\d+), (\d+), (\d+)\)',src["catalog"])}
    template_rows = []
    for original in snapshot["03_装备模板"][1:]:
        row = original[:8]
        if row[1].startswith("Equipment."):
            _, group, slot = row[1].split(".")
            row[2], row[3], row[4] = SETS[group]+slot_names[slot], SETS[group], slot_names[slot]
            stats = curves[slot]
            extra = [20,stats.get("MaxHealth",0),stats.get("Attack",0),stats.get("Defense",0),stats.get("Speed",0),"模板名称已核对；批准减半设计仍待落实"]
        else:
            name, slot, stats = legacy[row[1]]
            row[2],row[4] = name,slot_names[slot]
            extra = ["历史快照",stats[0],stats[2],stats[3],0,"旧档兼容；内力快照不增加有效内力"]
        template_rows.append(row+extra)
    sheet(wb,"03_装备模板",snapshot["03_装备模板"][0][:8]+["代码曲线封顶等级","代码基础生命","代码基础攻击","代码基础防御","代码基础速度","当前核对状态"],template_rows,[8,43,23,12,12,15,64,28,18,16,16,16,16,45],"D2")

    bonus_pattern = r'MakeBonus\(TEXT\("([^"]+)"\), TEXT\("([^"]+)"\), EGameXXKEquipmentSet::(\w+), (\d+), K::(\w+), S::(\w+), H::(\w+), U::(\w+), (\d+)(?:, (\d+))?(?:, (\d+))?\)'
    bonuses = re.findall(bonus_pattern,src["sets"])
    assert len(bonuses) == 18
    set_rows = []
    old_by_id = {r[0]:r for r in snapshot["05_套装效果"][1:]}
    for ident, desc, group, count, kind, scope, hook, unit, value, limit, secondary in bonuses:
        notes = "按穿戴者记录。" if scope == "Owner" else "全队唯一来源，按现行选择规则去重。"
        if group in {"QingNang","ZhuiFeng"}: notes += "高件数执行完整合并效果，不把2/4/6件描述重复各算一遍。"
        if group == "ShiGu": notes += "按现有每卡/目标、每回合毒爆规则结算；区别随机词缀。"
        set_rows.append([ident,SETS[group],int(count),"穿戴者" if scope=="Owner" else "全队唯一",old_by_id[ident][4],desc,"已接入（源码核对；本轮未重跑）",notes,ref(src,"sets",f'TEXT("{ident}")')])
    sheet(wb,"05_套装效果",["描述符ID","套装","件数","作用域","触发","当前代码完整效果","当前状态","去重与计数说明","源码定位"],set_rows,[24,12,8,17,32,79,37,72,65],"C2")

    terms = [
        ["主动出牌","玩家实际支付并打出的牌；自动重放、协战、反击、DoT和任务衍生效果不伪造主动出牌计数。"],
        ["地势牌","包含改变地势或触发地势效果的主动牌；按当前卡牌定义识别。"],
        ["全队唯一","同名团队套装只生效一个来源；不按装备件数或多个穿戴者重复触发。"],
        ["破绽","当前界面使用的名称，对应Vulnerability；不新增一个同义的易伤/破甲状态。"],
        ["分系随机词缀","每件装备随机生成的效果；进入效果列表不代表已有战斗结算入口。"],
        ["2/4/6件套装","按穿戴件数启用的独立机制，和同系随机词缀分别核对。"],
        ["源码已接入","已定位计算或触发入口；不等于本轮做过实机A/B或平衡验收。"],
        ["审查建议","尚未写入运行时代码的设计方向；阈值或上限待定时明确留待确认。"],
        ["词缀显示数值","基点值除以100显示为两位小数百分比；整数照原值显示；最大内力词缀在统一提示中隐藏。"],
    ]
    sheet(wb,"05A_套装术语",["术语","当前口径"],terms,[24,120],"A2")

    drop_rows = []
    for row in original_drop_bands[1:]:
        note = "待实现：按关卡等级带掉落尚未接入" if row[4] <= 100 else "待实现：101～135级还超过当前实例上限"
        drop_rows.append(row+["按宝箱SourceItemLevel生成；普通游历结算记录玩家等级",100,note])
    sheet(wb,"07_掉落等级带",original_drop_bands[0]+["当前代码掉落规则","当前实例/宝箱等级上限","当前状态"],drop_rows,[12,12,12,14,14,14,78,24,54],"D2")
    assert "NormalGoldMultiplier = 10" in src["training"]
    gold_growth = float(re.search(r"PostNormalGoldGrowthPerStage\s*=\s*([\d.]+)",src["training"]).group(1))
    gold_rows=[]
    for difficulty_index,difficulty in enumerate(["普通","困难","地狱"]):
        for n in range(1,10):
            growth=(difficulty_index-1)*9+n
            gold=(18+3*n)*10 if difficulty_index==0 else math.floor(450*gold_growth**growth+0.5)
            gold_rows.append([difficulty,f"{(n-1)//3+1}-{(n-1)%3+1}",(difficulty_index*9+n)*5,gold,10*(difficulty_index+1)+2*n,math.floor(gold*4.5+0.5),f"已接入：普通×10；困难至地狱连续每关×{gold_growth:g}",ref(src,"training","constexpr int32 NormalGoldMultiplier")])
    sheet(wb,"07A_挂机金币现状",["难度","关卡","敌人等级","每波基础金币","每波基础经验","满在线金币天赋每波金币","当前状态","源码定位"],gold_rows,[12,12,15,21,21,31,70,68],"C2")
    sale_bases=[5000,12500,30000,75000,200000,500000,1250000,3000000,7500000,20000000]
    tool_rows=[[i+1,name,sale_bases[i],10,9**i,math.floor(9**i*3.5+0.5),"金币：原方案×5已选，尚未实装；工具经验成长方案待确认"] for i,name in enumerate(QUALITY)]
    sheet(wb,"07B_分解与工具经验",["阶","品质","已选分解基础价提案","当前每件基础金币","当前每件基础工具经验","当前单件满工具天赋经验","方案/实现状态"],tool_rows,[8,13,26,24,27,31,87],"C2")
    sheet(wb,"07C_经济与工具边界",["项目","当前代码","本轮选择或待办","状态/说明"],[
        ["分解基础金币","每件10金币；工具金币天赋最高×3.5。","分品质基础价采用原方案×5，见07B；尚未实现。","与当前实际数值分栏记录。"],
        ["分解金币修正","当前未按装备等级或词缀计价。","候选：基础价×等级系数×[1+5%×(平均词缀阶-1)+10%×平均成色]×工具金币天赋。等级先限制1～100，等级系数=1+(等级-1)/99，范围1～2。","词缀重设计后再确认成色换算；不是当前到账公式。"],
        ["分解材料与宝石","每件1强化石、1洗炼砂；随装备删除镶嵌宝石。","宝石退回与强化投入返还属于待办，未写入代码。","不能把建议退还写成已经实现。"],
        ["工具等级","上限10；升下一级经验=当前等级×100；1→10共4500经验。","是否重算满级周期尚待选择。","不把建议的7～10天写成已批准目标。"],
        ["工具经验","品质经验=9^(品质阶-1)；分解/合成/强化/洗炼等共用。","分解预览、具体到账经验提示和升级结果显示待补。","当前不朽基础经验6561，超过1→10总需求。"],
        ["工具满级","满级后剩余经验清零；结果字段仍可记录公式奖励。","显示时需区分计算奖励与实际可累计经验。","当前行为。"],
        ["工具合成等级","工具档位1生成1～10级；档位n生成(n-1)×10～n×10级。","装备合成需9件同品质；合成产物等级由所选工具档位决定。","当前行为。"],
        ["商店价格","装备包100000；高级箱100000；普通箱25000；宝石包100000。","当前装备包品质概率普通70%／稀有25%／珍稀5%。","当前行为。"],
        ["商店产物等级","读取工具等级本身；工具上限10，所以常规商店产物最高10级。","不同于工具合成档位的最高100级，不能把两者混写。","当前映射，未擅自修正。"],
        ["天赋目标","用户期望全天在线约45天点满；这不是实测完成周期。","挂机金币已改每关×1.25；加入分解×5后的时间仍是带条件估算。","分解价和工具经验未实装，不能给出已达成45天的结论。"],
        ["金币存储","PlayerGold为int32，最大2147483647；工具结算有溢出拒绝。","×5方案的9件极品宇宙最高约1953000000，仍须检查与已有余额相加。","大额收益实现前核对完整资金链。"],
    ],[25,85,100,53],"B2")

    differences=[
        ["词缀玩家显示","统一提示已使用效果标签+数值；目录与旧控件仍留短名。","本表35条名称全部同步统一提示；旧控件另列。","显示来源已核对",FILES["tooltip"]],
        ["分系随机词缀接入","30条均可生成并投影；仅护甲获得量定位到实际消费者。","逐条审查29条未接入效果；不是18条套装都失效。","29条未找到消费者；1条已接入",FILES["battle"]],
        ["分系词缀强度","存档与生成区间保持原值；随机产甲汇总后已采用75%模型。","α=0.5，E=75S/(75+S)；固定套装独立计算。其余29条仍需触发定义与接入。","产甲已写入代码，实测证据见本轮验收",FILES["bonus_rules"]],
        ["装备基础贡献","当前模板曲线仍在20级封顶，实例上限100。","原批准的生命/攻击/防御贡献约减半与百级设计基准继续保留。","待实现，未用旧代码覆盖设计目标",FILES["catalog"]],
        ["宝石数值",f"固定攻防：{current_bonuses}；生命{current_health}倍。","17类×10品质；属性百分比曲线×0.25%，机制曲线×0.5%；75%边际。每类单图，品质Shader。",f"170项中{sum(r[9] is None or r[10] != 0 for r in gem_rows)}项目录数值待同步，详见02/02A表",FILES["gems"]],
        ["至宝孔位","每件2孔，六件共12孔。","与批准配装基准一致。","一致（源码核对）",FILES["gems"]],
        ["装备内力","有效最大内力不受装备影响；统一提示隐藏相关旧词缀。","保留旧数据解析，禁止新增有效上限。","已接入忽略规则",FILES["equipment"]],
        ["掉落等级带","宝箱取SourceItemLevel；常规结算记录玩家等级；均限制1～100。","批准的上一关+1至当前关等级与101～135装备尚未落实。","待实现，07表保留设计并列现状",FILES["training"]],
        ["六套2/4/6件效果","18条描述和消费者按当前源码核对；当前存档版本为"+str(version)+"。","完整效果直述，去掉高件数继承的阅读歧义。","已接入；本轮未重跑战斗验证",FILES["sets"]],
        ["百级最终属性","本轮未用完整配装/词缀/强化条件重做运行时同值验证。","01、06表仅按新至宝宝石贡献修订：2026-09-04原设计+240生命/+48攻击/+48防御，原值另列。","设计基准，非当前实测属性","docs/design/2026-09-09-gem-quality-marginal-curve.md"],
        ["分解与工具经验","分解仍10金币；品质工具经验仍9的幂；提示未显示具体到账经验。","分解基础价原方案×5已选；工具经验与预览仍待补。","未实装方案，详见07B/07C",FILES["tools"]],
    ]
    sheet(wb,"08_实现差异",["项目","当前代码事实","目标/本轮文档处理","状态","主要源码"],differences,[24,82,91,50,69],"B2")
    source_rows=[[path,"本轮读取的当前工作区源码",digest] for path,digest in source_hashes.items()]
    source_rows += [["工作区分支",branch,"本次文档更新已获当前分支授权"],["HEAD",commit,"源码哈希以当前未提交工作区为准"],["核对时间",stamp.isoformat(timespec="seconds"),"仅源码/工作簿核对，本轮未运行UBT/PIE"]]
    for path in ["docs/superpowers/specs/2026-09-03-card-monster-progression-rebalance-design.md","docs/superpowers/specs/2026-09-04-xuanjia-shanhe-set-design.md"]:
        source_rows.append([path,"既有批准设计来源",sha(ROOT/path)])
    sheet(wb,"09_来源校验",["源文件/项目","用途/值","SHA256或说明"],source_rows,[84,73,76],"A2")
    intro=[
        ["更新日期",date+"；按当前工作区代码核对后更新。"],
        ["工作簿定位","现行代码事实与批准设计基准共同记录；审查建议、待实现项均明确标记。"],
        ["词缀命名","名称采用当前统一装备提示的直接效果描述；不再以目录内部两字别名作为表中名称。"],
        ["当前显示状态","桌面背包和商店结果卡已使用统一效果文本；旧商店/旧伙伴控件仍有短名路径，详见04C。"],
        ["当前接入状态","35条目录=3条通用+2条兼容+30条分系。分系中1条已定位数值消费者，29条未找到消费者；不是套装18条失效。"],
        ["触发与强度审查","04A逐条记录现状与建议；α=0.5与75%上限公式/60组例子/曲线见04D～04F。产甲已接入共享结算，其余29条仍明确待办；详细属性见04G。"],
        ["百级批准基准","01表保留100级、六件至宝与4攻/4防/4生命配装，仅按新宝石贡献修订；原2026-09-04属性另列对照。"],
        ["宝石批准贡献","17类×10品质共170条。02A为固定值，02B为属性/机制百分比名义与单颗有效值，02C为收益及显示边界。新版至宝固定宝石128/128/640仍计入01表。"],
        ["装备基础调整","原批准的生命/攻击/防御贡献约减半继续作为设计目标；不能与本轮分系词缀审查混为一次全局减半。"],
        ["掉落等级","07表左侧仍为批准的1～135等级带，右侧明确当前SourceItemLevel/100级上限差异。"],
        ["套装与存档","18条2/4/6件效果按当前源码完整表述；当前存档版本"+str(version)+"。"],
        ["经济与工具","已实装挂机金币×10/后续×1.25；分解基础价原方案×5仍是未实装选择；工具经验方案待确认。"],
        ["核验边界","本轮同步宝石品质曲线、产甲75%模型与详细属性C++；本表核验源码和工作簿一致性，UBT/行为验收单独记录，不能代替29条未接入词缀的实装。"],
        ["维护方式","使用python scripts/update_equipment_master_workbook.py；--check只验证。禁止用历史全量导出覆盖本表的批准设计基准。"],
    ]
    sheet(wb,"00_说明",["项目","内容"],intro,[25,137],"A2")

    # Re-style the design targets after the narrowly scoped gem contribution revision.
    for name in ["01_百级最终属性","06_至宝配装基准"]:
        original=snapshot[name]
        widths=[19,18,15,15,15,17,17,40,62,26,26,26] if name.startswith("01_") else [26,12,62,50,25]
        sheet(wb,name,original[0],original[1:],widths,"B2")
    order=["00_说明","04_词缀目录","04A_触发与接入审查","04B_品质与数值基线","04D_叠加上限与边际","04E_一至六件边际","04F_递减曲线","04G_详细属性口径","04C_名称显示入口","05_套装效果","05A_套装术语","08_实现差异","01_百级最终属性","02_宝石总表","02A_宝石品质边际","03_装备模板","06_至宝配装基准","07_掉落等级带","07A_挂机金币现状","07B_分解与工具经验","07C_经济与工具边界","09_来源校验"]
    for position,name in enumerate(order):
        wb.move_sheet(name,position-wb.sheetnames.index(name))
    for ws in wb:
        ws.sheet_properties.tabColor = "244D57"
        if ws.title in {"01_百级最终属性", "06_至宝配装基准"}:
            ws.sheet_properties.tabColor = "4B8B68"
        elif ws.title in {"04A_触发与接入审查", "04B_品质与数值基线", "04D_叠加上限与边际", "04E_一至六件边际", "04F_递减曲线", "07B_分解与工具经验", "07C_经济与工具边界"}:
            ws.sheet_properties.tabColor = "D99A33"
        elif ws.title == "08_实现差异":
            ws.sheet_properties.tabColor = "BA6650"
    wb.active=wb.sheetnames.index("04_词缀目录")
    wb.properties.title="GameXXK 装备设计总表（"+date+"代码核对）"
    wb.properties.subject="效果直述、随机词缀接入审查、批准设计与运行时差异"
    wb.properties.modified=datetime.now(timezone.utc).replace(tzinfo=None)
    assert rows_of(wb["01_百级最终属性"]) == snapshot["01_百级最终属性"]
    assert rows_of(wb["06_至宝配装基准"]) == snapshot["06_至宝配装基准"]
    assert [wb["02_宝石总表"].cell(r,7).value for r in range(2,len(gem_rows)+2)] == [r[6] for r in gem_rows]
    assert [r[:6] for r in rows_of(wb["07_掉落等级带"])] == original_drop_bands
    apply_travel_money_update(wb)
    apply_training_economy_update(wb)
    counts=validate(wb,src)
    assert sha(BOOK)==original_hash, "Workbook changed during the review; reload instead of overwriting."
    assert all(sha(ROOT/path)==digest for path,digest in source_hashes.items()), "Source changed during the update; rerun the audit."
    report_dir=ROOT/"Saved/Diagnostics/EquipmentMasterTableUpdate"/stamp.strftime("%Y%m%d-%H%M%S")
    report_dir.mkdir(parents=True,exist_ok=False)
    backup=report_dir/"equipment-master-before.xlsx"
    shutil.copy2(BOOK,backup)
    with tempfile.NamedTemporaryFile(suffix=".xlsx",prefix="equipment-master-",dir=BOOK.parent,delete=False) as handle:
        temporary=Path(handle.name)
    wb.save(temporary)
    wb.close()
    reloaded=openpyxl.load_workbook(temporary)
    validate(reloaded,src)
    reloaded.close()
    assert sha(BOOK)==original_hash, "Workbook changed before replacement."
    temporary.replace(BOOK)
    report={"ok":True,"workbook":str(BOOK),"backup":str(backup),"before_sha256":original_hash,"after_sha256":sha(BOOK),"branch":branch,"head":commit,"date":date,"source_hashes":source_hashes,"preserved_non_gem_design_budgets":True,"gem_quality_revision":"2026-09-09",**counts}
    (report_dir/"report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps(report,ensure_ascii=False))


if __name__ == "__main__":
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check",action="store_true",help="Validate the workbook against current code without writing files")
    update(parser.parse_args().check)
