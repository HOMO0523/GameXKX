"""Update only approved travel-money rules in the existing equipment master."""
from pathlib import Path
from copy import copy
import hashlib
import json
import re
import shutil
import openpyxl
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "docs/design/2026-09-04-project-design-tables/GameXXK_装备设计总表_2026-09-04.xlsx"
SHEET = "07D_宝箱与行旅钱"

def currency_constants():
    source = (ROOT / "Source/GameXXK/Public/GameXXKTravelMoneyRules.h").read_text(encoding="utf-8")
    return {name: int(re.search(r"\b" + name + r"\s*=\s*(\d+)", source).group(1))
            for name in ["ChestDropQuantity", "ShopBundleQuantity", "ShopBundleGoldPrice", "DismantleGoldPerUnit"]}

def apply_travel_money_update(wb):
    constants = currency_constants()
    price, quantity, drop, redemption = (constants[name] for name in
        ["ShopBundleGoldPrice", "ShopBundleQuantity", "ChestDropQuantity", "DismantleGoldPerUnit"])
    if SHEET in wb: wb.remove(wb[SHEET])
    ws = wb.create_sheet(SHEET, wb.sheetnames.index("07C_经济与工具边界") + 1)
    rows = [
        ["普通宝箱", "随机装备", 1, 0.50, None, "普通品质；等级按SourceItemLevel", "一次只出池中一类奖励"],
        ["普通宝箱", "随机宝石", 1, 0.30, None, "普通品质；当前17种属性中随机", "宝石详情由02总表另行维护"],
        ["普通宝箱材料池", "强化石", 1, "=20%/3", None, "材料池共20%，三种等概率", "非额外赠送"],
        ["普通宝箱材料池", "洗练砂", 1, "=20%/3", None, "材料池共20%，三种等概率", "非额外赠送"],
        ["普通宝箱材料池", "行旅钱", drop, "=20%/3", None, "材料池共20%，三种等概率", "总概率1/15，约6.67%"],
        ["局外商店", "行旅钱", quantity, None, price, "固定价格；购买进入背包", f"{quantity}个一份，连续购买逐份结算"],
        ["分解", "行旅钱", 1, None, redemption, f"每个固定返还{redemption}金币", "不吃金币天赋；不获得工具经验"],
        ["局内行商", "卡牌强化／遗物／刷新", None, None, None, "只消耗背包＋仓库的行旅钱", "背包优先，不足部分从仓库扣；不使用局外金币"],
        ["局内产出", "行旅钱", 0, 0, 0, "开局、战斗、精英、首领、事件、篝火、遗物均不产出", "旧虚拟计数不转换为实物"],
        ["局内卡牌强化", "目标品质普通／稀有／史诗", None, None, None, "对应25／40／60行旅钱；普通升稀有40，稀有升史诗60", "以已保存商品价格为准"],
        ["局内遗物", "普通／稀有／史诗", None, None, None, "沿用70／100／140行旅钱", "以已保存商品价格为准"],
        ["资产", "Item.TravelMoney", None, None, None, "/Game/GameXXK/UI/Items/Currency/T_Item_TravelMoney", "用户选定：前大后小两枚胖铜钱，左下短红绳"],
        ["批准口径", "2026-09-10最终反馈", None, None, None, "20%出材料，放进池子里平均出", "取代此前10%单独掉落方案"],
        ["验收记录", "路线地图与行旅钱", None, None, None, "docs/production/2026-09-09-route-map-travel-money-acceptance.md", "具体构建与行为结果见验收记录"],
    ]
    ws.append(["系统", "内容", "数量", "每箱概率", "金币金额", "规则", "备注"])
    for row in rows: ws.append(row)
    for col, width in zip("ABCDEFG", [23, 28, 10, 14, 16, 70, 60]): ws.column_dimensions[col].width = width
    for row in ws:
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=11, color="263F43")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.fill = PatternFill("solid", fgColor="F1F5F4" if cell.row % 2 == 0 else "FFFFFF")
        ws.row_dimensions[row[0].row].height = 48
    for cell in ws[1]:
        cell.font = Font(name="Microsoft YaHei", size=11, color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="244D57")
    for row in range(2, 7): ws.cell(row, 4).number_format = "0.00%"
    ws.freeze_panes = "C2"; ws.auto_filter.ref = ws.dimensions
    ws.sheet_properties.tabColor = "D99A33"
    icon = Image(str(ROOT / "SourceArt/UI/Items/Currency/T_Item_TravelMoney.png"))
    icon.width = icon.height = 128; ws.add_image(icon, "I2")
    ws["I1"] = "已选图标（透明PNG）"
    bounds = wb["07C_经济与工具边界"]
    for row in bounds:
        if row[0].value == "商店价格":
            row[1].value = f"装备包100000；高级箱100000；普通箱25000；宝石包100000；行旅钱{quantity}个{price}金币。"
            row[3].value = "行旅钱价格与普通箱材料池见07D；本项同步2026-09-10。"
    entries = {
        "行旅钱": ["实物Item.TravelMoney，背包与仓库数量合计供局内使用。", f"普通箱材料池20%三选一，行旅钱一次{drop}个；局外{quantity}个{price}金币。", f"每个分解{redemption}金币，不吃天赋、不加工具经验；详见07D。"],
        "局内商店货币": ["买卡、买遗物、刷新均只扣实物行旅钱。", "局内停止产出；不再用局外金币付款。", "余额不足及交易失败均不扣款、不发货。"],
    }
    for name, values in entries.items():
        row_index = next((r[0].row for r in bounds if r[0].value == name), bounds.max_row + 1)
        for col, value in enumerate([name] + values, 1):
            cell = bounds.cell(row_index, col, value)
            cell._style = copy(bounds.cell(2, col)._style)
        bounds.row_dimensions[row_index].height = 65
    bounds.auto_filter.ref = bounds.dimensions

def update():
    original = hashlib.sha256(BOOK.read_bytes()).hexdigest()
    wb = openpyxl.load_workbook(BOOK)
    unaffected = {s.title: list(s.values) for s in wb if s.title not in {SHEET, "07C_经济与工具边界"}}
    apply_travel_money_update(wb)
    folder = ROOT / "Saved/RouteNodeArt"; folder.mkdir(parents=True, exist_ok=True)
    backup = folder / "equipment-master-before-travel-money.xlsx"
    if not backup.exists(): shutil.copy2(BOOK, backup)
    temp = BOOK.with_name(BOOK.stem + ".travel-money.tmp.xlsx")
    wb.save(temp); wb.close()
    check = openpyxl.load_workbook(temp)
    assert all(list(check[name].values) == rows for name, rows in unaffected.items())
    constants = currency_constants()
    assert check[SHEET]["C6"].value == constants["ChestDropQuantity"] and check[SHEET]["E7"].value == constants["ShopBundleGoldPrice"] and check[SHEET]["E8"].value == constants["DismantleGoldPerUnit"]
    assert len(check[SHEET]._images) == 1
    assert hashlib.sha256(check[SHEET]._images[0]._data()).hexdigest() == hashlib.sha256((ROOT / "SourceArt/UI/Items/Currency/T_Item_TravelMoney.png").read_bytes()).hexdigest()
    check.close()
    assert hashlib.sha256(BOOK.read_bytes()).hexdigest() == original, "Workbook changed concurrently; reload it."
    temp.replace(BOOK)
    report = {"workbook": str(BOOK), "updated_sheets": [SHEET, "07C_经济与工具边界"], "unrelated_sheets_preserved": len(unaffected), "sha256": hashlib.sha256(BOOK.read_bytes()).hexdigest()}
    (folder / "travel-money-workbook-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False))

if __name__ == "__main__": update()
