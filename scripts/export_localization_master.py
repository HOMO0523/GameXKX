"""Export the live text catalogue/rules and sync names without replacing design numbers."""
import copy
import json
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from update_ui_goal_workbooks import prepare, finish

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs/design/2026-09-10-ui-localization'
RULES=json.loads((OUT/'rules.json').read_text(encoding='utf-8'))
ENTRIES=json.loads((ROOT/'Content/Localization/GameXXK/strings.json').read_text(encoding='utf-8'))['entries']
GEM_NAMES={e['key'].removeprefix('Gem.ShortName.'):e for e in ENTRIES if e['key'].startswith('Gem.ShortName.')}
SETS=RULES['equipmentSets'];SLOTS=RULES['equipmentSlots']
BY_SOURCE={e['zh-Hans']:e['en'] for e in ENTRIES if e.get('usage')!='compact'}

def sheet(book,name,headers,widths,rows,note):
    ws=book.create_sheet(name);prepare(ws,'GameXXK · '+name,note,headers,widths)
    for row in rows:ws.append(row)
    finish(ws);return ws

def save_verified(book,path):
    temp=path.with_name(path.stem+'.pending.xlsx');book.save(temp)
    loaded=load_workbook(temp,read_only=True,data_only=False)
    assert loaded.sheetnames==book.sheetnames
    counts={ws.title:ws.max_row for ws in loaded.worksheets};loaded.close();temp.replace(path)
    return counts

def append_column(ws,label,values,width=28):
    heads={str(c.value):c.column for c in ws[1] if c.value is not None}
    col=heads.get(label,ws.max_column+1)
    cell=ws.cell(1,col,label);cell.fill=PatternFill('solid',fgColor='244C45');cell.font=Font(name='Microsoft YaHei',bold=True,color='FFFFFF')
    ws.column_dimensions[get_column_letter(col)].width=width
    for row,value in values.items():
        c=ws.cell(row,col,value);c.alignment=Alignment(wrap_text=True,vertical='center');c.font=Font(name='Microsoft YaHei',size=10)
    if ws.auto_filter.ref:ws.auto_filter.ref=f'A1:{get_column_letter(ws.max_column)}{ws.max_row}'

def sync_design_names():
    results=[]
    for kind in ('装备','卡牌','怪物与阶段数值'):
        path=ROOT/f'docs/design/2026-09-04-project-design-tables/GameXXK_{kind}设计总表_2026-09-04.xlsx'
        wb=load_workbook(path);preserved={};gem_name_cells=set()
        if kind=='装备':
            gem_sheet=wb['02_宝石总表']
            for row in range(2,gem_sheet.max_row+1):
                parts=str(gem_sheet.cell(row,2).value).split('.')
                if len(parts)==4 and parts[:2]==['Item','Gem'] and parts[2] in GEM_NAMES:
                    gem_name_cells.add((gem_sheet.title,gem_sheet.cell(row,3).coordinate))
        # Numeric/design columns remain exactly as they were. Existing translation
        # columns are refreshed and excluded from the preservation comparison.
        for ws in wb:
            for row in ws:
                for cell in row:
                    if str(ws.cell(1,cell.column).value).startswith('英文'):continue
                    if (ws.title,cell.coordinate) in gem_name_cells:continue
                    preserved[(ws.title,cell.coordinate)]=cell.value
        if kind=='装备':
            ws=wb['03_装备模板'];names={};set_names={};slot_names={}
            for row in range(2,ws.max_row+1):
                parts=str(ws.cell(row,2).value).split('.')
                if len(parts)==3 and parts[0]=='Equipment' and parts[1] in SETS and parts[2] in SLOTS:
                    set_names[row]=SETS[parts[1]][1];slot_names[row]=SLOTS[parts[2]][1]
                    names[row]=set_names[row]+' '+slot_names[row]
                else:names[row]=BY_SOURCE.get(str(ws.cell(row,3).value),'待复核旧兼容名称')
            append_column(ws,'英文名称',names,28);append_column(ws,'英文套装',set_names);append_column(ws,'英文部位',slot_names)
            gem_english={}
            for row in range(2,gem_sheet.max_row+1):
                parts=str(gem_sheet.cell(row,2).value).split('.')
                if (gem_sheet.title,gem_sheet.cell(row,3).coordinate) in gem_name_cells:
                    gem_sheet.cell(row,3,GEM_NAMES[parts[2]]['zh-Hans'])
                    gem_english[row]=GEM_NAMES[parts[2]]['en']
            append_column(gem_sheet,'英文名称',gem_english,24)
        elif kind=='卡牌':
            ws=wb['01_卡牌总表'];heads={str(c.value):c.column for c in ws[1]}
            name_col=next((heads[h] for h in ('名称','卡牌名称','卡名') if h in heads),None)
            assert name_col,heads
            names={r:BY_SOURCE.get(str(ws.cell(r,name_col).value),'待复核') for r in range(2,ws.max_row+1)}
            append_column(ws,'英文名称',names,35)
        else:
            ws=next(s for s in wb if s.title.startswith('01_'));heads={str(c.value):c.column for c in ws[1]}
            name_col=next((heads[h] for h in ('名称','怪物名称','怪物') if h in heads),None)
            assert name_col,heads
            names={r:BY_SOURCE.get(str(ws.cell(r,name_col).value),'待复核') for r in range(2,ws.max_row+1)}
            append_column(ws,'英文名称',names,28)
        for (name,coordinate),value in preserved.items():assert wb[name][coordinate].value==value,(name,coordinate)
        # These are project design masters: keep existing sheets and all numbers.
        counts=save_verified(wb,path)
        check=load_workbook(path,read_only=False,data_only=False)
        for (name,coordinate),value in preserved.items():assert check[name][coordinate].value==value,(name,coordinate)
        check.close();results.append({'path':str(path),'names':len(names),'preservedCells':len(preserved),'sheets':counts})
    return results

def main():
    stamp=datetime.now().isoformat(timespec='seconds');wb=Workbook();wb.remove(wb.active)
    sheet(wb,'00_本地化规则',['规则ID','范围','规则','依据'],[12,24,115,32],RULES['rules'],stamp+' | '+RULES['status'])
    rows=[]
    for e in ENTRIES:
        rows.append([e.get('namespace','GameXXK'),e['key'],e['zh-Hans'],e['en'],e.get('usage','full'),e.get('nativeSource',''),e.get('nativePattern',''),'资源已登记；逐页实机状态见05'])
    sheet(wb,'01_全部文本',['命名空间','Key','中文','英文','用途','原生源','动态模板','验收边界'],[28,45,68,85,12,50,65,34],rows,'唯一文本资源：Content/Localization/GameXXK/strings.json；占位符、限定条件和Pill需保留。')
    sheet(wb,'02_简称与完整语义',['Key','中文短控件','英文短标签','英文完整名称','规则'],[40,28,24,48,65],
          [[e['key'],e['zh-Hans'],e['en'],BY_SOURCE.get(e['zh-Hans'],''),'仅用于短控件；完整Tooltip保留目的和机制'] for e in ENTRIES if e.get('usage')=='compact'],
          '优先清楚的短词，其次常用简称；不通过缩小到不可读字号来掩盖长度问题。')
    sheet(wb,'03_装备中英命名',['模板ID','中文名称','英文名称','中文套装','英文套装','中文部位','英文部位'],[45,30,30,20,24,20,22],
          [[f'Equipment.{sc}.{pc}',sv[0]+pv[0],sv[1]+' '+pv[1],sv[0],sv[1],pv[0],pv[1]] for sc,sv in SETS.items() for pc,pv in SLOTS.items()],
          '现代装备42个模板，英文严格两个单词；装备ID、数值、部位和套装效果不因命名而改变。')
    sheet(wb,'04_套装与词缀说明',['命名空间','Key','中文','英文'],[28,45,90,105],
          [[e.get('namespace','GameXXK'),e['key'],e['zh-Hans'],e['en']] for e in ENTRIES if e.get('namespace')=='EquipmentAffixUI' or e['key'].startswith('Equipment.Set.')],
          '18条套装效果 + 35个词缀标签；使用完整规则后渲染Pill，禁止中英后缀拼接。')
    pages=['工作台挂机条','背包属性','背包装备','背包卡组','伙伴选择','NPC选择','仓库','编队','天赋','工具分解','工具强化','工具洗炼','工具镶嵌','工具合成','历练地图与门票','商店','教程与界面引导','主线任务目录','剧情对白与选项','局内路线','事件营地行商','BattleBoard','卡牌普通ShiftCtrlTooltip','装备宝石Tooltip','通关失败讨伐结算','离线收益','设置','存档管理']
    qa=[]
    for page in pages:
        for language in RULES['languages']:
            for scale in (50,75,100):
                status='待逐页验证';evidence=''
                check=next((v for v in RULES.get('visualChecks',[]) if v['page']==page and v['language']==language and v['scale']==scale),None)
                if check:
                    status=check['status'];evidence=check['evidence']
                qa.append([page,language,scale,status,evidence])
    sheet(wb,'05_逐页检查',['界面','语言','HUD%','状态','截图/证据'],[30,18,12,55,110],qa,'通过范围仅按实际证据填写；用户截图仍显示多处漏译与出框，整体尚未完成。')
    ambiguous=defaultdict(list)
    for e in ENTRIES:
        if e.get('usage')!='compact':ambiguous[e['zh-Hans']].append(e)
    rows=[]
    for zh,group in ambiguous.items():
        if len({e['en'] for e in group})>1:
            rows.append([zh,' | '.join(sorted({e['en'] for e in group})),' | '.join(e.get('namespace','GameXXK')+'.'+e['key'] for e in group),'需要上下文Key；不可做全局字符串替换'])
    sheet(wb,'06_上下文译法审查',['中文源','不同英文','对应身份','规则'],[50,75,110,50],rows,'不同品质系统、动作和状态可能使用不同英文；通过明确身份区分。')
    sheet(wb,'07_未完成问题',['序号','问题'],[12,140],[[i+1,t] for i,t in enumerate(RULES['knownIssues'])],'此页列明未完成项，不能把词库行数当作实机完成度。')
    sheet(wb,'08_主线全文',['命名空间','Key','中文','英文'],[30,45,100,115],
          [[e.get('namespace','GameXXK'),e['key'],e['zh-Hans'],e['en']] for e in ENTRIES if e.get('namespace','').startswith('GameXXKMainStory')],
          '独立源MainStory/*.en.json；main-story.entries.json完成后由merge_localization_fragment.py统一合并。')
    sheet(wb,'09_敌人与招式',['Key','中文','英文'],[55,35,45],
          [[e['key'],e['zh-Hans'],e['en']] for e in ENTRIES if e['key'].startswith('Enemy.')],
          '对应怪物21种；普通、困难、地狱阶段招式保持ID与数值，仅优化显示名。')
    card_names=json.loads((ROOT/'Content/Localization/GameXXK/card-short-names.json').read_text(encoding='utf-8'))
    sheet(wb,'10_全角色卡牌短名',['Card ID','中文名','旧英文（仅对照）','当前英文名','单词数','字体预检宽度','UE字号30宽度'],[48,25,30,23,12,20,22],
          [[r['id'],r['zh-Hans'],r['previousEnglish'],r['en'],r['wordCount'],r['widthAt40px'],r.get('ueWidthAt30','待冷编译回归')] for r in card_names['cards']],
          '覆盖主角、所有伙伴、NPC及首领奖励，共173张；优先单词，最多双词。ID、品质、解锁和效果不变。')
    academy=json.loads((ROOT/'Content/Localization/GameXXK/academy-guide-copy.json').read_text(encoding='utf-8'))
    sheet(wb,'11_角色机制教学',['Course ID','课节','中文标题','英文标题','中文机制','英文机制'],[34,12,23,25,85,110],
          [[r['course'],r['index']+1,r['zhTitle'],r['enTitle'],r['zhMechanism'],r['enMechanism']] for r in academy['lessons']],
          '13角色/29节；提示展示机制+当前操作+实际目标进度，独立关卡教学与既有首通账本保持。验收进度见项目状态。')
    by_key={e['key']:e for e in ENTRIES}
    guide_source=(ROOT/'Source/GameXXK/Private/UI/GameXXKInterfaceHelpWidget.cpp').read_text(encoding='utf-8')
    steps=re.findall(r'Add\(TEXT\("([^"]+)"\),TEXT\("([^"]+)"\),TEXT\("([^"]+)"\),TEXT\("([^"]+)"\)([^;]*)\);',guide_source)
    sheet(wb,'12_界面入门操作',['步骤','进度ID','目标控件','准备界面','完成方式','中文指令','英文指令'],[10,44,46,24,24,65,85],
          [[i+1,'UI.Basics.V1.'+s[0],s[2],s[3],'悬停1秒后确认' if 'false,true' in s[4] else '阅读确认' if 'false' in s[4] else '真实点击',by_key['UI.Step.'+s[0]]['zh-Hans'],by_key['UI.Step.'+s[0]]['en']] for i,s in enumerate(steps)],
          '操作步骤不能靠下一步跳过；读档继续首个未完成步骤；不消费资产，不清教学奖励账本。')
    pacing_path=ROOT/'docs/design/2026-09-11-academy-pacing.json'
    if pacing_path.exists():
        pacing=json.loads(pacing_path.read_text(encoding='utf-8'))
        sheet(wb,'13_教学演示数值',['对象','最大HP','开场HP','ATK','DEF','实际回合/规则'],[42,15,15,15,15,76],
              [['主角 / Hero',300,210,70,35,'借用教学角色；治疗演示保留受伤状态'],
               ['伙伴 / Partner',300,210,70,35,'借用教学角色'],['NPC',300,210,70,35,'借用教学角色'],
               ['资源 / Resources · Enemy',pacing['heroEnemyHP'][0],pacing['heroEnemyHP'][0],12,0,'实际2回合'],
               ['防护 / Defense · Enemy',pacing['heroEnemyHP'][1],pacing['heroEnemyHP'][1],12,0,'实际3回合'],
               ['四式 / Spell task · Enemy',300,300,12,0,'实际1回合；完成四牌任务与重放'],
               ['其余教学 / Other courses · Enemy',300,300,12,0,'29节目标/胜利与最大HP检查全部通过']],
              '仅角色教学演示；所有角色最大HP300，敌人不超过300。见2026-09-11-academy-short-demo.md。')
    role_table=ROOT/'docs/design/2026-09-12-body-font-rollout/font-roles.tsv'
    role_rows=[]
    for raw in role_table.read_text(encoding='utf-8').splitlines():
        line=raw.rstrip()
        if not line.strip() or line.lstrip().startswith('#'):continue
        parts=[p.strip() for p in line.split('\t')]
        if len(parts)<3:raise SystemExit('bad role row: '+repr(raw))
        role_rows.append([parts[0],int(parts[1]),parts[2],parts[3] if len(parts)>3 else ''])
    role_counts={kind:sum(1 for r in role_rows if r[2]==kind) for kind in ('Title','Body','Manual')}
    check=RULES['rules'][1][2]
    sheet(wb,'14_文本排版角色',['调用文件','原行号','角色','理由'],[64,10,12,86],role_rows,
          'L02双字体：标题保留江湖古风体，其余文本用荆南圆体（KeinannMaruPOP，2026-09-12替换芝士奶盖乌龙宋）。'
          f"共{len(role_rows)}处调用点：Title {role_counts['Title']}、Body {role_counts['Body']}、"
          f"Manual {role_counts['Manual']}（由带角色参数的helper分派）。"
          'Manual行在源码中经TitleFont/BodyFont或Font(Role,…)分派，静态守卫见 '
          'scripts/gamexxk_font_roles_check.py。主界面大按钮（导航盘、教程/任务、挑战/游历）按用户要求归Title。'
          '规则原文：'+check)
    path=OUT/'GameXXK_本地化总表_2026-09-10.xlsx';counts=save_verified(wb,path)
    assert counts['01_全部文本']==len(ENTRIES)+4
    master=sync_design_names()
    report={'updatedAt':stamp,'path':str(path),'textEntries':len(ENTRIES),'sheets':counts,'designMasterSync':master,'overallStatus':'执行中，未完成'}
    (OUT/'export-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(report,ensure_ascii=True))

if __name__=='__main__':main()
