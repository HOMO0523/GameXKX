"""Refresh card/enemy masters from a fresh v40 runtime export; preserve approved stage budgets."""
from __future__ import annotations
import argparse, csv, hashlib, json, re, shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook
from update_equipment_master_workbook import sheet

ROOT=Path(__file__).resolve().parents[1]
TABLES=ROOT/'docs/design/2026-09-04-project-design-tables'
CARD=TABLES/'GameXXK_卡牌设计总表_2026-09-04.xlsx'
ENEMY=TABLES/'GameXXK_怪物与阶段数值设计总表_2026-09-04.xlsx'
RESISTANCE=ROOT/'docs/design/2026-09-09-resistance-discussion/innate-resistances-proposal.csv'
QUALITY={'Common':'普通','Rare':'稀有','Epic':'史诗'}
ORDER={'普通':1,'稀有':2,'史诗':3}
ROLES={'Hero':'通用','Blade':'刀客','Guard':'守卫','Healer':'医师','Hunter':'猎人','Sorcerer':'法师','FormationMaster':'阵师','Invalid':'通用'}
ELEMENT={'None':'物理','Fire':'火焰法术','Frost':'冰霜法术','Lightning':'雷击法术'}
TARGET={'None':'无目标','Self':'自身','LowestHealthParty':'生命最低角色','RandomLivingParty':'随机角色','AllLivingParty':'全体角色','AllEnemyAllies':'全体怪物','LowestHealthEnemyAlly':'生命最低怪物','MarkedParty':'带标记角色','PreyTarget':'锁定猎物','MarkedPartyElseRandom':'标记者优先，否则随机角色','HighestManaParty':'内力最高角色'}

def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def values(ws):return list(ws.iter_rows(values_only=True))
def dictionaries(ws):
    rows=values(ws);return [dict(zip(rows[0],row)) for row in rows[1:]]
def put(wb,name,headers,rows,widths=None):return sheet(wb,name,headers,rows,widths or [24]*len(headers),'C2')
def append_columns(wb,name,extra_headers,by_key,key_column):
    old=values(wb[name]);headers=list(old[0]);new_rows=[]
    for h in extra_headers:
        if h not in headers:headers.append(h)
    for original in old[1:]:
        row=dict(zip(old[0],original));row.update(by_key[row[key_column]])
        new_rows.append([row.get(h) for h in headers])
    put(wb,name,headers,new_rows,[20]*len(headers))

def status_names():
    source=(ROOT/'Source/GameXXK/Private/GameXXKCardText.cpp').read_text(encoding='utf-8')
    a=source.index('FString DescribeStatus(');b=source.index('\n\t}',a)
    return dict(re.findall(r'case EGameXXKCardStatus::(\w+): return TEXT\("([^"]+)"\)',source[a:b]))

def channels(row):
    d=row['definition'];effects=d.get('effects',[])+d.get('chargeEffects',[])+d.get('finishEffects',[])
    if row.get('inherits_previous_element'):return '承接前一张法师序牌；无对应元素时物理'
    out=[]
    for e in effects:
        kind=e['type'];status=e.get('status','None')
        if kind=='LightningPerTargetStatusSnapshot':out.append('雷击法术')
        elif kind=='DamageAllPercentAttackPerConsumedArmor':out.append('冰霜法术')
        elif kind in {'DamagePercentAttack','DamagePercentAttackPlusArmor','DamagePercentAttackPerTargetStatus','EachLivingAllyAttackSelectedTarget'}:
            label=row.get('damage_element','物理');out.append(label if label=='物理' else label+'法术')
        elif kind=='DamageFlat':out.append('固定伤害' if row.get('damage_element','物理')=='物理' else row['damage_element']+'法术')
        if status=='Burn':out.append('灼烧：火抗减免')
        elif status in {'Bleed','Poison','DamageOverTime'}:out.append('独立持续伤害')
        elif status in {'Counter','Block'}:out.append('条件物理反击')
    return '；'.join(dict.fromkeys(out)) or '无即时伤害／后续效果见说明'

def card_group(row,names,old_groups):
    d=row['definition'];npc=d.get('npcId')
    if npc and npc!='None':return '任务NPC·'+names.get(npc,npc)
    return old_groups.get(row['id']) or ('主角·'+ROLES.get(d.get('linkedRole'),'通用') if d['owner']=='Hero' else '伙伴·'+ROLES.get(d['role'],d['role']))

def update_cards(runtime,names):
    wb=load_workbook(CARD);old=values(wb['01_卡牌总表']);groups={r[2]:r[1] for r in old[1:]}
    grouped=defaultdict(list)
    for r in runtime['cards']:grouped[r['id']].append(r)
    ids=[r[2] for r in old[1:] if r[2] in grouped]+sorted(set(grouped)-{r[2] for r in old[1:]})
    masters=[];variants=[];pills=[]
    for index,ident in enumerate(ids,1):
        rows=sorted(grouped[ident],key=lambda r:ORDER[r['quality']]);first=rows[0];d=first['definition'];group=card_group(first,names,groups)
        masters.append([index,group,ident,first['name'],QUALITY[d['baseQuality']],'／'.join(r['quality'] for r in rows),first['target'],
            f"{d['energyCost']}气 · {d['manaCost']}内",first['compact'],first['detail'],first['pills'],'当前共享说明由代码生成','已对照运行时',
            '属性和抗性按今日规则；角色临时加成另算','源攻击100、防御100、队伍等级100的文本示例','每品质7地形执行；精确机制按专项验证',channels(first)])
        for r in rows:
            definition=r['definition'];variants.append([len(variants)+1,group,ident,r['name'],QUALITY[definition['baseQuality']],r['quality'],
                definition['energyCost'],definition['manaCost'],'本列为基础费用；当局减费/增费另算',f"{definition['energyCost']}气 · {definition['manaCost']}内",
                r['target'],r['compact'],r['detail'],r['pills'],'Ctrl查看当前品质术语','已对照运行时','今日物理/法术与抗性分类',
                f"{r['passed_terrains']}/7地形执行",channels(r)])
            pills.append(['当前品质',group,ident,r['name'],r['quality'],'','完整术语说明','',r['pills'],'由当前代码生成'])
    put(wb,'01_卡牌总表',list(old[0])[:16]+['今日伤害通道'],masters,[8,23,43,20,12,20,20,23,65,85,70,45,23,52,56,60,40])
    headers=list(values(wb['02_品质版本'])[0])[:18]+['今日伤害通道']
    put(wb,'02_品质版本',headers,variants,[8,23,43,20,12,12,10,10,40,22,25,70,85,80,40,23,48,40,42])
    branches=[]
    for r in runtime.get('branches',[]):
        branches.append([len(branches)+1,groups.get(r['id'],'伙伴·法师'),r['id'],r['name'],r['quality'],r['branch'],r['target'],r['compact'],r['detail'],r['pills'],'锁定分支，当前代码完整说明'])
    put(wb,'03_分支效果',['序号','分组','CardId','卡名','品质','分支','对象','简述','详述','Pill说明','共享注释'],branches,[8,24,44,22,14,13,25,75,95,85,55])
    put(wb,'04_Pill逐卡',['范围','分组','CardId','卡名','品质','分支','Pill','合并成员','短说明','共享注释'],pills,[20,25,44,22,13,15,26,25,100,65])
    common=[['物理伤害','伤害通道','逐包固定减防御，随后护甲吸收。','物理宝石不增幅法术'],['火/冰/雷法术','伤害通道','对应抗性减免，跳过防御；临时护甲仍吸收。','三抗独立，最高75%'],['抗性弱点','抗性','负抗性提高对应法术伤害；最低-25%。','-10%火抗承受110%火焰'],['万法归一','继承','按前一张法师序牌的火/冰/雷元素；无对应元素时物理，重放沿用快照。','不改变原倍率'],['吸取内力','资源效果','金钱鼠8点，巨蟾6点；均不随难度改变。','巨蟾命中后一次；不足扣至0']]
    previous=values(wb['05_通用术语']);replaced={r[0] for r in common}
    put(wb,'05_通用术语',list(previous[0]),[list(r) for r in previous[1:] if r[0] not in replaced]+common,[28,24,100,75])
    put(wb,'08_公式口径',['项目','公式','边界','示例/说明'],[
        ['物理伤害','max(1,D-有效防御)→护甲→生命','D先计入对应增伤；物理与法术独立','100伤害、30防御、20护甲：扣生命50'],
        ['法术伤害','D×(1-R/100)→护甲→生命','对应三抗，跳过防御','100伤害、20%抗性、20护甲：扣生命60'],
        ['抗性边际','R0+(75-R0)S/(75-R0+S)，减抗后夹到[-25,75]','S为名义抗性百分点，目前来自抗性宝石','基础10%＋宇宙86：最终47.02%'],
        ['伤害宝石边际','75S/(75+S)','同次真正适用的物理/反击或元素/持续收益合池一次','不把物理宝石叠到法术'],
        ['灼烧','来源火焰/持续增伤→目标火抗→生命','原始层数不变，抗性只计算一次','自然按施加者；引燃按触发者'],
        ['内力吸取','min(当前内力,8或6)','资源效果，不伪造生命伤害，不受伤害增幅/抗性影响','金钱鼠最高内力者；巨蟾实际受击者'],
    ],[25,80,85,75])
    put(wb,'09_覆盖核对',['项目','数量'],[['玩家卡',len(masters)],['品质版本',len(variants)],['分支展示',len(branches)],['当前品质术语',len(pills)],['7地形执行次数',runtime['executions']],['宝石类型',17],['存档版本',runtime['save_version']]],[30,65])
    put(wb,'00_说明',['项目','内容'],[['更新口径','今日编译模块实际导出；原文件名保留以维持链接'],['玩家卡与品质',f"{len(masters)}张，{len(variants)}个品质版本"],['新增规则','物理/三系法术、对应抗性、万法继承与8/6吸元已同步'],['数值示例','角色源攻击100、防御100、队伍等级100；示例不是任意敌人的最终扣血'],['品质术语','04页每个品质版本一条完整当前术语，替代旧拆行数量口径'],['验收边界','运行可执行与文本对照不代表任意配队平衡；精确机制与实机报告见生产记录'],['源文件','Saved/Automation/DesignTableRuntime/runtime.json'],['存档版本',runtime['save_version']]],[28,130])
    wb.save(CARD);return {'cards':len(masters),'variants':len(variants),'branches':len(branches)}

def effect_text(e,difficulties,status):
    def amount(field):
        data=e.get(field,{})
        return '／'.join(str(data.get(d,0)) for d in difficulties)
    kind=e['type'];target=TARGET.get(e.get('target'),e.get('target',''));condition='命中后：' if e.get('bRequiresPreviousDirectHit') else ''
    st=status.get(e.get('status'),e.get('status',''))
    if kind=='DirectDamage':
        text=f"{amount('attackPercentByDifficulty')}%攻击的{ELEMENT.get(e.get('damageElement','None'),'物理')}伤害×{e.get('hitCount',1)}"
        if e.get('status') not in ['None','Invalid',None]:text+=f"，附加{st} {amount('statusAmountByDifficulty')}"+('系数' if e['status'] in ['Burn','Bleed','Poison','DamageOverTime'] else '层')
    elif kind=='ApplyStatus':text=f"施加{st} {amount('statusAmountByDifficulty')}"+('系数（按等级换算）' if e['status'] in ['Burn','Bleed','Poison','DamageOverTime'] else '层')
    elif kind=='AddArmorDefensePercent':text=f"获得防御的{amount('defensePercentByDifficulty')}%护甲"
    elif kind=='HealMaxHealthPercent':text=f"恢复最大生命的{amount('resourceAmountByDifficulty')}%"
    elif kind=='DrainMana':text=f"吸取{amount('resourceAmountByDifficulty')}点当前内力，不足扣至0"
    elif kind=='QueueNextRoundEnergyPenalty':text=f"下回合共享气力减少{amount('resourceAmountByDifficulty')}"
    elif kind=='IncreaseNextCardEnergy':text=f"下一张牌气力消耗增加{amount('resourceAmountByDifficulty')}"
    elif kind in ['RemovePositiveStatus','RemoveNegativeStatus']:text=f"移除{amount('resourceAmountByDifficulty')}层"+('增益' if kind=='RemovePositiveStatus' else '负面')
    elif kind=='RefreshHealingAmplification':text=f"刷新下次恢复增幅{amount('resourceAmountByDifficulty')}%"
    elif kind=='TriggerDamageOverTime':text=f"触发{st}{e.get('hitCount',1)}次，保持原状态结算规则"
    elif kind=='ConsumeWealthForHealing':text=f"消耗至多{e.get('maxConsumedStacks',0)}层财富，每层恢复最大生命的{amount('resourceAmountByDifficulty')}%"
    elif kind=='ModifyAttack':
        text=f"增加来源攻击的{amount('attackPercentByDifficulty')}%" if any(e.get('attackPercentByDifficulty',{}).values()) else f"下一次直接攻击增加{amount('resourceAmountByDifficulty')}点"
    else:raise ValueError('Unhandled live enemy effect: '+kind)
    if e.get('sourceStatusForFlatMagnitude') not in ['None','Invalid',None]:text+=f"；每层{status.get(e['sourceStatusForFlatMagnitude'],e['sourceStatusForFlatMagnitude'])}另加{amount('resourceAmountByDifficulty')}点，上限{e.get('maxConsumedStacks',0)}层"
    if e.get('consumedStatus') not in ['None','Invalid',None] and kind!='ConsumeWealthForHealing':text+=f"；消耗最多{e.get('maxConsumedStacks',0)}层{status.get(e['consumedStatus'],e['consumedStatus'])}，每层另加{e.get('magnitudePerConsumedStack',0)}点"
    if e.get('bAssignsPersistentTarget'):text+='；锁定对应猎物'
    if e.get('bClearsPersistentTargetAfterResolve'):text+='；结算后解除锁定'
    return condition+target+'：'+text

def update_enemies(runtime,profiles):
    wb=load_workbook(ENEMY);status=status_names();enemies={e['id']:e for e in runtime['enemies']}
    updates={}
    for ident,e in enemies.items():
        r=profiles[ident];v={'名称':e['displayName'],'基础生命':e['baseHP'],'生命/级':round(e['hPPerLevel'],3),'基础攻击':e['baseAttack'],'攻击/级':round(e['attackPerLevel'],3),'基础防御':e['baseDefense'],'防御/级':round(e['defensePerLevel'],3),'速度':e['speed'],'火抗%':r['fire'],'冰抗%':r['frost'],'雷抗%':r['lightning']}
        for level in [100,125,135]:
            for cn,key in [('生命','maxHP'),('攻击','attack'),('防御','defense')]:v[f'L{level}{cn}']=e[f'level{level}'][key]
        updates[ident]=v
    append_columns(wb,'01_怪物属性',['火抗%','冰抗%','雷抗%'],updates,'怪物ID')
    old_notes={}
    for name in ['05_普通怪意图','06_阶段一意图','07_阶段二意图','08_阶段三意图']:
        for row in dictionaries(wb[name]):old_notes[(row['怪物'],row['意图'])]=next((v for k,v in row.items() if '批准效果' in k or k=='原条件与设计说明'),'')
    phase_rows=defaultdict(list);mana=[]
    for e in runtime['enemies']:
        for phase in e['phases']:
            number=phase['phaseNumber'];diffs=['normal','hard','hell'] if number==1 else ['hard','hell'] if number==2 else ['hell']
            for index,card in enumerate(phase['intents'],1):
                effects=card['effects'];text='\n'.join(effect_text(effect,diffs,status) for effect in effects)
                if card.get('chargeRounds',0):text+=f"\n蓄力{card['chargeRounds']}回合"
                if card.get('healingCooldownRounds',0):text+=f"\n治疗冷却{card['healingCooldownRounds']}回合"
                element='／'.join(dict.fromkeys(ELEMENT.get(x.get('damageElement','None'),'物理') for x in effects if x['type']=='DirectDamage')) or '资源／状态效果'
                key='normal' if e['tier']=='Normal' else str(number)
                phase_rows[key].append([e['displayName'],number,index,card['displayName'],text,element,old_notes.get((e['displayName'],card['displayName']),''),e['id'],card['id']])
                for effect in effects:
                    if effect['type']=='DrainMana':mana.append([e['displayName'],number,card['displayName'],*[effect['resourceAmountByDifficulty'][d] for d in ['normal','hard','hell']],TARGET[effect['target']],bool(effect.get('bRequiresPreviousDirectHit')),card['id']])
    for key,name in [('normal','05_普通怪意图'),('1','06_阶段一意图'),('2','07_阶段二意图'),('3','08_阶段三意图')]:
        put(wb,name,['怪物','阶段','顺序','意图','当前目录效果（普通／困难／地狱按本阶段范围）','伤害通道','原条件与设计说明','怪物ID','意图ID'],phase_rows[key],[22,10,10,28,105,28,100,43,45])
    put(wb,'12_抗性与吸元',['项目','当前规则'],[['抗性','各怪基础三抗见01；不随等级/难度普涨。火冰雷法术按对应抗性减伤，跳过防御，仍由护甲吸收。'],['抗性上限','正向加成边际递减，上限75%、下限-25%；负值是元素弱点。'],['元素特色','白猿投石冰霜、盘角鹿角击雷击、既有灼烧招式火焰。'],['吸取内力','金钱鼠8/8/8，巨蟾6/6/6；不足扣至0；当前内力影响后续出牌合法性。'],['既有编制','02～04的27关、189编制和125级地狱3-1基准保留；没有改关卡等级或原伤害倍率。'],['解释边界','当前目录效果为原始倍率/系数；目标状态条件与原设计说明并列；最终生命损失按护甲与抗性结算。']],[27,145])
    put(wb,'13_内力吸取清单',['怪物','阶段','招式','普通','困难','地狱','对象','需要前段命中','意图ID'],mana,[23,10,28,13,13,13,46,24,43])
    put(wb,'10_实现状态',['项目','当前实现','验证口径'],[['三抗配置','21只怪各有独立基础三抗','实际战斗快照与当前编译模块导出对照'],['元素伤害','物理/三系法术独立减伤，元素传递至提示和命中表现','专项数值/来源测试与2D实机'],['内力吸取','金钱鼠8点、巨蟾6点，三难度一致','最高内力、援护、闪避、不足和存档检查'],['阶段与编制','原阶段数、27关和189编制保留','既有351预演与阶段回归继续验证']],[28,95,100])
    wb.save(ENEMY);return {'enemies':len(enemies),'intent_rows':sum(map(len,phase_rows.values())),'mana_rows':len(mana)}

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--runtime',default='Saved/Automation/DesignTableRuntime/runtime.json');args=parser.parse_args()
    path=ROOT/args.runtime;runtime=json.loads(path.read_text(encoding='utf-8-sig'))
    assert runtime.get('schema_version')==2 and runtime.get('save_version')==40 and runtime.get('gem_type_count')==17,'Require the newly compiled, current runtime export'
    assert len(runtime['cards'])==419 and all(r['passed_terrains']==7 for r in runtime['cards'])
    assert len(runtime['enemies'])==21 and len(runtime['gems'])==170
    with RESISTANCE.open(encoding='utf-8-sig',newline='') as f:approved=list(csv.reader(f))[1:]
    profiles={r['key']:r for r in runtime['resistance_profiles']};assert len(profiles)==34
    for row in approved:
        assert [profiles[row[1]][x] for x in ['fire','frost','lightning']]==[float(x) for x in row[3:6]],row
    names={r[1]:r[2] for r in approved};backup=ROOT/'Saved/Diagnostics/DesignTablesToday'/datetime.now().strftime('%Y%m%d-%H%M%S');backup.mkdir(parents=True)
    for p in [CARD,ENEMY]:shutil.copy2(p,backup/p.name)
    old=load_workbook(ENEMY,read_only=True,data_only=True);preserved={n:values(old[n]) for n in ['02_27关总表','03_189编制','04_地狱3-1']};old.close()
    report={'cards':update_cards(runtime,names),'enemies':update_enemies(runtime,profiles),'runtime_sha256':sha(path),'backup':str(backup)}
    current=load_workbook(ENEMY,read_only=True,data_only=True)
    assert all(values(current[n])==rows for n,rows in preserved.items());current.close()
    for p in [CARD,ENEMY]:
        check=load_workbook(p,read_only=True,data_only=True);assert all(s.max_row>1 for s in check);check.close()
    report['files']={str(p.relative_to(ROOT)):sha(p) for p in [CARD,ENEMY]}
    (backup/'report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(report,ensure_ascii=False))

if __name__=='__main__':main()
