"""Add the missing draw-to-hand brief, clearly labelled as a proposed sound."""
import hashlib
import json
import shutil
import subprocess
import zipfile
from pathlib import Path
import numpy as np
from openpyxl import load_workbook
from essential_sfx_requirement_catalog import SPECS,folder
from build_essential_sfx_packages import ff,wav,sheets

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'Saved/Codex/CardDrawRequirement-20260915'
SOURCE=ROOT/'SourceArt/Audio/References/CardDraw'
SPEC=next(s for s in SPECS if s['cue']=='CardDraw')
PACKAGE=ROOT/'Deliverables/音效需求'/folder(SPEC)
NAME='SFX_CardDraw_ref_v01'


def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def write(path,value):path.write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf-8')


def reference():
    SOURCE.mkdir(parents=True,exist_ok=True)
    original=OUT/'bookFlip1.ogg'
    # Decode to floating point before gain adjustment to avoid clipping a
    # compressed source whose reconstructed samples can exceed full scale.
    raw=ff(['-i',original,'-ar','48000','-ac','1','-f','f32le','-'])
    samples=np.frombuffer(raw,'<f4').astype(np.float64)
    active=np.flatnonzero(np.abs(samples)>.08*np.abs(samples).max())
    begin=max(0,int(active[0])-96);part=samples[begin:begin+7680].copy()
    assert len(part)==7680
    part[:96]*=np.linspace(0,1,96);part[-960:]*=np.linspace(1,0,960)
    gain=.36/float(np.abs(part).max());part*=gain
    dest=SOURCE/(NAME+'.wav');wav(dest,part[:,None]);shutil.copy2(original,SOURCE/'bookFlip1.ogg')
    record={'cue':'CardDraw','variant':1,'name':NAME,'author':'Kenney','pack':'kenney-rpg',
            'source':'https://kenney.nl/assets/rpg-audio','license':'CC0-1.0',
            'license_url':'https://creativecommons.org/publicdomain/zero/1.0/',
            'file':'参考音频/'+NAME+'.wav','source_file':'来源与授权/kenney-rpg/bookFlip1.ogg',
            'project_file':dest.relative_to(ROOT).as_posix(),'project_source_file':(SOURCE/'bookFlip1.ogg').relative_to(ROOT).as_posix(),
            'sha256':sha(dest),'source_sha256':sha(original),'seconds':.16,'sample_rate':48000,'channels':1,'pcm_bits':16,
            'primary':True,'used_in_video':True,'reference_only':True,'role':'新增抽牌参考配音；尚未接入游戏',
            'processing':{'source_seconds':len(samples)/48000,'excerpt_start_seconds':begin/48000,'seconds':.16,'fade_in_ms':2,'fade_out_ms':20,'peak':.36,'gain':gain},
            'source_archive_sha256':sha(OUT/'kenney_rpg-audio.zip')}
    write(SOURCE/'manifest.json',record)
    (PACKAGE/'参考音频').mkdir(parents=True,exist_ok=True)
    (PACKAGE/'来源与授权/kenney-rpg').mkdir(parents=True,exist_ok=True)
    shutil.copy2(dest,PACKAGE/record['file']);shutil.copy2(original,PACKAGE/record['source_file'])
    for p in (ROOT/'SourceArt/Audio/Essential/Licenses/kenney-rpg').iterdir():
        if p.is_file():shutil.copy2(p,PACKAGE/'来源与授权/kenney-rpg'/p.name)
    write(PACKAGE/'来源与授权/音频清单.json',[record])
    return record,part


def video(part):
    source=ROOT/'Saved/Codex/EssentialSfxRequirements-20260914/defeat02-raw.mp4'
    frames=[397,762];sections=[];events=[];clips=[]
    for index,frame in enumerate(frames):
        start=frame-60;target=OUT/f'clip-{index}.mp4'
        ff(['-i',source,'-vf',f'trim=start_frame={start}:end_frame={start+150},setpts=PTS-STARTPTS','-an','-c:v','libx264','-crf','18','-preset','fast','-pix_fmt','yuv420p','-r','60','-fps_mode','cfr',target])
        clips.append(target)
        offset=index*2.5
        sections.append({'name':f'round-draw-{index+1}','video':str(source),'offset':offset,'duration':2.5,'trim_frame':start,'raw_first_visible_frame':frame})
        events.append({'cue':'CardDraw','file':NAME+'.wav','seconds':offset+1,'frame_contains_start':index*150+60,'target_frame':index*150+60,
                       'primary':True,'section':index,'label':f'第{index+1}次回合补牌 · 整批入手',
                       'reference_only':True,'timing_basis':'first visible incoming-card frame, verified against adjacent frames'})
    (OUT/'concat.txt').write_text('\n'.join(f"file '{p.name}'" for p in clips),encoding='utf-8')
    ff(['-f','concat','-safe','0','-i',OUT/'concat.txt','-c','copy',OUT/'picture.mp4'])
    mixed=np.zeros((5*48000,2),dtype=np.float64)
    for event in events:
        start=round(event['seconds']*48000);mixed[start:start+len(part)]+=part[:,None]*.282
    wav(OUT/'reference-mix.wav',mixed)
    (OUT/'caption.ass').write_text('''[Script Info]
ScriptType: v4.00+
PlayResX: 1280
PlayResY: 720
WrapStyle: 2
[V4+ Styles]
Format: Name, Fontname, Fontsize, PrimaryColour, SecondaryColour, OutlineColour, BackColour, Bold, Italic, Underline, StrikeOut, ScaleX, ScaleY, Spacing, Angle, BorderStyle, Outline, Shadow, Alignment, MarginL, MarginR, MarginV, Encoding
Style: Title,Microsoft YaHei,28,&H00F1F3EC,&H00000000,&H00000000,&H00000000,-1,0,0,0,100,100,0,0,1,0,0,7,0,0,0,1
Style: Body,Microsoft YaHei,18,&H00D0D8D1,&H00000000,&H00000000,&H00000000,0,0,0,0,100,100,0,0,1,0,0,7,0,0,0,1
Style: Cue,Microsoft YaHei,20,&H008FD8B3,&H00000000,&H00000000,&H00000000,0,0,0,0,100,100,0,0,1,0,0,7,0,0,0,1
[Events]
Format: Layer, Start, End, Style, Name, MarginL, MarginR, MarginV, Effect, Text
Dialogue: 0,0:00:00.00,0:00:05.00,Title,,0,0,0,,{\\pos(28,19)}15 / 抽牌 · 音效需求
Dialogue: 0,0:00:00.00,0:00:05.00,Body,,0,0,0,,{\\pos(820,26)}参考配音 · 尚未接入游戏
Dialogue: 0,0:00:00.00,0:00:02.50,Body,,0,0,0,,{\\pos(28,651)}第一次回合补牌｜卡牌进入手牌；同批多张合并播放一次。
Dialogue: 0,0:00:02.50,0:00:05.00,Body,,0,0,0,,{\\pos(28,651)}第二次回合补牌｜复用同一条轻短抽牌音，其他战斗声静音。
Dialogue: 0,0:00:01.00,0:00:01.65,Cue,,0,0,0,,{\\pos(780,689)}入手起点 1.000 s / F0060
Dialogue: 0,0:00:03.50,0:00:04.15,Cue,,0,0,0,,{\\pos(780,689)}入手起点 3.500 s / F0210
''',encoding='utf-8-sig')
    shutil.copy2('C:/Windows/Fonts/consola.ttf',OUT/'counter.ttf')
    vf=r"subtitles=caption.ass,drawtext=fontfile=counter.ttf:fontsize=18:fontcolor=0xaebfb5:x=28:y=691:text='F%{eif\:n\:d\:4}  /  60 fps'"
    ff(['-i','picture.mp4','-i','reference-mix.wav','-map','0:v:0','-map','1:a:0','-vf',vf,'-c:v','libx264','-crf','18','-preset','fast','-threads','4','-pix_fmt','yuv420p','-c:a','aac','-b:a','192k','-t','5','-movflags','+faststart',PACKAGE/'01_演示_参考配音.mp4'],cwd=OUT)
    ff(['-i',PACKAGE/'01_演示_参考配音.mp4','-map','0:v:0','-an','-c:v','copy','-movflags','+faststart',PACKAGE/'02_演示_无声版.mp4'])
    plan={'cue':'CardDraw','title':'抽牌','seconds':5,'frame_count':300,'reference_only':True,'sections':sections,'events':events}
    write(ROOT/'Saved/Codex/EssentialSfxRequirements-20260914/CardDraw-delivery-plan.json',plan)
    write(OUT/'delivery-plan.json',plan)
    return plan


def build():
    OUT.mkdir(parents=True,exist_ok=True);PACKAGE.mkdir(parents=True,exist_ok=True)
    record,part=reference();plan=video(part)
    sheets(SPEC,plan,[record],PACKAGE)
    workbook=load_workbook(PACKAGE/'03_时间轴需求.xlsx')
    sheet=workbook['时间轴需求'];sheet.cell(1,3).value='参考配音落点秒';sheet.cell(1,4).value='画面入手帧'
    for row in sheet.iter_rows(min_row=2):row[6].value='新增参考配音，待接入'
    overview=workbook['制作需求']
    for row in overview.iter_rows():
        if row[0].value=='时间含义':row[1].value='画面原始帧397/762出现新卡，剪接后对应F0060/F0210。声音按这两个画面落点添加参考配音，并非游戏当前抽牌实录。'
        if row[0].value=='状态':row[1].value='新增抽牌需求；当前没有独立CardDraw触发，参考声音尚未导入游戏。'
        if row[0].value=='数量':row[1].value='1条共用音效。单张一次，同批多张合并一次，减少密集抽牌噪音。'
    workbook.save(PACKAGE/'03_时间轴需求.xlsx')
    (PACKAGE/'04_时间轴速览.txt').write_text('''抽牌音效需求 v001 — 新增，待接入
视频：5秒 / 60fps / F0000–F0299
声轨为参考配音，其他战斗声静音。

第1段：[0.000, 2.500)秒 / [F0000, F0150)
第一批卡牌进入手牌：1.000秒 / F0060
参考音频：SFX_CardDraw_ref_v01.wav，播放一次。

第2段：[2.500, 5.000)秒 / [F0150, F0300)
第二批卡牌进入手牌：3.500秒 / F0210
参考音频：SFX_CardDraw_ref_v01.wav，复用一次。

要求：轻短翻纸、滑入；单张一次，同批多张合并一次。
手牌重排、悬停、无实际入手不响。抽牌和出牌提交分开。
''',encoding='utf-8-sig')
    (PACKAGE/'00_先看这里.md').write_text('''# 抽牌 · 新增音效需求

打开`00_音效需求.html`查看、试听和导入新WAV试配。此项**尚未接入游戏**，视频声轨为CC0素材制作的**参考配音**；其他战斗声音静音。

抽牌是新卡进入手牌，区别于“出牌”卡牌提交飞出。建议轻短的翻纸/滑入感，0.08–0.20秒。单张抽牌一次，同批多张合并播放一次，避免密集连响。起手、回合补牌、技能抽牌可复用；本视频演示两次回合批量补牌。

视频画面复用已录到的真实战斗，分别在原始F0397、F0762首次出现新入手卡牌，已与前后帧核对。每段截取150帧，最终总计300帧/5秒/60fps，入手点为F0060（1.000秒）、F0210（3.500秒）。此为成片帧索引，不是引擎内部帧；原始录像可能包含重复帧。

参考声音`SFX_CardDraw_ref_v01.wav`来自Kenney RPG Audio的`bookFlip1.ogg`，截取0.16秒并调整增益、淡入淡出。原文件、授权原文和来源记录随包提供。[作者发布页](https://kenney.nl/assets/rpg-audio)列明CC0；音频授权不覆盖本项目画面或以后导入的成品。

参考WAV为48kHz/单声道/16-bit，由OGG素材处理，不是24-bit录音母带。制作交付建议48kHz/单声道/24-bit PCM，听审后再接入游戏。网页导入和导出不会修改游戏声音或存档。
''',encoding='utf-8')
    streams=[ff(['-i',p,'-map','0:v:0','-c','copy','-bsf:v','h264_mp4toannexb','-f','h264','-']) for p in [PACKAGE/'01_演示_参考配音.mp4',PACKAGE/'02_演示_无声版.mp4']]
    assert hashlib.sha256(streams[0]).digest()==hashlib.sha256(streams[1]).digest()
    frames=[line for line in ff(['-i',PACKAGE/'01_演示_参考配音.mp4','-map','0:v:0','-f','framemd5','-']).decode().splitlines() if line and not line.startswith('#')]
    assert len(frames)==300
    manifest={'package':folder(SPEC),'primaryCue':'CardDraw','reference_only':True,'integration_status':'not_integrated',
              'video':{'width':1280,'height':720,'fps':60,'frame_count':300,'seconds':5,'same_encoded_video_stream':True,'soundtrack':'reference_dub'},
              'sound_events':plan['events'],'files':[{'file':p.relative_to(PACKAGE).as_posix(),'bytes':p.stat().st_size,'sha256':sha(p)} for p in sorted(PACKAGE.rglob('*')) if p.is_file() and p.name!='05_文件核验.json']}
    write(PACKAGE/'05_文件核验.json',manifest)
    print(json.dumps({'package':str(PACKAGE),'frames':300,'draw_markers':[60,210],'status':'reference_dub_not_integrated'},ensure_ascii=False))


if __name__=='__main__':build()
