"""Create the remaining 13 timecoded audio briefs from verified local recordings."""
import argparse
import hashlib
import json
import math
import re
import shutil
import subprocess
import textwrap
import wave
import zipfile
from pathlib import Path
import numpy as np
import imageio_ffmpeg
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from essential_sfx_requirement_catalog import SPECS, folder

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'Saved/Codex/EssentialSfxRequirements-20260914'
DEST=ROOT/'Deliverables/音效需求'
LIB=ROOT/'SourceArt/Audio/Essential'
FFMPEG=imageio_ffmpeg.get_ffmpeg_exe()
MANIFEST=json.loads((LIB/'manifest.json').read_text(encoding='utf-8'))
RECORDS={r['name']+'.wav':r for r in MANIFEST['files']}


def load(path):return json.loads(path.read_text(encoding='utf-8-sig'))
def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def write(path,value):path.write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf-8')
def ff(args,cwd=None):
    r=subprocess.run([FFMPEG,'-hide_banner','-loglevel','error','-y',*map(str,args)],cwd=cwd,capture_output=True)
    if r.returncode:raise RuntimeError(r.stderr.decode('utf-8','replace')[-3000:])
    return r.stdout
def audio(path):
    with wave.open(str(path)) as w:
        assert w.getsampwidth()==2 and w.getframerate()==48000
        return np.frombuffer(w.readframes(w.getnframes()),'<i2').reshape(-1,w.getnchannels()).astype(np.float64)/32768
def wav(path,data):
    with wave.open(str(path),'wb') as w:
        w.setparams((data.shape[1],2,48000,0,'NONE','not compressed'));w.writeframes(np.clip(np.rint(data*32768),-32768,32767).astype('<i2').tobytes())


def match(data, reference, hint=None, window=.8):
    x=data.mean(axis=1)[::4];y=reference.mean(axis=1)[::4]
    lower=max(0,int((hint-window)*12000)) if hint is not None else 0
    upper=min(len(x),int((hint+window)*12000)+len(y)) if hint is not None else len(x)
    x=x[lower:upper];assert len(x)>=len(y)
    n=1<<(len(x)+len(y)-1).bit_length()
    corr=np.fft.irfft(np.fft.rfft(x,n)*np.fft.rfft(y[::-1],n),n)[len(y)-1:len(x)]
    c=np.r_[0.,np.cumsum(x*x)];energy=np.maximum(0,c[len(y):]-c[:-len(y)])
    denom=np.sqrt(energy*np.sum(y*y));scores=np.divide(corr,denom,out=np.zeros_like(corr),where=denom>.0001)
    at=int(scores.argmax())
    return (lower+at)/12000,float(scores[at])


def separate_matches(data,events):
    """Fit overlapping known waves jointly before checking each individual cue."""
    mono=data.mean(axis=1)[::4]
    refs=[audio(ROOT/RECORDS[e['file']]['file']).mean(axis=1)[::4] for e in events]
    for _ in range(3):
        columns=[]
        for e,ref in zip(events,refs):
            at=round(e['audio_seconds']*12000);column=np.zeros(len(mono));size=min(len(ref),len(mono)-at)
            column[at:at+size]=ref[:size];columns.append(column)
        design=np.stack(columns,axis=1);gains=np.linalg.lstsq(design,mono,rcond=None)[0]
        total=design@gains
        for i,(e,ref) in enumerate(zip(events,refs)):
            isolated=mono-total+design[:,i]*gains[i]
            # match() downsamples by four, so repeat the low-rate signal and
            # reference to reuse its normalized local correlation consistently.
            at,corr=match(np.repeat(isolated[:,None],4,axis=0),np.repeat(ref[:,None],4,axis=0),e['audio_seconds'],.08)
            e['audio_seconds']=at;e['correlation']=corr;e['fitted_gain']=float(gains[i])
    return events


def recording(name):
    cached=OUT/(name+'-wave-analysis.json')
    if cached.exists() and load(cached).get('analysis_version')==2:return load(cached)
    report=load(OUT/(name+'-capture.json'));start=load(OUT/(name+'-start.json'));stop=load(OUT/(name+'-stop.json'))
    data=audio(OUT/(name+'.wav'));events=[];counts={}
    for observed in stop['events']:
        for cue,count in observed['audio']['played'].items():
            before=counts.get(cue,0)
            variants=[r for r in MANIFEST['files'] if r['cue']==cue]
            for index in range(before,count):
                r=variants[index%len(variants)]
                at,correlation=match(data,audio(ROOT/r['file']),observed['wall']-start['wall'])
                events.append({'cue':cue,'file':r['name']+'.wav','audio_seconds':at,
                               'raw_seconds':start['wall']-report['video_epoch']+at,'correlation':correlation,'mixed_correlation':correlation,
                               'cue_ordinal':index+1})
            counts[cue]=count
    assert counts==stop['state']['audio']['played'],name
    events=separate_matches(data,events)
    for e in events:
        e['raw_seconds']=start['wall']-report['video_epoch']+e['audio_seconds']
        assert e['correlation']>.85,(name,e)
    result={'analysis_version':2,'name':name,'video':str(OUT/(name+'-raw.mp4')),'audio':str(OUT/(name+'.wav')),
            'video_epoch':report['video_epoch'],'audio_epoch':start['wall'],'events':sorted(events,key=lambda e:e['raw_seconds']),
            'capture_api':report['capture_api'],'layout':report.get('layout','battle')}
    write(cached,result);return result


def button_recording(name):
    old=ROOT/'Saved/Codex/ToolSfxRequirementSample-20260914'
    start=load(old/(name+'-audio-start.json'))['wall_before']
    epoch=float(re.search(r'start: ([0-9.]+)',(old/(name+'-ffmpeg.log')).read_text(encoding='utf-8'))[1])
    data=audio(old/(name+'.wav'));at,corr=match(data,audio(LIB/'Waves/SFX_Button_v01.wav'))
    assert corr>.98
    return {'name':name,'video':str(old/(name+'-raw.mp4')),'audio':str(old/(name+'.wav')),'video_epoch':epoch,
            'audio_epoch':start,'events':[{'cue':'Button','file':'SFX_Button_v01.wav','audio_seconds':at,'raw_seconds':start-epoch+at,'correlation':corr,'cue_ordinal':1}],
            'capture_api':'verified 2026-09-14 mouse recording reused','layout':'side'}


def plan(spec):
    recordings=[button_recording(n) for n in ['reforge_generate01','reforge_keep01']] if spec['cue']=='Button' else [recording(spec['clip'])]
    sections=[];events=[];offset=0
    for capture in recordings:
        target=[e for e in capture['events'] if e['cue']==spec['cue']];assert target,spec['cue']
        begin=max(0,math.floor((min(e['raw_seconds'] for e in target)-spec.get('pre',2))*60))
        duration=5 if spec['cue']=='Button' else max(6,math.ceil(max(e['raw_seconds'] for e in target)-begin/60+2.4))
        meta=imageio_ffmpeg.read_frames(capture['video']);video_meta=next(meta);meta.close()
        assert video_meta['duration']>=begin/60+duration-.04,(capture['name'],video_meta['duration'],begin,duration)
        section={**capture,'trim_frame':begin,'duration':duration,'offset':offset}
        sections.append(section)
        for event in capture['events']:
            t=event['raw_seconds']-begin/60+offset
            if 0<=t<offset+duration and t>=offset:
                events.append({**event,'seconds':t,'frame_contains_start':math.floor(t*60),'target_frame':round(t*60),
                               'primary':event['cue']==spec['cue'],'section':len(sections)-1})
        offset+=duration
    return {'cue':spec['cue'],'title':spec['title'],'seconds':offset,'frame_count':offset*60,'sections':sections,'events':events}


def ass_time(s):
    c=round(s*100);return f'{c//360000}:{c//6000%60:02}:{c//100%60:02}.{c%100:02}'


def caption_file(spec,p,work):
    lines=['[Script Info]','ScriptType: v4.00+','PlayResX: 1280','PlayResY: 720','WrapStyle: 2','','[V4+ Styles]',
           'Format: Name, Fontname, Fontsize, PrimaryColour, SecondaryColour, OutlineColour, BackColour, Bold, Italic, Underline, StrikeOut, ScaleX, ScaleY, Spacing, Angle, BorderStyle, Outline, Shadow, Alignment, MarginL, MarginR, MarginV, Encoding',
           'Style: Title,Microsoft YaHei,28,&H00F1F3EC,&H00000000,&H00000000,&H00000000,-1,0,0,0,100,100,0,0,1,0,0,7,0,0,0,1',
           'Style: Body,Microsoft YaHei,18,&H00D0D8D1,&H00000000,&H00000000,&H00000000,0,0,0,0,100,100,0,0,1,0,0,7,0,0,0,1',
           'Style: Accent,Microsoft YaHei,20,&H009AE4B9,&H00000000,&H00000000,&H00000000,0,0,0,0,100,100,0,0,1,0,0,7,0,0,0,1',
           '','[Events]','Format: Layer, Start, End, Style, Name, MarginL, MarginR, MarginV, Effect, Text']
    def text(style,x,y,label,start=0,end=None):lines.append(f'Dialogue: 0,{ass_time(start)},{ass_time(p["seconds"] if end is None else end)},{style},,0,0,0,,{{\\pos({x},{y})}}{label}')
    def wrap_chinese(value):
        lines=[]
        for line in textwrap.wrap(value,21):
            while line and line[0] in '，。；：！？、）》' and lines:
                lines[-1]+=line[0];line=line[1:]
            if line:lines.append(line)
        return '\\N'.join(lines)
    if spec.get('layout')=='side':
        text('Body',48,35,f'SFX {spec["number"]:02}  /  音效制作需求 v001')
        text('Title',48,82,spec['title'])
        for i,section in enumerate(p['sections']):
            title=['按下洗炼按钮','选择原属性'][i] if spec['cue']=='Button' else spec['scene']
            text('Accent',48,238,title,section['offset'],section['offset']+section['duration'])
        text('Body',48,302,wrap_chinese(spec['brief']))
        text('Body',48,423,wrap_chinese(spec['rule']))
        text('Body',48,580,f'同类共用 · {len([e for e in p["events"] if e["primary"]])} 个参考落点')
    else:
        text('Title',28,19,f'{spec["number"]:02} / {spec["title"]} · 音效需求')
        text('Body',800,26,'v001  ·  60 fps  ·  当前游戏实录')
        text('Body',28,651,spec['scene']+'｜'+spec['rule'])
    for event in p['events']:
        if event['primary']:
            y=528 if spec.get('layout')=='side' else 689
            x=48 if spec.get('layout')=='side' else 740
            text('Accent',x,y,f'本类起音 {event["seconds"]:.3f} s / F{event["target_frame"]:04}',max(0,event['seconds']-.05),min(p['seconds'],event['seconds']+.65))
    (work/'caption.ass').write_text('\n'.join(lines),encoding='utf-8-sig')


def media(spec,p,package):
    work=OUT/('build-'+spec['cue']);work.mkdir(exist_ok=True)
    cuts=[];sounds=[]
    for i,section in enumerate(p['sections']):
        target=work/f'part-{i}.mp4';begin=section['trim_frame'];duration=section['duration']
        ff(['-i',section['video'],'-vf',f'trim=start_frame={begin}:end_frame={begin+duration*60},setpts=PTS-STARTPTS',
            '-an','-c:v','libx264','-preset','fast','-crf','18','-threads','4','-pix_fmt','yuv420p','-r','60','-fps_mode','cfr',target])
        cuts.append(target)
        data=audio(Path(section['audio']));at=round((begin/60+section['video_epoch']-section['audio_epoch'])*48000)
        leading=max(0,-at);assert leading<48000
        part=data[max(0,at):max(0,at+duration*48000)]
        if leading:
            assert np.max(np.abs(data[:4800]))<.003
            part=np.pad(part,((leading,0),(0,0)))
        if len(part)<duration*48000:
            assert duration*48000-len(part)<24000 and np.max(np.abs(part[-4800:]))<.003
            part=np.pad(part,((0,duration*48000-len(part)),(0,0)))
        sounds.append(part)
    wav(work/'mix.wav',np.concatenate(sounds))
    (work/'concat.txt').write_text('\n'.join(f"file '{f.name}'" for f in cuts),encoding='utf-8')
    ff(['-f','concat','-safe','0','-i',work/'concat.txt','-c','copy',work/'joined.mp4'])
    caption_file(spec,p,work);shutil.copy2('C:/Windows/Fonts/consola.ttf',work/'counter.ttf')
    vf=r"subtitles=caption.ass,drawtext=fontfile=counter.ttf:fontsize=18:fontcolor=0xaebfb5:x=28:y=691:text='F%{eif\:n\:d\:4}  /  60 fps'"
    ff(['-i','joined.mp4','-i','mix.wav','-map','0:v:0','-map','1:a:0','-vf',vf,'-c:v','libx264','-preset','fast','-crf','18','-threads','4','-pix_fmt','yuv420p','-c:a','aac','-b:a','192k','-t',p['seconds'],'-r','60','-fps_mode','cfr','-movflags','+faststart',package/'01_演示_现有效果.mp4'],cwd=work)
    ff(['-i',package/'01_演示_现有效果.mp4','-map','0:v:0','-an','-c:v','copy','-movflags','+faststart',package/'02_演示_无声版.mp4'])
    ff(['-i',package/'01_演示_现有效果.mp4','-frames:v','1',work/'preview.png'])


def sources(spec,p,package):
    used={e['file'] for e in p['events']};names=used|{name for name,r in RECORDS.items() if r['cue']==spec['cue']}
    records=[]
    for name in sorted(names):
        r=dict(RECORDS[name]);wave_file=ROOT/r['file'];original=ROOT/r['source_file']
        assert sha(wave_file)==r['sha256'] and sha(original)==r['source_sha256']
        dest=package/'参考音频'/name;dest.parent.mkdir(exist_ok=True);shutil.copy2(wave_file,dest)
        pack=package/'来源与授权'/r['pack'];pack.mkdir(parents=True,exist_ok=True);shutil.copy2(original,pack/original.name)
        for source in (LIB/'Licenses'/r['pack']).iterdir():
            if source.is_file():shutil.copy2(source,pack/source.name)
        r.update(project_file=r['file'],project_source_file=r['source_file'],file=dest.relative_to(package).as_posix(),source_file=(pack/original.name).relative_to(package).as_posix(),
                 primary=r['cue']==spec['cue'],used_in_video=name in used,role=('本类主音参考' if r['cue']==spec['cue'] else '随片其他声音')+(' · 实录使用' if name in used else ' · 同类备选变体'))
        records.append(r)
    write(package/'来源与授权/音频清单.json',records);return records


def sheets(spec,p,records,package):
    book=Workbook();overview=book.active;overview.title='制作需求'
    for row in [('音效需求',spec['title']),('制作方向',spec['brief']),('建议时长',f"{spec['duration'][0]}–{spec['duration'][1]} 秒"),('触发条件',spec['rule']),('保持安静',spec['silent']),('数量', '先交付1条主版本；现有同类变体均附作参考。按钮/工具各共用1条。'),('交付建议',f"SFX_{spec['cue']}_v02.wav；48kHz、单声道、24-bit PCM，不循环。"),('视频',f"{p['seconds']} 秒 / 60 fps / {p['frame_count']}帧；F0000起，区间含起不含止。"),('时间含义','实际起音来自录音波形；建议落点取最近视频帧。捕获帧可能重复，不等同于引擎内部帧。'),('试配','HTML中导入新WAV即可按本类落点试配；可保留随片声音。导出评审HTML保存成品与备注。'),('状态','制作/调整需求；新音效尚未制作或接入，最终听感待人工确认。')]:overview.append(row)
    timeline=book.create_sheet('时间轴需求');timeline.append(['编号','画面/声音','实际起音秒（约）','起音所在帧','建议落点帧','音频文件','本包角色','播放规则'])
    for i,e in enumerate(sorted(p['events'],key=lambda x:x['seconds'])):timeline.append([f"{spec['number']:02}-{i+1:02}",spec['scene'] if e['primary'] else e['cue'],round(e['seconds'],4),e['frame_contains_start'],e['target_frame'],e['file'],'本类音效' if e['primary'] else '随片参考',spec['rule'] if e['primary'] else '保留当前触发位置，作为对照。'])
    sheet=book.create_sheet('音频清单');sheet.append(['文件','角色','时长秒','格式','作者','授权','来源','SHA256'])
    for r in records:
        sheet.append([r['name']+'.wav',r['role'],r['seconds'],f"{r['sample_rate']}Hz / {r['channels']}ch / {r['pcm_bits']}-bit",r['author'],r['license'],r['source'],r['sha256']]);sheet.cell(sheet.max_row,1).hyperlink=r['file'];sheet.cell(sheet.max_row,7).hyperlink=r['source']
    widths=[[22,110],[15,36,23,17,18,30,18,65],[31,36,14,27,34,15,53,68]]
    for s,cols in zip(book,widths):
        s.freeze_panes='B2';s.sheet_view.showGridLines=False
        if s!=overview:s.auto_filter.ref=s.dimensions
        for i,width in enumerate(cols,1):s.column_dimensions[get_column_letter(i)].width=width
        for row in s:
            s.row_dimensions[row[0].row].height=52 if s==overview else 60
            for c in row:
                c.font=Font(name='Microsoft YaHei',size=11,color='26343D',bold=c.row==1);c.alignment=Alignment(wrap_text=True,vertical='center');c.border=Border(bottom=Side(style='hair',color='DDE3E6'))
                if c.row==1:c.fill=PatternFill('solid',fgColor='26343D');c.font=Font(name='Microsoft YaHei',size=11,color='FFFFFF',bold=True)
                elif c.row%2==0:c.fill=PatternFill('solid',fgColor='F0F5F3')
        s.page_setup.orientation='landscape';s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0
    path=package/'03_时间轴需求.xlsx';book.save(path);check=load_workbook(path,read_only=True);assert len(check.sheetnames)==3;check.close()


def docs(spec,p,records,package):
    table=['编号 | 秒（约） | 起音帧 | 文件 | 角色']
    for i,e in enumerate(sorted(p['events'],key=lambda x:x['seconds'])):table.append(f"{i+1:02} | {e['seconds']:.4f} | F{e['frame_contains_start']:04} | {e['file']} | "+('本类音效' if e['primary'] else '随片参考'))
    (package/'04_时间轴速览.txt').write_text(f"{spec['title']}音效需求 v001\n{p['seconds']}秒 / 60fps / {p['frame_count']}帧\n\n{spec['rule']}\n不触发：{spec['silent']}\n\n"+'\n'.join(table)+'\n',encoding='utf-8-sig')
    author_lines='\n'.join(f"- {r['name']}.wav：{r['author']}，{r['license']}；{r['role']}。" for r in records)
    reused='按钮两段来自已验收工具样包的实际鼠标操作；重新剪接与标注。' if spec['cue']=='Button' else '当前纯2D实机画面；用临时开发场景准备条件，并通过现有UI动作入口触发演出，未更改游戏运行代码。'
    (package/'00_先看这里.md').write_text(f"""# {spec['title']} · 音效制作需求 v001

双击 `00_音效需求.html` 查看。HTML内嵌视频、音频和需求，单独发送即可离线查看；可以导入新WAV试配并导出包含成品、版本及备注的评审HTML。

制作方向：{spec['brief']} 建议 {spec['duration'][0]}–{spec['duration'][1]} 秒，单次、不循环。先交付一条主版本；现有同类变体作为参考，不扩展成新音效类别。

触发：{spec['rule']}

保持安静：{spec['silent']}

建议交付 `SFX_{spec['cue']}_v02.wav`，48kHz、单声道、24-bit PCM。参考文件为48kHz/单声道/16-bit，由原始素材处理，不是24-bit录音母带。新音效尚未制作、接入或主观听审。

## 视频和时间

{p['seconds']}秒，1280×720，60fps固定输出，{p['frame_count']}帧。F0000起；帧F对应[F/60,(F+1)/60)秒。区间含起点不含终点。实际录制可能包含重复帧，不等同于引擎内部逻辑帧。

{reused} 声轨为游戏混音器实录，保留静音并以采集时钟对齐；时间轴以参考波形匹配的实际起音标注，建议落点取最近视频帧并结合画面核对。没有以点击命令发送时间冒充发声时间。重剪、变速或修改帧率后须重新标注。

有声与无声版使用相同的视频流。网页新音试配会静音原声轨，用新WAV替换本类落点，并可重放随片参考音频；这是网页试配，不是游戏接入结果。原始成品字节会随导出文件完整保留。

## 文件与来源

`03_时间轴需求.xlsx`含制作需求、逐次发声和素材清单。`04_时间轴速览.txt`为简版。WAV、原始源文件、作者发布页快照和授权记录均在本包。

{author_lines}

以上授权只适用于对应参考音频，不覆盖游戏画面或后来导入的成品。来源URLs及处理参数见`来源与授权/音频清单.json`。
""",encoding='utf-8')


def verify(spec,p,package):
    voiced=package/'01_演示_现有效果.mp4';silent=package/'02_演示_无声版.mp4'
    elementary=[ff(['-i',path,'-map','0:v:0','-c','copy','-bsf:v','h264_mp4toannexb','-f','h264','-']) for path in [voiced,silent]]
    assert hashlib.sha256(elementary[0]).digest()==hashlib.sha256(elementary[1]).digest()
    frames=[line for line in ff(['-i',voiced,'-map','0:v:0','-f','framemd5','-']).decode().splitlines() if line and not line.startswith('#')]
    assert len(frames)==p['frame_count'],(spec['cue'],len(frames))
    final_wav=OUT/('build-'+spec['cue'])/'final-audio.wav';ff(['-i',voiced,'-vn','-c:a','pcm_s16le',final_wav]);data=audio(final_wav)
    # Asset identification is verified against the original mixer recording.
    # Check the AAC export against the complete pre-encode PCM mix; fitting each
    # overlapping source again would confuse codec residue with source identity.
    original=audio(OUT/('build-'+spec['cue'])/'mix.wav').mean(axis=1)
    exported=data.mean(axis=1);size=min(len(original),len(exported))
    x=original[:size:4];y=exported[:size:4]
    export_correlation=float(np.dot(x,y)/np.sqrt(np.dot(x,x)*np.dot(y,y)))
    assert export_correlation>.99,(spec['cue'],export_correlation)
    files=[{'file':path.relative_to(package).as_posix(),'bytes':path.stat().st_size,'sha256':sha(path)} for path in sorted(package.rglob('*')) if path.is_file() and path.name!='05_文件核验.json']
    result={'package':folder(spec),'primaryCue':spec['cue'],'video':{'width':1280,'height':720,'fps':60,'frame_count':p['frame_count'],'seconds':p['seconds'],'same_encoded_video_stream':True},'audio_export_correlation':export_correlation,'sound_events':p['events'],'files':files}
    write(package/'05_文件核验.json',result)
    return result


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--cue',default='');parser.add_argument('--start-number',type=int,default=1);parser.add_argument('--analysis-only',action='store_true');args=parser.parse_args()
    for spec in SPECS:
        if (args.cue and args.cue!=spec['cue']) or spec['number']<args.start_number:continue
        if spec.get('reference_only'):
            if spec['number']>=16:
                if args.analysis_only:
                    print(spec['cue'],'reference dub; see supplemental frame evidence',flush=True)
                else:
                    from build_extra_sfx_packages import build as build_extra,load as load_extra,OUT as extra_out
                    build_extra(spec,load_extra(extra_out/'edit-list.json')[spec['cue']])
                continue
            if args.analysis_only:
                print(spec['cue'],'proposed soundtrack; see CardDraw delivery plan',flush=True)
            else:
                from build_card_draw_requirement import build
                build()
            continue
        p=plan(spec);write(OUT/(spec['cue']+'-delivery-plan.json'),p)
        if args.analysis_only:
            print(spec['cue'],p['seconds'],[(e['file'],round(e['seconds'],3),round(e['correlation'],3)) for e in p['events']],flush=True);continue
        package=DEST/folder(spec);package.mkdir(parents=True,exist_ok=True)
        media(spec,p,package);records=sources(spec,p,package);sheets(spec,p,records,package);docs(spec,p,records,package);verify(spec,p,package)
        print(json.dumps({'cue':spec['cue'],'seconds':p['seconds'],'events':len(p['events']),'files':len(records),'status':'media_and_brief_ready'},ensure_ascii=False),flush=True)

if __name__=='__main__':main()
