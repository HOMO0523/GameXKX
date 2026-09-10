"""Export explicit design state for the relic/hunt expansion without marking it implemented."""
import json
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook

from update_ui_goal_workbooks import ROOT, prepare, finish


def load(path):
    return json.loads((ROOT / path).read_text(encoding="utf-8"))


def save_verified(book, path, expected):
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.stem + ".pending.xlsx")
    book.save(temporary)
    check = load_workbook(temporary, read_only=True, data_only=False)
    actual = {sheet.title: sheet.max_row - 4 for sheet in check}
    for name, count in expected.items():
        assert actual[name] == count, (name, actual[name], count)
    check.close()
    temporary.replace(path)
    return {"path": str(path), "rows": actual, "validation": "reloaded-and-counted", "runtimeVerification": "pending"}


def relic_book():
    design = load("docs/design/2026-09-10-relic-redesign/relic-design.json")
    progress = {row["id"]: row for row in load("docs/production/2026-09-10-ui-guidance-localization-state.json")["milestones"]}["G5R"]
    common = load("docs/design/2026-09-10-relic-redesign/common-relic-translations.json")["relics"]
    rows = common + design["relics"]
    assert Counter(row["quality"] for row in rows) == {"Common": 16, "Rare": 20, "Epic": 10}
    assert len({row["id"] for row in rows}) == 46
    art = {row["slug"]: row for row in load("SourceArt/UI/Relics/gemstyle-20260910/manifest.json")["icons"]}
    book = Workbook()
    ws = book.active
    ws.title = "01_遗物总表"
    prepare(ws, "GameXXK · 遗物设计总表", "16普通/特殊 + 20稀有 + 10史诗；" + progress["implementation"] + "。完整实牌和平衡验收未完成。",
            ["遗物ID", "名称（中文）", "Name (English)", "品质", "中文短描述", "English description", "触发", "主数值", "次数/次参数", "配合方向", "实现状态", "验证状态", "图标状态", "图标源路径"],
            [31, 19, 31, 13, 68, 94, 22, 15, 18, 48, 35, 30, 24, 85])
    qualities = {"Common": "普通/特殊", "Rare": "稀有", "Epic": "史诗"}
    for row in rows:
        icon = art.get(row["slug"])
        high = row["quality"] != "Common"
        ws.append([row["id"], row["name"]["zh-Hans"], row["name"]["en"], qualities[row["quality"]],
                   row["description"]["zh-Hans"], row["description"]["en"], row.get("trigger", "沿用当前规则"),
                   row.get("primary", "见原规则"), row.get("secondary", "见原规则"), row.get("synergy", "普通机制保持"),
                   progress["implementation"] if high else "保持既有运行时",
                   progress["validation"] if high else "普通防护/回复相邻回归通过；完整验收待完成",
                   "已导入，重启加载通过" if icon and icon.get("imported") else "已整理，未导入" if icon else "沿用特殊护符" if row["slug"] == "LifeSavingTalisman" else "重绘待完成",
                   icon["icon"] if icon else ""])
    finish(ws)
    rules = book.create_sheet("02_实现与验收边界")
    prepare(rules, "遗物运行合同", "详细依据：docs/design/2026-09-10-relic-redesign/runtime-contract.md",
            ["类别", "要求", "状态"], [24, 130, 30])
    for category, text in [
        ("计数与兼容", "保留旧31个稳定ID，普通/特殊16件不改机制；重做15件并新增15件高品质，目标目录46。"),
        ("触发证据", "核对真实治疗/护甲/伤害/蓄力/地势变化；不让遗物新增的结果再次触发其他遗物。"),
        ("次数与存档", "每回合/每人限次需要可存档账本；重放、读档或重复回调不得刷新次数。"),
        ("真实反应", "格挡与反击必须创建已有反应系统记录，不能只增加状态图标。"),
        ("完整结算", "覆盖弃牌/洞察/检索等待选后续、怪物阶段与死亡目标，使用实际战斗入口。"),
        ("数值验证", "按项目当前等级/队伍/关卡进行同种子有无遗物对照，不使用近似数值模型冒充实战。"),
        ("美术与本地化", "45个不同物件按宝石风格重绘；图标/卡面/短描述/tooltip/品质Shader与中英文本一致。")]:
        rules.append([category, text, "待完整验证"])
    finish(rules)
    return save_verified(book, ROOT / "docs/design/2026-09-10-relic-redesign/GameXXK_遗物设计总表_2026-09-10.xlsx",
                         {"01_遗物总表": 46, "02_实现与验收边界": 7})


def hunt_book():
    data = load("docs/design/2026-09-10-hunt-stage-design/hunt-design.json")
    accepted = data["userApproved"]
    book = Workbook()
    stages = book.active
    stages.title = "01_讨伐关卡"
    prepare(stages, "GameXXK · 三难度3-4讨伐设计", "用户需求及已确认规则；当前未接入运行时。保留原27关ID与等级，不通过改9关常量把旧难度等级整体移位。",
            ["Stage ID", "名称", "对照关卡", "战斗等级", "连续三章参考", "门票ID", "门票品质", "生命倍率", "金币倍率", "经验倍率", "胜利耗令", "失败耗令", "必得讨伐箱", "状态"],
            [28, 20, 28, 14, 63, 30, 16, 15, 15, 15, 15, 15, 18, 32])
    for row in accepted["stages"]:
        stages.append([row["id"], row["name"]["zh-Hans"], row["sourceStage"], row["combatLevel"], " → ".join(row["routeReferences"]),
                       row["ticketId"], row["ticketQualityName"], 1.5, 1.5, 1.5, 1, 0, 1, "设计已登记，未接入/未验证"])
    finish(stages)
    odds = book.create_sheet("02_宝箱品质概率")
    prepare(odds, "已批准 · 宝箱品质概率", "装备/宝石分支的条件品质概率；各高品质档精确4倍。先判定门票替代，因此不是每个箱子的无条件出货概率。",
            ["品质", "高级箱概率", "讨伐箱概率", "概率倍数", "批准状态", "实现状态"], [24, 25, 25, 23, 30, 35])
    names = {"Rare": "稀有", "Epic": "珍稀", "Legendary": "传奇", "Immortal": "不朽"}
    tables = accepted["chestQualityWeights"]
    assert sum(tables["Advanced"].values()) == sum(tables["Hunt"].values()) == 10000
    for key in ("Rare", "Epic", "Legendary", "Immortal"):
        advanced, hunt = tables["Advanced"][key], tables["Hunt"][key]
        if key != "Rare": assert hunt == advanced * 4
        odds.append([names[key], advanced / 10000, hunt / 10000, hunt / advanced, "用户已批准", "尚未接入运行时"])
    finish(odds)
    for row in range(5, 9):
        odds.cell(row, 2).number_format = odds.cell(row, 3).number_format = "0.00%"
        odds.cell(row, 4).number_format = '0.00"×"'
    tickets = book.create_sheet("03_门票与奖励规则")
    prepare(tickets, "讨伐令与结算", "来源难度与替代奖励均已得到用户确认；实现和验证状态另列。",
            ["规则", "设计值", "状态", "验证要求"], [29, 100, 30, 95])
    ticket_rows = [
        ["3-3首通", "每个难度首次挑战通关3-3，必得对应品质讨伐令1枚", "用户已要求", "首次/重复通关与旧存档策略"],
        ["普通箱令概率", "2%", "用户已要求", "固定种子统计、背包满原子失败"],
        ["高级箱令概率", "8%", "用户已要求", "固定种子统计、背包满原子失败"],
        ["掉出令的品质", "按取得时保存的箱来源难度：普通→普通令，困难→珍稀令，地狱→不朽令。商店箱记录购买时已解锁的所选难度；旧来源不明按普通。", "用户已批准", "切换开箱时难度不得改变来源品质"],
        ["令额外/替代掉落", "先判定2%/8%；命中时仅给1令，替代原装备/宝石/材料；未命中才进入原类别池。", "用户已批准", "命中不可同时发原奖励，满包失败不耗箱"],
        ["无票禁入", "没有对应品质令，挑战和游历均不可进入", "用户已要求", "UI禁用及运行时双重校验"],
        ["挑战胜利", "消费1令，必得1个红金讨伐宝箱", "用户已要求", "重复结算/读档不得重复扣票或发奖"],
        ["挑战失败", "不消耗讨伐令", "用户已要求", "失败、退出、异常恢复的原票保护"],
        ["讨伐游历", "需要对应令，消耗1令且必得讨伐宝箱", "用户已要求", "票耗尽停止循环，离线与在线结算一致"],
        ["HUD第三箱", "现有宝箱右侧增加讨伐箱区域，接入相同左右键", "用户已要求", "展开/折叠/三档缩放/计数与开箱"],
        ["困难/地狱强度", "困难约1.5倍，地狱约2倍；避免重复叠加旧倍率", "用户已要求，作用字段核查中", "挑战、游历、意图、tooltip与怪物总表一致"]
    ]
    for row in ticket_rows: tickets.append(row)
    finish(tickets)
    return save_verified(book, ROOT / "docs/design/2026-09-10-hunt-stage-design/GameXXK_讨伐与宝箱设计总表_2026-09-10.xlsx",
                         {"01_讨伐关卡": 3, "02_宝箱品质概率": 4, "03_门票与奖励规则": len(ticket_rows)})


def main():
    reports = [relic_book(), hunt_book()]
    output = ROOT / "docs/design/2026-09-10-project-progress/design-workbook-report.json"
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps({"workbooks": reports}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(reports, ensure_ascii=False))


if __name__ == "__main__":
    main()
