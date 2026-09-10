"""Record observed results, preserving the larger unfinished goal and design budgets."""
from pathlib import Path
from datetime import datetime
import json
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'Saved/Codex/UIGuidanceLocalization-20260910'
def load(path):return json.loads(path.read_text(encoding='utf-8-sig'))
def main():
    core=load(OUT/'wind-affix-final-tests.json')['results']
    adjacent=load(OUT/'wind-adjacent-tests.json')['results']
    assert core['total']==5 and core['failed']==0 and core['passed']==5
    assert adjacent['total']==69 and adjacent['failed']==0 and adjacent['passed']==69
    source=ROOT/'docs/production/2026-09-10-ui-guidance-localization-state.json'
    state=load(source);by={m['id']:m for m in state['milestones']}
    by['G4W'].update(validation='核心5/5、相邻69/69通过：逐张与同名叠加、火/冰/雷/物理/重箭、重放不计数/回合重置、两件实际投影、十品质生成洗炼、旧档和待选结果转换。',
        implementation='乘风固定1%已接真实伤害。5项随机资源词缀退休；旧装备定向转换，强化/孔/归属/其他词缀保留。洗炼使用同一候选池决定启用，窄行短标签、完整白字江湖体Tooltip。',
        progress='乘风及产甲已生效；其余24项分系词缀的未实现状态保留在总表，不冒充已接入。',next='实际工具窄行与完整Tooltip视觉复核；其余词缀独立继续审查。')
    by['G4R'].update(validation='逐箱报告检查通过：批量30箱与逐次30箱内容次序相同、30条真实UI历史、中文生成后切英文、保存失败无虚假报告且宝箱保留。',
        implementation='每个成功开箱产生稳定收据和独立报告；按当前语言重新构建具体名称/数量/品质/等级/去向，长信息放悬停。',
        progress='早期中文残留告警实际为测试遍历临时FText产生的悬空字符串误报，已纠正；历史仍沿用200条界面消息，不属于跨重启日志。',next='消息区真实悬停视觉复核；跨重启日志随完整存档目标继续。')
    by['G3']['validation']='69项回归包含23步中英真实按钮/悬停/阅读流程、继续与重学；半透明白字已取得实机截图，完整视觉复核尚未完成。'
    by['G3A']['validation']='独立Academy六项在69项相邻检查中通过，覆盖13课程29节目标与真实获胜；新版教学界面实拍仍待完成。'
    state['updatedAt']=datetime.now().isoformat(timespec='seconds')
    state['runtimeControl'].update(status='GUI编辑器18765/PIE运行。短暂界面检查已结束，当前玩家状态按新建的本次快照严格恢复，Dev会话关闭，中文100%。',
        buildLog='Saved/Codex/UIGuidanceLocalization-20260910/wind-test-final-build.log',
        latestCore='Saved/Codex/UIGuidanceLocalization-20260910/wind-affix-final-tests.json',
        latestAdjacent='Saved/Codex/UIGuidanceLocalization-20260910/wind-adjacent-tests.json',
        currentReviewBaseline='Saved/Codex/UIGuidanceLocalization-20260910/ui-review/review-original.json')
    source.write_text(json.dumps(state,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    file=ROOT/'docs/design/2026-09-04-project-design-tables/GameXXK_装备设计总表_2026-09-04.xlsx'
    wb=load_workbook(file)
    for row in wb['04_词缀目录'].iter_rows(min_row=2):
        if row[8].value=='Affix.ZhuiFeng.LateCardDamage':row[3].value='真实伤害/生成/洗炼/旧档核心检查通过；实机Tooltip待复核'
    for row in wb['04A_触发与接入审查'].iter_rows(min_row=2):
        if row[8].value=='Affix.ZhuiFeng.LateCardDamage':row[3].value='已接入；实际伤害矩阵通过';row[7].value='核心5/5与相邻69/69'
    for row in wb['04H_追风递增与退休'].iter_rows(min_row=5):
        row[3].value='行为检查通过；视觉另行验收'
    for row in wb['08_实现差异'].iter_rows(min_row=2):
        if row[0].value=='分系随机词缀接入':
            row[1].value='当前26个分系随机项；产甲和乘风已接入并验证，其余24项待接入。';row[3].value='2项已验证；24项待接入'
    for row in wb['04D_叠加上限与边际'].iter_rows(min_row=2):
        if row[0].value=='追风乘风例外':row[3].value='用户确认；实际出牌/叠加/回合检查通过'
    tmp=file.with_suffix('.pending.xlsx');wb.save(tmp);check=load_workbook(tmp,read_only=True);assert '04H_追风递增与退休' in check;check.close();tmp.replace(file)
    doc=ROOT/'docs/design/2026-09-11-affix-wind-ramp.md'
    text=doc.read_text(encoding='utf-8').replace('状态：实现与验证中；','状态：代码与定向行为验证通过，实际界面复核未完成；')
    marker='\n## 当前证据\n'
    if marker in text:text=text.split(marker)[0]
    text+=marker+'''
- 冷UBT：首轮2337.78秒、收尾187.02秒、报告结构补编298.10秒、最终测试补编92.69秒，均成功。
- 核心5/5、相邻69/69通过（两组有重复项，不作为74个独立测试计数）。
- 真实存档31份在隔离自动化后哈希不变；图形试玩后又创建本次新快照保护并严格恢复，不能恢复早前392730金币/4箱基线。
- 实机洗炼窄行、报告悬停和角色教学外观仍待复核。一般窄行显示“连打直伤 +1% / Chain DMG +1%”，完整限定条件在Tooltip；字号16测宽中文138、英文151，低于214内容宽。
- 初轮伤害失败来自自动目标/纯防护牌的测试前提；迁移失败来自未初始化三人队伍。报告“中文残留”经完整输出确认是测试遍历临时FText的生命周期误报，已改为先复制FString再检查。
- 开箱历史目前为界面内最近200条消息，跨重启日志仍属于未完成的完整挂机存档目标。
'''
    doc.write_text(text,encoding='utf-8')
    print(json.dumps({'core':core['passed'],'adjacent':adjacent['passed'],'overallGoal':'unfinished','visual':'pending'}))

if __name__=='__main__':main()
